#!/usr/bin/env python3
"""Check the web page (``ratesplot/web.py``): its charts against the command line's, and how it behaves.

Two parts, both run by default:

* ``charts``: for a few regression cases, the page's own drawing is hashed
  against the command line's. The page reads the case as it reads its
  address (``cli.sort_tokens``, then ``frontend.starting_choices``) and turns
  its controls' choices into a config (``web._config_for``), which must draw
  what the command line's does (``options.drawn_config`` for each nation); then it draws with ``web._draw``. The command line
  renders the same case through ``tools/verify_charts.run`` into
  ``out/verify_web/NAME/``. Each chart is IDENTICAL or DIFFER: on the same
  machine the page must draw, byte for byte, what the command line and the
  window draw.
* ``page``: Streamlit's ``AppTest`` runs ``streamlit_app.py`` headless in this
  process and drives it as a visitor would: a chart in the address, the rules
  between controls, a bad field, Reset, a bad address, the date buttons with
  Back and Unzoom, the calendars, the size boxes and their aspect lock. Each
  check prints ``ok`` or ``FAIL``.

A drag on a chart is not tested here: AppTest runs no JavaScript, so the chart
component never sends a box. The rule a box goes through,
``frontend.zoom_updates``, is the window's own; test a drag in a browser. So
is the hover read-out (batch 11.2): here, the data the chart component gets
after its PNG, unpacked as the page's script unpacks them, against the lines
the window reads (``hover.chart_lines``); its rules are the window's
(``tools/verify_gui.py``, group ``hover``); look at it in a browser.

Downloads come from the regression tool's disk cache (``--cache``, default
``out/download_cache``, shared with ``tools/regress_charts.py``), and the
charts part pins ``--e:2026-09-01`` as that tool does. Several page checks
draw the default charts, which end today: their intraday quotes are fetched
live, so the page part needs the network.

Usage (from the project root)::

    python tools/verify_web.py                 # both parts (a few minutes once the cache is warm)
    python tools/verify_web.py --only charts   # or --only page
    python tools/verify_web.py --list          # the chart cases
    python tools/verify_web.py --synthetic     # made-up data, no network (tools/synthetic_sources.py)

Exit code 1 on any DIFFER or FAIL.
"""

from __future__ import annotations

import argparse
import contextlib
import datetime
import hashlib
import io
import json
import logging
import re
import struct
import sys
import urllib.parse
import zlib
from dataclasses import replace
from pathlib import Path
from typing import Callable

# The page's buttons are labelled "◀ Back" and "Later ▶"; a console or pipe
# in the Windows code page cannot encode them. Escape what it cannot show
# rather than fail. (Before ratesplot is imported: its output router writes
# through the stream that is sys.stdout now.)
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="backslashreplace")

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / "tools")]

import verify_charts  # noqa: E402  (first: it sets the Agg backend before pyplot is imported)
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
import streamlit.testing.v1.app_test as app_test_module  # noqa: E402
import synthetic_sources  # noqa: E402
from regress_charts import CASES, PINNED_END, install_disk_cache  # noqa: E402
from streamlit.runtime.media_file_manager import MediaFileManager  # noqa: E402
from streamlit.runtime.memory_media_file_storage import MemoryMediaFileStorage  # noqa: E402
from streamlit.testing.v1 import AppTest  # noqa: E402

import ratesplot  # noqa: E402
from ratesplot import cli, frontend, hover, http, options, plotting, web  # noqa: E402
from ratesplot.cli import sort_tokens  # noqa: E402
from ratesplot.config import NATIONS, WEB_PAGE_ADDRESS, PlotConfig  # noqa: E402
from ratesplot.frontend import move_dates_updates, starting_choices, version_text  # noqa: E402
from ratesplot.options import command_line_tokens, config_from_choices, drawn_config  # noqa: E402

# AppTest touches Streamlit's state before each page's first run, outside any
# script run, and Streamlit logs "missing ScriptRunContext! This warning can be
# ignored when running in bare mode" each time, which it is here. A filter,
# not a level: Streamlit resets its loggers' levels when a page runs.
logging.getLogger("streamlit.runtime.scriptrunner_utils.script_run_context").addFilter(
    lambda record: "missing ScriptRunContext" not in record.getMessage()
)

APP = str(ROOT / "streamlit_app.py")
TIMEOUT_SECONDS = 180  # per AppTest run: a drawing whose downloads are not cached yet takes about 15 s


class _KeptMediaStorage(MemoryMediaFileStorage):
    """AppTest's store of a run's files (the download buttons' data), kept for reading after the run.

    AppTest makes a new store for each run and drops the runtime that holds
    it when the run ends, so a download button's file could not be read
    afterwards (Streamlit 1.64): each new store is remembered here instead.
    """

    last: MemoryMediaFileStorage | None = None

    def __init__(self, media_endpoint: str) -> None:
        super().__init__(media_endpoint)
        _KeptMediaStorage.last = self


app_test_module.MemoryMediaFileStorage = _KeptMediaStorage


class _KeptMediaFileManager(MediaFileManager):
    """AppTest's manager of a run's files, kept as its store is: it holds the downloads made only on a click (batch 11.1)."""

    last: MediaFileManager | None = None

    def __init__(self, storage) -> None:
        super().__init__(storage)
        _KeptMediaFileManager.last = self


app_test_module.MediaFileManager = _KeptMediaFileManager

# Regression cases drawn both ways: the defaults, each measure, levels, -l,
# --reg with and without a tolerance, the history to 1867, a small canvas.
CHART_CASES = (
    "default", "c2015", "u_reg_r", "c_lv_r", "c_l_lv", "u_reg_tol3", "c_reg1867", "c_p", "c1100",
    # Batch 1 (2026-10-06): --no-yields read from the address, and the yield axis's other curves.
    "u_noy", "all_rates",
    # Batch 2: values for one nation's chart.
    "nat_limits", "nat_terms",
    # Batch 3: the UK, alone with its curves, with its own values, and all three nations.
    "uk_rates", "uk_nat", "all3",
    # Batch 4: Germany, alone with its curves, with its own values, and all four nations.
    "de_rates", "de_nat", "all4",
    # Batch 9.1: both policy terms on every nation's chart, the U.S. target range a band.
    "policy_both",
    # Batch 10.4: the UK's net debt alone. Batch 10.5: each nation's own levels.
    "uk_net_only", "nat_levels",
    # Batch 18.3: Top and Bottom in U.S. dollars on all four charts.
    "nat_fx",
)

NATION_KEYS = tuple(nation.key for nation in NATIONS)  # in drawing order, as the page's tabs
VIEW_BUTTONS = ("◀ Back", "Unzoom", "◀ Earlier", "Later ▶", "Zoom out")

Check = Callable[..., None]


# ---------------------------------------------------------------------------
# Part 1: the page's charts against the command line's
# ---------------------------------------------------------------------------


def sha256(data: bytes) -> bytes:
    return hashlib.sha256(data).digest()


def check_charts(cases: tuple[str, ...], outdir: Path, verbose: bool = False) -> bool:
    """Draw each case as the page does and as the command line does; True if every chart is identical."""
    all_same = True
    for name in cases:
        tokens = [PINNED_END, *CASES[name]]
        payloads, flags, unmatched, misspelled = sort_tokens(tokens)
        if unmatched or misspelled:
            print(f"{name:12} FAIL: the case does not parse ({unmatched + misspelled})")
            all_same = False
            continue
        # As the page does: the address laid over the defaults, then the controls' choices to a config.
        config = config_from_choices(payloads, flags)
        choices, _notes = starting_choices(config, payloads, flags, {})
        page_config, problems = web._config_for(choices)
        # Every nation drawn by default (as before batch 3), and any other the case draws: a
        # nation drawn only when asked for, and not asked for, keeps fields the page does not write.
        compared = [n.key for n in NATIONS if n.shown_by_default or getattr(config, n.show_field)]
        if page_config is None or any(drawn_config(page_config, k) != drawn_config(config, k) for k in compared):
            print(f"{name:12} FAIL: the page's config differs from the command line's ({problems})")
            all_same = False
            continue
        drawing = web._draw(page_config, tuple(command_line_tokens(*choices)))
        if drawing.error is not None or drawing.failed:
            print(f"{name:12} FAIL: the page's drawing failed: {drawing.error or drawing.failed}")
            all_same = False
            continue
        output = io.StringIO()
        with contextlib.nullcontext() if verbose else contextlib.redirect_stdout(output):
            written = verify_charts.run(outdir / name, tokens)
        keys = [key for key in NATION_KEYS if key in drawing.pngs]
        verdicts = []
        if len(keys) != len(written):
            verdicts.append(f"FAIL: the page drew {len(keys)} chart(s), the command line {len(written)}")
            all_same = False
        for key, path in zip(keys, written):
            same = sha256(drawing.pngs[key]) == sha256(path.read_bytes())
            verdicts.append(f"{key}:{'IDENTICAL' if same else 'DIFFER'}")
            all_same &= same
        print(f"{name:12} {' '.join(CASES[name]):28} {' '.join(verdicts)}  (page {drawing.seconds:.1f} s)")
        # The command line's warnings (data coverage, fallbacks), as tools/regress_charts.py echoes them.
        for line in output.getvalue().splitlines():
            if line.strip().startswith("Warning:") and "requested start" not in line:
                print(f"    {line.strip()}")
    return all_same


# ---------------------------------------------------------------------------
# Part 2: the page's behaviour, through Streamlit's AppTest
# ---------------------------------------------------------------------------


@contextlib.contextmanager
def topic_taken_away(name: str):
    """Option ``name`` with no topic for the while: a control with nothing more to say, now that every control names one (batch 29.3).

    ``Option`` is frozen, so its field is set past the dataclass's guard and
    put back after. The page's script reads the same ``options`` module, so
    a page run inside the block lays the control out with no More.
    """
    option = options.by_name(name)
    topic = option.topic
    object.__setattr__(option, "topic", "")
    try:
        yield
    finally:
        object.__setattr__(option, "topic", topic)


def new_page(chart: str | None = None, heard: str | None = "") -> AppTest:
    """A fresh visitor's page, with ``chart`` in its address (None: no address), run once.

    ``heard``: the cookie's value as the page's script would send it (batch
    24.6, ``web._READER_JS``), encoded; "" for nothing remembered. AppTest runs
    no script, so it is given here; None gives none, and a page with no
    chart to draw waits for it.
    """
    at = AppTest.from_file(APP, default_timeout=TIMEOUT_SECONDS)
    if chart is not None:
        at.query_params["chart"] = chart
    if heard is not None:
        at.session_state["heard"] = heard
    at.run()
    return at


def flags(at: AppTest) -> dict[str, bool]:
    return {box.key.split(":", 1)[1]: box.value for box in at.checkbox if box.key and box.key.startswith("flag:")}


def chart_elements(at: AppTest) -> list:
    """The page's chart components (a ``bidi_component`` to AppTest, as the settings' cookie writer is, batch 24.6)."""
    return [element for element in at.get("bidi_component") if "-chart:" in element.proto.id]


def charts_shown(at: AppTest) -> int:
    """The charts shown: each is the page's chart component."""
    return len(chart_elements(at))


def memory(at: AppTest) -> dict | None:
    """What the page's cookie writer was given in its last run (batch 24.6): its cookie's name, line and seconds; None if not laid out."""
    found = [element for element in at.get("bidi_component") if element.proto.id.endswith("-memory")]
    return json.loads(found[0].proto.json) if found else None


def reader_asked(at: AppTest) -> bool:
    """True if the page's last run laid out its cookie reader (batch 24.6), waiting for the browser's answer."""
    return any(element.proto.id.endswith("-reader") for element in at.get("bidi_component"))


def sidebar_keeper(at: AppTest) -> dict | None:
    """What the sidebar width's keeper was given in the page's last run (batch 31.1); None if not laid out."""
    found = [element for element in at.get("bidi_component") if element.proto.id.endswith("-sidebar_width")]
    return json.loads(found[0].proto.json) if found else None


def button(at: AppTest, label: str):
    return next(b for b in at.button if b.label == label)


def downloaded(at: AppTest, label: str) -> tuple[str, str, bytes]:
    """The file a download button gives, from the page's last run: (file name, type, bytes)."""
    found = next(b for b in at.get("download_button") if b.proto.label == label)
    file = _KeptMediaStorage.last.get_file(found.proto.url.rsplit("/", 1)[1])
    return file.filename, file.mimetype, file.content


def clicked_download(at: AppTest, key: str) -> tuple[str, str, bytes]:
    """The file a download button made on the click gives (``key``: its widget key), from the page's last run.

    AppTest cannot click a download button; this does what Streamlit does
    on the click (``MediaFileManager.execute_deferred``: the page's callable
    run, its bytes stored under the button's file name and type).
    """
    found = next(b for b in at.get("download_button") if b.proto.id.endswith(f"-{key}"))
    url = _KeptMediaFileManager.last.execute_deferred(found.proto.deferred_file_id)
    file = _KeptMediaStorage.last.get_file(url.rsplit("/", 1)[1])
    return file.filename, file.mimetype, file.content


def enabled(at: AppTest) -> dict[str, bool]:
    return {b.label: not b.disabled for b in at.button if b.label in VIEW_BUTTONS}


def address(at: AppTest) -> str:
    """The chart's command line in the page's address (AppTest may hold a query value as a list)."""
    value = at.query_params.get("chart", "")
    return value[0] if isinstance(value, list) else value


def png_size_bytes(width: int, height: int) -> bytes:
    """A PNG header's width and height (bytes 16-24), for comparing with a drawn chart."""
    return width.to_bytes(4, "big") + height.to_bytes(4, "big")


def _is_date(placeholder: str) -> bool:
    """``YYYY-MM-DD``, any date (the automatic start's year follows config.DEFAULT_START_FLOOR)."""
    try:
        datetime.date.fromisoformat(placeholder)
    except ValueError:
        return False
    return True


def _is_number(placeholder: str, suffixes: str = "") -> bool:
    """A number as the zoom writes it, with one of ``suffixes`` allowed at its end ("2.85t", "176%")."""
    try:
        float(placeholder.rstrip(suffixes) if suffixes else placeholder)
    except ValueError:
        return False
    return True


def check_address_and_rules(check: Check) -> None:
    """A chart in the address; the rules between controls; a bad field; levels and tolerance; Reset."""
    at = new_page("-c -r --reg --e:2026-09-01")
    check("no exception", not at.exception, at.exception)
    f = flags(at)
    check("address: Canada only, -r, --reg", f["cdn"] and not f["us"] and f["relative"] and f["regression"], f)
    check("GDP box disabled under -r", at.checkbox(key="flag:gdp").disabled)
    check("tolerance field enabled with --reg", not at.text_input(key="value:regression-tolerance").disabled)
    check("one chart drawn", charts_shown(at) == 1, charts_shown(at))
    check("address kept", address(at) == "--C --relative --regression --end:2026-09-01", address(at))
    code = [c.value for c in at.code]
    check("command line shown", any(v.startswith("python ratesplot.py --C --relative") for v in code), code[:1])
    # Blank fields show the value in effect in grey (Terry, 2026-10-08, decision 18).
    placeholder = at.text_input(key="value:start").placeholder
    check("start placeholder: the date found", _is_date(placeholder), placeholder)
    placeholder = at.text_input(key="value:end").placeholder
    check("end placeholder: the date drawn to", placeholder == "2026-09-01", placeholder)
    placeholder = at.text_input(key="value:regression-tolerance").placeholder
    check("tolerance placeholder: the default", placeholder == "0.3", placeholder)
    limits = {name: at.text_input(key=f"value:cdn:{name}").placeholder for name in ("min", "max", "top", "bottom")}
    check(
        "Canada's limits as drawn (yields in percent, -r in % of GDP)",
        all(_is_number(limits[name]) for name in ("min", "max"))
        and all(limits[name].endswith("%") and _is_number(limits[name], "%") for name in ("top", "bottom"))
        and float(limits["min"]) < float(limits["max"])
        and float(limits["bottom"].rstrip("%")) < float(limits["top"].rstrip("%")),
        limits,
    )
    # An unticked nation's section is not laid out (batch 28.1); before, its
    # fields showed "automatic", having no chart of their own.
    shown = {widget.key for widget in at.text_input}
    check("an unticked nation's section not shown", "value:us:max" not in shown and "value:cdn:max" in shown)

    # One measure at a time; changing it clears Top and Bottom.
    at.text_input(key="value:cdn:top").input("200%")
    at.run()
    at.checkbox(key="flag:per-capita").check()
    at.run()
    f = flags(at)
    check("ticking -p unticks -r", f["per-capita"] and not f["relative"], f)
    top = at.text_input(key="value:cdn:top").value
    check("Top cleared by the measure change", top == "", top)
    check("address follows", address(at) == "--C --per-capita --regression --end:2026-09-01", address(at))

    # The last country cannot be unticked.
    at.checkbox(key="flag:cdn").uncheck()
    at.run()
    check("Canada ticked again", flags(at)["cdn"])
    check("notice shown", any("At least one country" in i.value for i in at.info), [i.value for i in at.info])

    # A bad field: not drawn, the previous chart stays, the address is not changed.
    at.text_input(key="value:start").input("20x5")
    at.run()
    check("error for a bad start", any("Not drawn" in e.value for e in at.error), [e.value for e in at.error])
    check("previous chart still shown", charts_shown(at) == 1)
    check("address unchanged", address(at) == "--C --per-capita --regression --end:2026-09-01", address(at))
    at.text_input(key="value:start").input("2001")
    at.run()
    check(
        "drawn again from 2001",
        not at.error and address(at) == "--C --per-capita --regression --start:2001 --end:2026-09-01",
        (address(at), [e.value for e in at.error]),
    )

    # The level boxes (Canada's own since batch 10.5) and the tolerance.
    at.checkbox(key="level:cdn:f").check()
    at.checkbox(key="level:cdn:p").check()
    at.text_input(key="value:regression-tolerance").input("3")
    at.run()
    check("levels and tolerance in the address", "--cdn:debt:fp" in address(at) and "--reg:3" in address(at), address(at))
    at.checkbox(key="flag:regression").uncheck()
    at.run()
    check(
        "tolerance greyed and left out when --reg is off",
        at.text_input(key="value:regression-tolerance").disabled and "reg" not in address(at),
        address(at),
    )

    # Refresh (Terry, 2026-10-08: one button, not Redraw and Reload data) draws the same choices again.
    labels = [b.label for b in at.button]
    check("a Refresh button, no Redraw", "Refresh" in labels and "Redraw" not in labels, labels)
    finished = at.session_state["drawing"].finished
    button(at, "Refresh").click()
    at.run()
    check("Refresh draws again", at.session_state["drawing"].finished != finished)

    # Reset: the defaults, but the nation boxes as they were (batch 17.5).
    button(at, "Reset").click()
    at.run()
    f = flags(at)
    check("reset: Canada still alone, no measure", f["cdn"] and not f["us"] and not f["relative"] and not f["per-capita"], f)
    check("reset: one chart", charts_shown(at) == 1, charts_shown(at))


def check_bad_addresses(check: Check) -> None:
    """An address the program cannot read: reported, and the defaults drawn."""
    bad = new_page("pizza -r:1 --e:2026-09-01")
    warnings = [w.value for w in bad.warning]
    check("bad address reported", any("pizza: not an option" in w and "takes no value" in w for w in warnings), warnings)
    check("defaults drawn", charts_shown(bad) == 2, charts_shown(bad))
    # Contradictory options (-r with -p): reported by config_from_choices.
    both = new_page("-r -p --e:2026-09-01")
    check("-r -p reported", bool(both.warning), [w.value for w in both.warning])


def check_view_buttons(check: Check) -> None:
    """Back, Unzoom and the date buttons, on a pinned window; Reset undone by Back; Later off at today."""
    view = new_page("-c -s:2001 --e:2019")
    expected = {"◀ Back": False, "Unzoom": False, "◀ Earlier": True, "Later ▶": True, "Zoom out": True}
    check("view: nothing to undo; the dates can move", enabled(view) == expected, enabled(view))
    png = chart_elements(view)[0].proto.bytes
    check(
        "the chart component gets the PNG, 2048x1536",
        png[:8] == b"\x89PNG\r\n\x1a\n" and png[16:24] == png_size_bytes(2048, 1536),
        len(png),
    )

    def expect(start: str, end: str, **move: float) -> tuple[str, str, str]:
        """The address a date move should give (by ``frontend.move_dates_updates``), and the new start and end."""
        u = move_dates_updates(pd.Timestamp(start), pd.Timestamp(end), **move)
        return f"--C --start:{u['start']}" + (f" --end:{u['end']}" if u["end"] else ""), u["start"], u["end"]

    e1, s1, n1 = expect("2001-01-01", "2019-01-01", shift=-0.5)
    button(view, "◀ Earlier").click()
    view.run()
    check("Earlier: half a window back", address(view) == e1, (address(view), e1))
    check("Back and Unzoom now possible", enabled(view)["◀ Back"] and enabled(view)["Unzoom"], enabled(view))
    e2, _s2, _n2 = expect(s1, n1, scale=2.0)
    button(view, "Zoom out").click()
    view.run()
    check("Zoom out: twice the span", address(view) == e2, (address(view), e2))
    button(view, "◀ Back").click()
    view.run()
    check("Back: one step", address(view) == e1, address(view))
    e3, s3, n3 = expect(s1, n1, shift=0.5)
    button(view, "Later ▶").click()
    view.run()
    check("Later: half a window on", address(view) == e3, (address(view), e3))
    button(view, "Unzoom").click()
    view.run()
    # Earlier then Later came back to the same dates, so the fields as first
    # typed are the same chart: nothing is redrawn and the address keeps e3.
    fields = (view.text_input(key="value:start").value, view.text_input(key="value:end").value)
    check(
        "Unzoom: the fields as first typed; Unzoom greyed, Back possible",
        fields == ("2001", "2019") and not enabled(view)["Unzoom"] and enabled(view)["◀ Back"],
        (fields, enabled(view)),
    )
    # Unzoom is one more step (Terry, 2026-10-08: Back after Unzoom did nothing).
    button(view, "◀ Back").click()
    view.run()
    fields = (view.text_input(key="value:start").value, view.text_input(key="value:end").value)
    check("Back after Unzoom: the view before it", fields == (s3, n3) and enabled(view)["Unzoom"], (fields, enabled(view)))
    button(view, "Reset").click()
    view.run()
    # The defaults for the nations ticked: Canada alone stays (batch 17.5).
    check("Reset draws the defaults, Canada's alone", address(view) == "--C" and charts_shown(view) == 1, address(view))
    button(view, "◀ Back").click()
    view.run()
    check("Back undoes Reset", address(view) == e3, (address(view), e3))
    today = new_page()
    check("at today, Later is off", not enabled(today)["Later ▶"] and enabled(today)["◀ Earlier"], enabled(today))


def check_calendars(check: Check) -> None:
    """The calendar beside a date field: it opens at the field's date and writes a picked day into it."""
    cal = new_page("-c -s:2001-03 --e:2026-09-01")
    opened = cal.date_input(key="calendar:start").value
    check("calendar opens at the field's date", opened == datetime.date(2001, 3, 1), opened)
    cal.date_input(key="calendar:start").set_value(datetime.date(2005, 6, 15))
    cal.run()
    check(
        "a picked day is written in the field",
        cal.text_input(key="value:start").value == "2005-06-15" and "--start:2005-06-15" in address(cal),
        address(cal),
    )
    cal.button(key="blank:start").click()
    cal.run()
    check("Blank empties the field", cal.text_input(key="value:start").value == "" and "--start" not in address(cal), address(cal))
    # Before the calendar's first day (config.EARLIEST_DATE, 1600 since batch 28.3; 1680 before).
    cal.text_input(key="value:start").input("1599")
    cal.run()
    check(
        "a date before the calendar's first day leaves it empty, no error",
        cal.date_input(key="calendar:start").value is None and not cal.exception,
        cal.date_input(key="calendar:start").value,
    )
    # After today (2026-10-28, batch 32): an End later than today shows a change already announced for a day in it.
    later = datetime.date.today() + datetime.timedelta(days=3)
    cal = new_page("-c --e:2026-09-01")  # a fresh page: the Start 1599 above, refused, holds back every change
    cal.text_input(key="value:end").input(f"{later:%Y-%m-%d}")
    cal.run()
    check("an End after today: the calendar opens at it, and reaches 2100-12-31",
          cal.date_input(key="calendar:end").value == later and web._CALENDAR_LAST == datetime.date(2100, 12, 31),
          (cal.date_input(key="calendar:end").value, web._CALENDAR_LAST))
    cal.date_input(key="calendar:end").set_value(later + datetime.timedelta(days=1))
    cal.run()
    check("a day after today picked in the calendar: written in the field",
          f"--end:{later + datetime.timedelta(days=1):%Y-%m-%d}" in address(cal) and not cal.exception, address(cal))


def check_size_boxes(check: Check) -> None:
    """Width and Height with the aspect-ratio lock: the WxH text stays the source of truth."""

    def boxes(at: AppTest) -> tuple[str, str]:
        return at.text_input(key="size:width").value, at.text_input(key="size:height").value

    size = new_page("-c -s:2015 --e:2026-09-01 -dim:1600x900")
    check(
        "boxes show the address's size; the lock is on",
        boxes(size) == ("1600", "900") and size.checkbox(key="size:lock").value,
        boxes(size),
    )
    size.text_input(key="size:width").input("2400")
    size.run()
    check(
        "locked: width 2400 gives height 1350, drawn",
        boxes(size) == ("2400", "1350")
        and "--dimensions:2400x1350" in address(size)
        and chart_elements(size)[0].proto.bytes[16:24] == png_size_bytes(2400, 1350),
        (boxes(size), address(size)),
    )
    size.text_input(key="size:height").input("1000")
    size.run()
    check("locked: height 1000 gives width 1778 (the loaded 16:9, no drift)", boxes(size) == ("1778", "1000"), boxes(size))
    size.checkbox(key="size:lock").uncheck()
    size.run()
    size.text_input(key="size:width").input("2000")
    size.run()
    check(
        "unlocked: the height stays",
        boxes(size) == ("2000", "1000") and "--dimensions:2000x1000" in address(size),
        (boxes(size), address(size)),
    )
    size.checkbox(key="size:lock").check()
    size.run()
    size.text_input(key="size:height").input("900")
    size.run()
    check("locked again: the shape of when it was ticked (2:1)", boxes(size) == ("1800", "900"), boxes(size))
    size.checkbox(key="size:lock").uncheck()
    size.run()
    button(size, "Reset").click()
    size.run()
    check(
        "Reset: the default size in the boxes, the lock ticked again (Terry, 2026-10-12)",
        boxes(size) == ("2048", "1536") and "dimensions" not in address(size) and size.checkbox(key="size:lock").value,
        (boxes(size), address(size)),
    )
    button(size, "◀ Back").click()
    size.run()
    check(
        "Back: the boxes follow, and the lock as it was",
        boxes(size) == ("1800", "900")
        and "--dimensions:1800x900" in address(size)
        and not size.checkbox(key="size:lock").value,
        (boxes(size), address(size)),
    )
    one = new_page("-c -s:2015 --e:2026-09-01 -dim:1100")
    check(
        "one dimension: the implied 4:3 shown, the address as given",
        boxes(one) == ("1100", "825") and "--dimensions:1100" in address(one) and "1100x" not in address(one),
        (boxes(one), address(one)),
    )
    one.text_input(key="size:width").input("12a")
    one.run()
    check("a bad box: reported, not drawn", any("Size" in e.value for e in one.error), [e.value for e in one.error])


def check_choice_boxes(check: Check) -> None:
    """The tick boxes of the list options (yield terms, mortgage terms, spreads): the text stays the source of truth."""

    def ticked(at: AppTest, name: str) -> list[str]:
        prefix = f"choice:{name}:"
        return [box.key[len(prefix):] for box in at.checkbox if box.key and box.key.startswith(prefix) and box.value]

    page = new_page("-u --e:2026-09-01 --no-y:30 --mo:15 --sp:10-2,7-1m")
    check(
        "the address's lists ticked; a pair not offered gets its own box",
        ticked(page, "us:yields") == ["3m", "2y", "5y", "10y"]
        and ticked(page, "us:mortgage-terms") == ["15"]
        and ticked(page, "us:spread-pairs") == ["10y-2y", "7y-1m"],
        (ticked(page, "us:yields"), ticked(page, "us:mortgage-terms"), ticked(page, "us:spread-pairs")),
    )
    page.checkbox(key="choice:us:yields:7y").check()
    page.run()
    # Written for the U.S. alone since batch 9.3: Canada has a 7-year now, so
    # one list for both would add it to Canada's chart too.
    check("ticking 7y adds it to the address", "--us:yields:3m,2y,5y,7y,10y" in address(page), address(page))
    page.checkbox(key="choice:us:mortgage-terms:15").uncheck()
    page.run()
    check(
        "the last mortgage box cannot be unticked",
        ticked(page, "us:mortgage-terms") == ["15"] and any("untick “Mortgage rates”" in str(item.value) for item in list(page.info) + list(page.warning)),
        ticked(page, "us:mortgage-terms"),
    )
    page.checkbox(key="flag:spreads").uncheck()
    page.run()
    check(
        "spread boxes greyed while --spreads is off, and left out of the address",
        page.checkbox(key="choice:us:spread-pairs:10y-2y").disabled and "--spreads" not in address(page),
        address(page),
    )


def check_nation_sections(check: Check) -> None:
    """Batch 2: each nation's own section: only its terms, its own limits, and the address they write."""
    page = new_page("--e:2026-09-01 --us:max:9 --cdn:y:2,10 --mo")
    check("no exception", not page.exception, page.exception)
    keys = {box.key for box in page.checkbox if box.key}
    check(
        "Canada offers its terms (all but 20y since batch 9.3) and its mortgage terms only",
        "choice:cdn:yields:30y" in keys and "choice:cdn:yields:7y" in keys and "choice:cdn:yields:20y" not in keys
        and "choice:cdn:mortgage-terms:5" in keys and "choice:cdn:mortgage-terms:30" not in keys,
        sorted(key for key in keys if key.startswith("choice:cdn:")),
    )
    check(
        "the address's values in each nation's fields",
        page.text_input(key="value:us:max").value == "9" and page.text_input(key="value:cdn:max").value == ""
        and page.checkbox(key="choice:cdn:yields:2y").value and not page.checkbox(key="choice:cdn:yields:30y").value
        and page.checkbox(key="choice:us:yields:30y").value,
        (page.text_input(key="value:us:max").value, page.text_input(key="value:cdn:max").value),
    )
    check("address kept", address(page) == "--mortgages --cdn:yields:2y,10y --end:2026-09-01 --us:max:9", address(page))
    page.text_input(key="value:cdn:max").input("9")
    page.run()
    check("the same value for both is written once", "--max:9" in address(page) and "--us:max" not in address(page), address(page))
    page.text_input(key="value:cdn:min").input("x")
    page.run()
    # The label "Min" since batch 18.1, its "%" right of the field.
    check("a bad nation field: reported with its nation", any("Canada, Min:" in e.value for e in page.error), [e.value for e in page.error])


def check_uk(check: Check) -> None:
    """Batch 3: the UK's chart from the address; its section with its own terms and defaults; a second nation ticked."""
    page = new_page("--nations:gb --e:2026-09-01 --uk:max:9")
    check("no exception", not page.exception, page.exception)
    f = flags(page)
    check("address: the UK only", f["uk"] and not f["cdn"] and not f["us"], f)
    check("one chart drawn", charts_shown(page) == 1, charts_shown(page))
    keys = {box.key for box in page.checkbox if box.key}
    check(
        "the UK offers its three terms and its mortgage terms only",
        {"choice:uk:yields:5y", "choice:uk:yields:20y", "choice:uk:mortgage-terms:2f", "choice:uk:mortgage-terms:svr"} <= keys
        and "choice:uk:yields:2y" not in keys and "choice:uk:mortgage-terms:5" not in keys,
        sorted(key for key in keys if key.startswith("choice:uk:")),
    )
    check(
        "its own defaults ticked: all three terms, all five mortgage terms (batch 10.2; the tracker, 9.5)",
        all(page.checkbox(key=f"choice:uk:yields:{term}").value for term in ("5y", "10y", "20y"))
        and all(page.checkbox(key=f"choice:uk:mortgage-terms:{term}").value for term in ("2f", "3f", "5f", "svr", "tracker")),
    )
    check("its value in its own field", page.text_input(key="value:uk:max").value == "9", page.text_input(key="value:uk:max").value)
    check("address kept", address(page) == "--nations:gb --end:2026-09-01 --uk:max:9", address(page))
    page.checkbox(key="flag:cdn").check()
    page.run()
    check("Canada ticked too: two charts", charts_shown(page) == 2, charts_shown(page))
    check("address names both by --nations", address(page) == "--nations:ca,gb --end:2026-09-01 --uk:max:9", address(page))


def check_de(check: Check) -> None:
    """Batch 4: Germany's chart from the address (by --nations:de); its section; its values written with --ger:, never --de:."""
    page = new_page("--na:de --e:2026-09-01 --ger:max:6 --mo")
    check("no exception", not page.exception, page.exception)
    f = flags(page)
    check("address: Germany only", f["germany"] and not f["cdn"] and not f["us"] and not f["uk"], f)
    check("one chart drawn", charts_shown(page) == 1, charts_shown(page))
    keys = {box.key for box in page.checkbox if box.key}
    check(
        "Germany offers its seven terms and its mortgage bands only",
        {f"choice:de:yields:{term}" for term in ("1y", "2y", "5y", "7y", "10y", "20y", "30y")}
        | {"choice:de:mortgage-terms:1-5", "choice:de:mortgage-terms:over10"} <= keys
        and "choice:de:yields:3m" not in keys and "choice:de:mortgage-terms:5" not in keys,
        sorted(key for key in keys if key.startswith("choice:de:")),
    )
    check(
        "its own defaults ticked: 2 5 10 30, all three bands (batch 10.2)",
        all(page.checkbox(key=f"choice:de:yields:{term}").value for term in ("2y", "5y", "10y", "30y"))
        and not page.checkbox(key="choice:de:yields:7y").value
        and all(page.checkbox(key=f"choice:de:mortgage-terms:{term}").value for term in ("1-5", "5-10", "over10")),
    )
    check("its value in its own field", page.text_input(key="value:de:max").value == "6", page.text_input(key="value:de:max").value)
    check("address written with --ger:, not --de: (which is --debt:)",
          address(page) == "--nations:de --mortgages --end:2026-09-01 --ger:max:6", address(page))
    page.checkbox(key="choice:de:mortgage-terms:1-5").uncheck()
    page.run()
    tokens = address(page).split()
    drawn = drawn_config(config_from_choices(*sort_tokens(tokens)[:2]), "de").mortgage_terms
    check("a band unticked: the address draws Germany's other two", drawn == ("5-10", "over10"), (address(page), drawn))
    page.checkbox(key="flag:uk").check()
    page.run()
    check("the UK ticked too: two charts", charts_shown(page) == 2, charts_shown(page))


def check_none_lists(check: Check) -> None:
    """2026-10-07: a list that leaves a nation none of its terms: its boxes all unticked, the address kept, a warning logged."""

    def ticked(at: AppTest, name: str) -> list[str]:
        prefix = f"choice:{name}:"
        return [box.key[len(prefix):] for box in at.checkbox if box.key and box.key.startswith(prefix) and box.value]

    # Canada lacks only the 20-year since batch 9.3 (until then, -c --y:7,20).
    page = new_page("-c --y:20 --e:2026-09-01")
    check("no exception", not page.exception, page.exception)
    check("Canada's yield boxes all unticked (it has no 20y)", ticked(page, "cdn:yields") == [], ticked(page, "cdn:yields"))
    check("no box 'none'", "choice:cdn:yields:none" not in {box.key for box in page.checkbox}, "")
    check("address kept", address(page) == "--C --yields:20y --end:2026-09-01", address(page))
    check("one chart drawn", charts_shown(page) == 1, charts_shown(page))
    warned = [line.strip() for c in page.code for line in c.value.splitlines() if "Warning: no yield" in line]
    check("the warning in the log", any("on the Canadian chart" in line for line in warned), warned)
    page.checkbox(key="choice:cdn:yields:10y").check()
    page.run()
    drawn = drawn_config(config_from_choices(*sort_tokens(address(page).split())[:2]), "cdn").yield_terms
    check("10y ticked: the address draws Canada's 10-year", drawn == ("10y",), (address(page), drawn))
    own = new_page("--cdn:y:none --e:2026-09-01")
    check(
        "a nation's own none: its boxes unticked, the other's at its default, the address kept",
        ticked(own, "cdn:yields") == [] and ticked(own, "us:yields") == ["3m", "2y", "5y", "10y", "30y"]
        and address(own) == "--cdn:yields:none --end:2026-09-01",
        (ticked(own, "cdn:yields"), address(own)),
    )
    bad = new_page("--yields:none --e:2026-09-01")
    warnings = [w.value for w in bad.warning]
    check("a list of none for every chart reported, naming --no-yield", any("--no-yield" in w for w in warnings), warnings)


def check_greyed(check: Check) -> None:
    """Batch 5 (Terry, 2026-10-08): what has no curve to act on is greyed, its tick or text kept, and left out of the address."""

    def disabled(at: AppTest) -> set[str]:
        widgets = [*at.checkbox, *at.text_input]
        keys = {"flag:relative", "flag:per-capita", "flag:regression", "flag:gdp", "level:us:f", "value:regression-tolerance", "value:us:top"}
        return {w.key for w in widgets if w.key in keys and w.disabled}

    page = new_page("-u --yield --debt --reg:2 --d:fp --us:top:40t -s:2001 --e:2026-09-01")
    check("no exception", not page.exception, page.exception)
    check("debt drawn: nothing greyed", disabled(page) == set(), disabled(page))
    first = address(page)
    page.text_input(key="value:regression-tolerance").input("x")
    page.run()
    check("a bad tolerance stops the drawing", any("Not drawn" in e.value for e in page.error), [e.value for e in page.error])
    page.checkbox(key="flag:debt").uncheck()
    page.run()
    every = {"flag:relative", "flag:per-capita", "flag:regression", "level:us:f", "value:regression-tolerance", "value:us:top"}
    check("no right-axis curve: all greyed but GDP", disabled(page) == every, disabled(page))
    check("a bad text in a greyed field stops nothing", not page.error and charts_shown(page) == 1, [e.value for e in page.error])
    check("ticks and texts kept", flags(page)["regression"] and page.checkbox(key="level:us:f").value
          and page.text_input(key="value:us:top").value == "40t" and page.text_input(key="value:regression-tolerance").value == "x")
    check("left out of the address, --regression kept", address(page) == "--U --yield --regression --start:2001 --end:2026-09-01", address(page))
    page.checkbox(key="flag:gdp").check()
    page.run()
    check("GDP: -p, --reg, its tolerance and Top active; -r and the levels greyed",
          disabled(page) == {"flag:relative", "level:us:f"}, disabled(page))
    page.text_input(key="value:regression-tolerance").input("2")
    page.checkbox(key="flag:gdp").uncheck()
    page.checkbox(key="flag:debt").check()
    page.run()
    check("debt again: the chart as at first", address(page) == first, (address(page), first))
    rel = new_page("-u -r --gdp -s:2001 --e:2026-09-01")
    check(
        "-r with GDP only: -r greyed and ticked, GDP active (the tolerance greyed: --reg is off)",
        flags(rel)["relative"] and disabled(rel) == {"flag:relative", "level:us:f", "value:regression-tolerance"},
        (flags(rel)["relative"], disabled(rel)),
    )


def check_wrong_fields(check: Check) -> None:
    """Batch 6: a wrong field's label red; of crossed limits, the box changed last; the "Not drawn" line naming it."""

    def red(at: AppTest, key: str) -> bool:
        return at.text_input(key=key).label.startswith(":red[")

    def errors(at: AppTest) -> list[str]:
        return [e.value for e in at.error]

    page = new_page("--e:2026-09-01 -c --cdn:max:3")
    check("no exception", not page.exception, page.exception)
    page.text_input(key="value:cdn:min").input("5")
    page.run()
    check(
        "Min typed above Max: not drawn, the line names Canada's Min",
        any("Canada, Min: yield min (5.0) must be less than max (3.0)" in e for e in errors(page)), errors(page),
    )
    check("Min's label red, Max's not", red(page, "value:cdn:min") and not red(page, "value:cdn:max"),
          (page.text_input(key="value:cdn:min").label, page.text_input(key="value:cdn:max").label))
    check("Min keeps its text while red", page.text_input(key="value:cdn:min").value == "5", page.text_input(key="value:cdn:min").value)
    page.text_input(key="value:cdn:max").input("5")
    page.run()
    check("Max then typed equal: Max is the one red (changed last)", red(page, "value:cdn:max") and not red(page, "value:cdn:min"), errors(page))
    page.text_input(key="value:cdn:max").input("6")
    page.run()
    check(
        "Max fixed: drawn, no label red",
        not page.error and charts_shown(page) == 1 and not any(box.label.startswith(":red[") for box in page.text_input),
        errors(page),
    )
    page.text_input(key="value:start").input("2026-09-20")
    page.run()
    check("Start typed after End: Start red, End not", red(page, "value:start") and not red(page, "value:end"), errors(page))
    page.text_input(key="value:start").input("")
    # Canada's: the U.S. chart is not shown, and since batch 28.1 nor is its section.
    page.text_input(key="value:cdn:top").input("-5")
    page.run()
    check("a Top at or below zero: red by itself", red(page, "value:cdn:top") and any("Canada, Top" in e for e in errors(page)), errors(page))
    page.text_input(key="value:cdn:top").input("")
    page.checkbox(key="size:lock").uncheck()
    page.run()
    page.text_input(key="size:width").input("600")
    page.run()
    check(
        "the width alone too small (lock off): only Width red",
        red(page, "size:width") and not red(page, "size:height") and any("Size" in e for e in errors(page)),
        (page.text_input(key="size:width").label, page.text_input(key="size:height").label, errors(page)),
    )


def check_layout(check: Check) -> None:
    """Batch 7 (Terry, 2026-10-08 and 2026-10-12): the sidebar's blocks in the window's order; Max above Min."""
    at = new_page("--e:2026-09-01")
    check("no exception", not at.exception, at.exception)
    titles = [header.value for header in at.subheader]
    wanted = [
        "Country selection", "Curve selection", "Measure of debt, GDP, and interest", "Rates on the yield axis",
        "Date window", "Final canvas size in pixels",
    ]
    check("the titled blocks in order", titles[: len(wanted)] == wanted, titles)
    order = [box.key.split(":", 1)[1] for box in at.checkbox if box.key and box.key.startswith("flag:")]
    wanted = [
        "relative", "per-capita", "regression", "yield", "spreads", "policy", "mortgages", "label", "current", "auto-refresh",
        "projections", "hover",
    ]
    keys = [box.key for box in at.checkbox if box.key]
    # By level left the column for each nation's section in batch 10.5.
    check(
        "the measure under the curves, then the regression segments, the four rates, the line labels and the latest data",
        keys.index("flag:interest") < keys.index("flag:relative") and [name for name in order if name in wanted] == wanted
        and not any(key.startswith("level:") and key.count(":") == 1 for key in keys),
        order,
    )
    # The limits only: the prime margin's field sits with the mortgage boxes, above them (batch 9.4).
    fields = [
        field.key for field in at.text_input
        if field.key and field.key.startswith("value:cdn:") and not field.key.endswith(":prime-margin")
    ]
    check(
        "Canada's limits row by row: Max beside Top, then Min beside Bottom",
        fields == ["value:cdn:max", "value:cdn:top", "value:cdn:min", "value:cdn:bottom"],
        fields,
    )
    # 2026-10-13: at the sidebar's narrow default width each list box's own
    # "?" left no room for its text (Canada's terms showed only the "?").
    helped = [box.key for box in at.checkbox if box.key and box.key.startswith("choice:") and box.help]
    check("a list's tick boxes carry no help of their own (it is on the list's label)", not helped, helped[:3])


def check_batch15(check: Check) -> None:
    """Batch 15 (Terry, 2026-10-13 and 2026-10-14): the two ticks' wording; Save log file."""
    at = new_page("--e:2026-09-01")
    check("no exception", not at.exception, at.exception)
    labels = {box.key: box.label for box in at.checkbox if box.key in ("flag:label", "flag:current")}
    check(
        "-l reads 'Show latest value for each line', --cur 'Intraday quotes on end date'",
        labels == {"flag:label": "Show latest value for each line", "flag:current": "Intraday quotes on end date"},
        labels,
    )
    name, kind, data = downloaded(at, "Save log file")
    log = next((c.value for c in at.code if c.value.startswith("--- drawing")), None)
    check(
        "Save log file: ratesplot_log_DATE_HHMM.txt, plain text",
        name.startswith("ratesplot_log_") and name.endswith(".txt") and len(name) == len("ratesplot_log_2026-10-14_1015.txt")
        and kind == "text/plain",
        (name, kind),
    )
    check("Save log file: the log as shown", log is not None and data.decode("utf-8") == log + "\n", (len(data), log and len(log)))
    # 15.5: a number typed without its 0 gets it, and the address follows.
    at = new_page("--reg --e:2026-09-01")
    at.text_input(key="value:regression-tolerance").input(".3")
    at.run()
    at.text_input(key="value:cdn:min").input("-.5")
    at.run()
    shown = (at.text_input(key="value:regression-tolerance").value, at.text_input(key="value:cdn:min").value)
    check("'.3' and '-.5' typed: '0.3' and '-0.5' in their boxes", shown == ("0.3", "-0.5"), shown)
    check("... and in the address", "0.3" in address(at) and "-0.5" in address(at) and ":.3" not in address(at), address(at))


def check_batch9(check: Check) -> None:
    """Batch 9.3: Canada's section offers its ten terms (all but the 20-year), the U.S.' all but the 3-year."""
    at = new_page("--e:2026-09-01")
    check("no exception", not at.exception, at.exception)
    keys = {box.key for box in at.checkbox if box.key}
    canada = {key for key in keys if key.startswith("choice:cdn:yields:")}
    united = {key for key in keys if key.startswith("choice:us:yields:")}
    wanted = {f"choice:cdn:yields:{term}" for term in ("1m", "3m", "6m", "1y", "2y", "3y", "5y", "7y", "10y", "30y")}
    check("Canada offers 1m 3m 6m 1y 2y 3y 5y 7y 10y 30y, not 20y", canada == wanted, sorted(canada))
    check("the U.S. offers all but 3y", "choice:us:yields:3y" not in united and len(united) == 10, sorted(united))
    ticked = {key.rsplit(":", 1)[1] for key in canada if at.checkbox(key=key).value}
    check("Canada's default five ticked, the new terms not", ticked == {"3m", "2y", "5y", "10y", "30y"}, sorted(ticked))
    at.checkbox(key="choice:cdn:yields:3y").check()
    at.run()
    tokens = address(at).split()
    drawn = drawn_config(config_from_choices(*sort_tokens(tokens)[:2]), "cdn").yield_terms
    check("3y ticked: the address draws Canada's 3-year with the five", drawn == ("3m", "2y", "3y", "5y", "10y", "30y"), (address(at), drawn))


def check_batch94(check: Check) -> None:
    """Batch 9.4: a prime margin field beside Canada's and the U.S.' prime box, greyed while that nation draws no prime rate."""
    at = new_page("--e:2026-09-01")
    check("no exception", not at.exception, at.exception)
    fields = {field.key for field in at.text_input if field.key and field.key.endswith(":prime-margin")}
    check("a margin field for Canada and the U.S. only", fields == {"value:cdn:prime-margin", "value:us:prime-margin"}, sorted(fields))
    boxes = {box.key for box in at.checkbox if box.key and box.key.endswith(":prime")}
    check("a prime box for each of them", boxes == {"choice:cdn:mortgage-terms:prime", "choice:us:mortgage-terms:prime"}, sorted(boxes))

    def greyed(page: AppTest) -> set[str]:
        return {key.split(":")[1] for key in fields if page.text_input(key=key).disabled}

    check("mortgages off: both greyed", greyed(at) == {"cdn", "us"}, greyed(at))
    at.checkbox(key="flag:mortgages").check()
    at.run()
    check("mortgages on, prime unticked: both greyed", greyed(at) == {"cdn", "us"}, greyed(at))
    at.checkbox(key="choice:cdn:mortgage-terms:prime").check()
    at.run()
    check("Canada's prime ticked: its field alone active", greyed(at) == {"us"}, greyed(at))
    at.text_input(key="value:cdn:prime-margin").input("-.25")
    at.run()
    drawn = drawn_config(config_from_choices(*sort_tokens(address(at).split())[:2]), "cdn")
    check("the address draws Canada's prime less 0.25",
          drawn.mortgage_terms == ("5", "prime") and drawn.prime_margin == -0.25, (address(at), drawn.mortgage_terms, drawn.prime_margin))
    check("no exception, one drawing per chart", not at.exception and charts_shown(at) == 2, at.exception)
    at.checkbox(key="choice:cdn:mortgage-terms:prime").uncheck()
    at.run()
    check("prime unticked: the margin greyed, kept, out of the address",
          greyed(at) == {"cdn", "us"} and at.text_input(key="value:cdn:prime-margin").value == "-0.25" and "prime" not in address(at),
          (greyed(at), address(at)))
    page = new_page("--prime:0.5 --e:2026-09-01")
    ticked = {key for key in boxes if page.checkbox(key=key).value}
    texts = {key: page.text_input(key=key).value for key in sorted(fields)}
    check("--prime:0.5: both prime boxes ticked, both fields 0.5",
          ticked == boxes and set(texts.values()) == {"0.5"} and flags(page)["mortgages"], (sorted(ticked), texts))


def check_batch95(check: Check) -> None:
    """Batch 9.5: the UK tracker's box, ticked by default, with its margin field beside it, greyed while the UK draws no tracker."""
    at = new_page("--nations:gb --e:2026-09-01")
    check("no exception", not at.exception, at.exception)
    fields = {field.key for field in at.text_input if field.key and field.key.endswith(":tracker-margin")}
    check("a tracker margin field for the UK only", fields == {"value:uk:tracker-margin"}, sorted(fields))
    box, field = "choice:uk:mortgage-terms:tracker", "value:uk:tracker-margin"
    check("the tracker ticked by default (the UK's own list)", at.checkbox(key=box).value)
    check("its grey value is the default margin, 1.75", at.text_input(key=field).placeholder == "1.75", at.text_input(key=field).placeholder)
    check("mortgages off: greyed", at.text_input(key=field).disabled)
    at.checkbox(key="flag:mortgages").check()
    at.run()
    check("mortgages on, the tracker ticked: active", not at.text_input(key=field).disabled)
    at.text_input(key=field).input("2")
    at.run()
    drawn = drawn_config(config_from_choices(*sort_tokens(address(at).split())[:2]), "uk")
    check("the address draws the UK's tracker at Bank Rate + 2",
          "tracker" in drawn.mortgage_terms and drawn.tracker_margin == 2.0, (address(at), drawn.mortgage_terms, drawn.tracker_margin))
    check("no exception, one drawing", not at.exception and charts_shown(at) == 1, at.exception)
    at.checkbox(key=box).uncheck()
    at.run()
    check("the tracker unticked: the margin greyed, kept, out of the address",
          at.text_input(key=field).disabled and at.text_input(key=field).value == "2" and "tracker" not in address(at), address(at))


def check_batch91(check: Check) -> None:
    """Batch 9.1: each nation's policy terms, target and market, as boxes; the target ticked; greyed while the policy rate is off."""
    # Every nation ticked, so every nation's section is laid out (batch 28.1).
    at = new_page("--na:ca,us,gb,de --e:2026-09-01")
    check("no exception", not at.exception, at.exception)
    boxes = {box.key for box in at.checkbox if box.key and ":policy-terms:" in box.key}
    wanted = {f"choice:{key}:policy-terms:{term}" for key in NATION_KEYS for term in ("target", "market")}
    check("every nation's section has target and market", boxes == wanted, sorted(boxes))
    labels = {key: at.checkbox(key=f"choice:{key}:policy-terms:market").label for key in NATION_KEYS}
    check("each nation's market box named for its rate (batch 24.5)",
          labels == {"cdn": "market (CORRA)", "us": "market (fed funds)", "uk": "market (SONIA)", "de": "market (€STR)"}, labels)
    targets = {f"choice:{key}:policy-terms:target" for key in NATION_KEYS}
    check("the target ticked in each, the market not", {key for key in boxes if at.checkbox(key=key).value} == targets)
    check("greyed while the policy rate is off", all(at.checkbox(key=key).disabled for key in boxes))
    at.checkbox(key="flag:policy").check()
    at.run()

    def drawn(page: AppTest, key: str):
        return drawn_config(config_from_choices(*sort_tokens(address(page).split())[:2]), key)

    check("policy on: the boxes active, the address draws the target",
          not any(at.checkbox(key=key).disabled for key in boxes) and drawn(at, "us").policy_rates
          and drawn(at, "us").policy_terms == ("target",), address(at))
    at.checkbox(key="choice:us:policy-terms:market").check()
    at.run()
    terms = {key: drawn(at, key).policy_terms for key in ("cdn", "us")}
    check("the U.S.' market ticked: the address draws both there, the target alone in Canada",
          terms == {"cdn": ("target",), "us": ("target", "market")}, (address(at), terms))
    at.checkbox(key="choice:cdn:policy-terms:target").uncheck()
    at.run()
    check("Canada's last policy box cannot be unticked",
          at.checkbox(key="choice:cdn:policy-terms:target").value
          and any("untick “Policy rate”" in str(item.value) for item in list(at.info) + list(at.warning)), address(at))
    check("no exception, one drawing per chart", not at.exception and charts_shown(at) == 4, at.exception)
    page = new_page("--po:market --na:ca,us,gb,de --e:2026-09-01")
    markets = {f"choice:{key}:policy-terms:market" for key in NATION_KEYS}
    check("--po:market: the market ticked in every section, the target not",
          {key for key in wanted if page.checkbox(key=key).value} == markets and flags(page)["policy"],
          sorted(key for key in wanted if page.checkbox(key=key).value))


def check_batch104(check: Check) -> None:
    """Batch 10.4: the UK's gross and net debt as two boxes, gross ticked; greyed without debt or with a level; no other nation's."""
    at = new_page("--nations:ca,gb --e:2026-09-01")
    check("no exception", not at.exception, at.exception)
    gross, net = "choice:uk:aggregate:gross", "choice:uk:aggregate:net"
    boxes = {box.key for box in at.checkbox if box.key and ":aggregate:" in box.key}
    check("the UK's section alone has them, two boxes", boxes == {gross, net}, sorted(boxes))
    check("gross ticked, net not (decision 33)", at.checkbox(key=gross).value and not at.checkbox(key=net).value)
    check("active while debt is drawn", not at.checkbox(key=gross).disabled and not at.checkbox(key=net).disabled)

    def drawn(page: AppTest) -> tuple[str, ...]:
        return drawn_config(config_from_choices(*sort_tokens(address(page).split())[:2]), "uk").debt_aggregates

    at.checkbox(key=net).check()
    at.run()
    check("net ticked: the address draws both on the UK's chart",
          drawn(at) == ("gross", "net") and "--aggregate:gross,net" in address(at), (address(at), drawn(at)))
    at.checkbox(key=gross).uncheck()
    at.run()
    check("gross unticked: net alone", drawn(at) == ("net",), (address(at), drawn(at)))
    at.checkbox(key=net).uncheck()
    at.run()
    check("the last box cannot be unticked, and the debt curve's box is named",
          at.checkbox(key=net).value
          and any("untick “Aggregate public debt”" in str(item.value) for item in list(at.info) + list(at.warning)), address(at))
    at.checkbox(key="flag:debt").uncheck()
    at.run()
    check("debt off: both greyed, ticks kept, the list out of the address",
          at.checkbox(key=gross).disabled and at.checkbox(key=net).disabled and at.checkbox(key=net).value
          and "--agg" not in address(at), address(at))
    at.checkbox(key="flag:debt").check()
    at.checkbox(key="level:uk:f").check()
    at.run()
    check("a level ticked: greyed (the levels replace the aggregate lines), the list out of the address",
          at.checkbox(key=net).disabled and "--agg" not in address(at), address(at))
    at.checkbox(key="level:uk:f").uncheck()
    at.run()
    check("the level unticked: active again, net alone again",
          not at.checkbox(key=net).disabled and drawn(at) == ("net",), (address(at), drawn(at)))
    check("no exception, two charts", not at.exception and charts_shown(at) == 2, (at.exception, charts_shown(at)))
    page = new_page("--nations:gb --agg:gross,net --e:2026-09-01")
    check("--agg:gross,net from the address: both ticked",
          page.checkbox(key=gross).value and page.checkbox(key=net).value, address(page))


def check_batch105(check: Check) -> None:
    """Batch 10.5: the levels in each nation's section, named for its governments; a nation's own written after its code."""
    at = new_page("--na:ca,us,gb,de --e:2026-09-01 -s:2005")
    check("no exception", not at.exception, at.exception)
    wanted = {
        "cdn": ["Federal", "Provincial", "Municipal", "Non-federal"],
        "us": ["Federal", "State", "Local", "State & Local"],
        "uk": ["Central", "Local"],
        "de": ["Federal", "Länder", "Local", "Länder & Local"],
    }
    captions = {
        key: [box.label for box in at.checkbox if box.key and box.key.startswith(f"level:{key}:")] for key in NATION_KEYS
    }
    check("each nation's levels named for its governments", captions == wanted, captions)
    check("none in the column", not any(box.key and box.key.count(":") == 1 and box.key.startswith("level:") for box in at.checkbox))

    def drawn(page: AppTest, key: str) -> str:
        return drawn_config(config_from_choices(*sort_tokens(address(page).split())[:2]), key).components

    at.checkbox(key="level:cdn:f").check()
    at.checkbox(key="level:cdn:p").check()
    at.run()
    check("Canada's federal and provincial: written after its code, every curve written",
          "--cdn:debt:fp" in address(at) and drawn(at, "cdn") == "fp" and drawn(at, "us") == "", address(at))
    for key in ("us", "uk", "de"):
        at.checkbox(key=f"level:{key}:f").check()
        if key != "uk":
            at.checkbox(key=f"level:{key}:p").check()
    at.run()
    check("the same levels on every chart (the UK has no provincial level): written once",
          "--debt:fp" in address(at) and "--cdn:debt" not in address(at) and "--uk:debt" not in address(at)
          and drawn(at, "uk") == "f", address(at))
    check("the UK's aggregate debts greyed by its level",
          at.checkbox(key="choice:uk:aggregate:gross").disabled and at.checkbox(key="choice:uk:aggregate:net").disabled)
    at.checkbox(key="flag:debt").uncheck()
    at.checkbox(key="flag:interest").uncheck()
    at.run()
    check("no debt or interest: every nation's levels greyed, ticks kept, out of the address",
          all(box.disabled for box in at.checkbox if box.key and box.key.startswith("level:"))
          and at.checkbox(key="level:cdn:f").value and ":debt:" not in address(at) and "--debt:" not in address(at), address(at))
    check("no exception, four charts", not at.exception and charts_shown(at) == 4, (at.exception, charts_shown(at)))
    page = new_page("--na:gb,de --debt:fn --ger:interest:pm --e:2026-09-01 -s:2005")
    ticked = {key: [box.key.rsplit(":", 1)[1] for box in page.checkbox if box.key and box.key.startswith(f"level:{key}:") and box.value]
              for key in ("uk", "de")}
    check("from the address: the UK's n ticks its Local box; Germany's own levels", ticked == {"uk": ["f", "m"], "de": ["p", "m"]}, ticked)
    check("its charts the address's", drawn(page, "uk") == "fm" and drawn(page, "de") == "pm", address(page))


def check_batch17(check: Check) -> None:
    """Batch 17.4 and 17.5: a nation ticked by hand gets its tab once drawn; Reset leaves the nation boxes alone."""
    at = new_page("--e:2026-09-01")
    check("no exception", not at.exception, at.exception)
    check("the first tab current at first", at.session_state["tab"] == "Canada", at.session_state["tab"])
    at.checkbox(key="flag:uk").check()
    at.run()
    check("UK ticked: three charts, its tab current", charts_shown(at) == 3 and at.session_state["tab"] == "United Kingdom",
          (charts_shown(at), at.session_state["tab"]))
    at.checkbox(key="flag:germany").check()
    at.run()
    check("Germany ticked: its tab current", at.session_state["tab"] == "Germany", at.session_state["tab"])
    at.checkbox(key="flag:germany").uncheck()
    at.run()
    check("the current nation unticked: the first tab current", at.session_state["tab"] == "Canada", at.session_state["tab"])
    at.session_state["tab"] = "United States"  # as a click on its tab leaves it
    at.run()
    at.checkbox(key="flag:uk").uncheck()
    at.run()
    check("another nation unticked: the current tab kept", at.session_state["tab"] == "United States", at.session_state["tab"])
    at.checkbox(key="flag:germany").check()
    at.text_input(key="value:cdn:max").input("9")
    at.run()
    at.checkbox(key="flag:us").uncheck()
    at.run()
    button(at, "Reset").click()
    at.run()
    f = flags(at)
    check("Reset: the nation boxes as they were (Canada, Germany)",
          (f["cdn"], f["us"], f["uk"], f["germany"]) == (True, False, False, True), {key: f[key] for key in ("cdn", "us", "uk", "germany")})
    check("Reset: each nation's own fields reset", at.text_input(key="value:cdn:max").value == "", at.text_input(key="value:cdn:max").value)
    check("no exception, two charts", not at.exception and charts_shown(at) == 2, (at.exception, charts_shown(at)))


def check_batch18(check: Check) -> None:
    """Batch 18.1 to 18.3: each limit's unit on its right, and the "All" tick that copies it to every nation, an amount with its currency."""
    at = new_page("--na:ca,us --e:2026-09-01 -s:2015")
    check("no exception", not at.exception, at.exception)
    labels = {name: at.text_input(key=f"value:us:{name}").label for name in ("max", "min", "top", "bottom")}
    check("the labels without %", labels == {"max": "Max", "min": "Min", "top": "Top", "bottom": "Bottom"}, labels)

    def units(page: AppTest) -> list[str]:
        """The unit symbols right of the limit fields, in order (each in its HTML block, ``web._limit_control``)."""
        found = (re.fullmatch(r'<div style="padding-bottom: [^"]*">(.*)</div>', m.value) for m in page.markdown)
        return [match.group(1) for match in found if match]

    # Only the ticked nations' sections are laid out (batch 28.1): Canada's and the U.S.'.
    shown = units(at)
    check("the units: % beside Max and Min, each nation's currency beside Top and Bottom",
          shown == ["%", "C$", "%", "C$", "%", "$", "%", "$"], shown)
    alls = [box.key for box in at.checkbox if box.key and box.key.startswith("all:")]
    check("an All box beside each of the 8 limits shown", len(alls) == 8, alls)
    check("every All box greyed while its field is blank", all(at.checkbox(key=key).disabled for key in alls))

    def values(name: str) -> list[str]:
        # The UK's and Germany's fields are not laid out, but keep their values (``web._keep_section``).
        return [at.session_state[f"value:{key}:{name}"] for key in NATION_KEYS]

    def ticked(name: str) -> list[bool]:
        return [bool(at.session_state[f"all:{key}:{name}"]) for key in NATION_KEYS]

    at.text_input(key="value:us:max").input("8")
    at.run()
    check("a value entered: its All active, the others greyed",
          not at.checkbox(key="all:us:max").disabled and at.checkbox(key="all:cdn:max").disabled)
    at.checkbox(key="all:us:max").check()
    at.run()
    check("All ticked: every nation's Max 8, every Max's All ticked",
          values("max") == ["8"] * 4 and all(ticked("max")), values("max"))
    check("All ticked: written once, --max:8", "--max:8" in address(at) and "--us:max" not in address(at), address(at))
    at.text_input(key="value:cdn:max").input("7")
    at.run()
    check("a value entered in another nation's field goes into every one", values("max") == ["7"] * 4, values("max"))
    at.text_input(key="value:cdn:max").input("")
    at.run()
    check("a field blanked: its All unticked in every section, the others' values kept",
          not any(ticked("max")) and values("max") == ["", "7", "7", "7"],
          values("max"))
    check("a field blanked: its own All greyed, the others' active",
          at.checkbox(key="all:cdn:max").disabled and not at.checkbox(key="all:us:max").disabled)
    at.text_input(key="value:us:top").input("40t")
    at.run()
    # Batch 18.3: an amount goes with its currency, which each chart converts.
    check("Top in dollars: its All active", not at.checkbox(key="all:us:top").disabled)
    at.checkbox(key="all:us:top").check()
    at.run()
    check("All on Top: the U.S.' 40t as USD40t in the other nations' fields", values("top") == ["USD40t", "40t", "USD40t", "USD40t"],
          values("top"))
    check("All on Top: written once, --top:USD40t", "--top:USD40t" in address(at) and "--us:top" not in address(at), address(at))
    check("no exception", not at.exception, at.exception)
    at.checkbox(key="flag:relative").check()
    at.run()
    at.text_input(key="value:us:top").input("150%")
    at.run()
    check("-r: Top's and Bottom's unit blank, Max's and Min's %", units(at) == ["%", "&nbsp;"] * 4, units(at))
    at.checkbox(key="all:us:top").check()
    at.run()
    check("-r: All on Top copies 150% to every nation", values("top") == ["150%"] * 4 and "--top:150%" in address(at), values("top"))
    at.checkbox(key="flag:relative").uncheck()
    at.run()
    check("the measure changed: Top cleared, its All unticked", values("top") == [""] * 4 and not at.checkbox(key="all:us:top").value,
          values("top"))
    at.checkbox(key="flag:cdn").uncheck()
    at.run()
    # Canada's section gone with its box (batch 28.1): the U.S.' All boxes alone are left.
    check("one nation: every All greyed",
          all(at.checkbox(key=key).disabled for key in alls if key.startswith("all:us:"))
          and not any(box.key.startswith("all:cdn:") for box in at.checkbox if box.key))
    check("no exception", not at.exception, at.exception)
    page = new_page("--na:ca,us --max:8 --e:2026-09-01 -s:2015")
    check("from an address with --max:8: Max's All ticked", page.checkbox(key="all:cdn:max").value, address(page))
    page = new_page("-u --max:8 --e:2026-09-01 -s:2015")
    check("one nation: not ticked", not page.checkbox(key="all:us:max").value, address(page))


def check_batch11(check: Check) -> None:
    """Batch 11.1: under each chart a menu of formats and Export chart, the file made on the click, the format shared by the tabs."""
    from openpyxl import load_workbook

    at = new_page("-s:2015 --e:2026-09-01")
    check("no exception", not at.exception, at.exception)
    captions = [caption for caption, _extension, _every in web.EXPORT_CHOICES]
    menus = {box.key: box for box in at.selectbox if box.key and box.key.startswith("export:")}
    check("a format menu under each chart, the eight formats in order, PNG first",
          sorted(menus) == ["export:cdn", "export:us"] and all(list(box.options) == captions and box.value == "PNG image"
                                                                for box in menus.values()),
          {key: (box.value, len(box.options)) for key, box in menus.items()})
    buttons = sorted(b.proto.id.rsplit("-", 1)[1] for b in at.get("download_button") if b.proto.label == web.EXPORT_LABEL)
    check("an Export chart button under each chart (Download until batch 31.4), its file made on the click",
          buttons == ["download:cdn", "download:us"] and all(b.proto.deferred_file_id for b in at.get("download_button")
                                                             if b.proto.label == web.EXPORT_LABEL), buttons)
    drawing = at.session_state["drawing"]
    finished, chart = drawing.finished, address(at)
    name, kind, data = clicked_download(at, "download:us")
    check("PNG: us_rates_START_END.png, the bytes drawn", (name, kind) == ("us_rates_2015-01-01_2026-09-01.png", "image/png")
          and data == drawing.pngs["us"], (name, kind))
    us_data = web.export.chart_data(drawing.figures["us"], drawing.drawn.for_nation("us"), "us")
    starts = {".jpg": b"\xff\xd8\xff", ".pdf": b"%PDF", ".svg": b"<?xml"}
    mimes = {".png": "image/png", ".jpg": "image/jpeg", ".pdf": "application/pdf", ".svg": "image/svg+xml", ".csv": "text/csv",
             ".tsv": "text/tab-separated-values", ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
    for caption, extension, every in web.EXPORT_CHOICES[1:]:
        at.selectbox(key="export:us").select(caption).run()
        name, kind, data = clicked_download(at, "download:us")
        wanted = f"{'' if every else 'us_'}rates_2015-01-01_2026-09-01{extension}"
        if extension in starts:
            right = data.startswith(starts[extension])
        elif extension == ".xlsx":
            book = load_workbook(io.BytesIO(data))
            right = book.sheetnames == (["Canada", "United States"] if every else ["United States"])
        else:
            right = data == web.export.delimited_bytes(us_data, "," if extension == ".csv" else "\t")
        check(f"{caption}: {wanted}, {mimes[extension]}", (name, kind) == (wanted, mimes[extension]) and right, (name, kind, data[:8]))
    check("the other tab's menu shows the format last chosen", at.selectbox(key="export:cdn").value == captions[-1],
          at.selectbox(key="export:cdn").value)
    name, _kind, data = clicked_download(at, "download:cdn")
    check("... and its Export chart gives its own chart's", name == "cdn_rates_2015-01-01_2026-09-01.tsv" and data
          == web.export.delimited_bytes(web.export.chart_data(drawing.figures["cdn"], drawing.drawn.for_nation("cdn"), "cdn"), "\t"),
          name)
    check("choosing formats drew nothing, the address unchanged",
          at.session_state["drawing"].finished == finished and address(at) == chart, (address(at), chart))


def chart_bytes(at: AppTest, key: str) -> bytes:
    """What a nation's chart component got in the page's last run (its key holds "chart:KEY:")."""
    return next(c for c in at.get("bidi_component") if f"chart:{key}:" in c.proto.id).proto.bytes


def check_batch112(check: Check) -> None:
    """Batch 11.2: the hover read-out's box; its data after each chart's PNG while it is on; unticked, nothing drawn again."""
    at = new_page("-s:2015 --e:2026-09-01")
    check("no exception", not at.exception, at.exception)
    box = at.checkbox(key="flag:hover")
    check("a box Hover read-out, ticked", box.label == "Hover read-out" and box.value, (box.label, box.value))
    drawing = at.session_state["drawing"]
    finished, chart = drawing.finished, address(at)
    sizes = {}
    for key in ("cdn", "us"):
        data = chart_bytes(at, key)
        png = web.png_part(data)
        sizes[key] = (len(png), len(data) - len(png))
        check(f"{key}: the PNG drawn, then the read-out's data", png == drawing.pngs[key] and len(data) > len(png), sizes[key])
        # The data as the page's script unpacks it, against the lines as the window reads them.
        sent = json.loads(zlib.decompress(data[len(png):]))
        lines = hover.chart_lines(drawing.figures[key])
        origin, per_day = sent["days_to_x"]
        same = sent["box"] == [round(value, 2) for value in lines.box] and [c["name"] for c in sent["curves"]] == [c.name for c in lines.curves]
        for sent_curve, curve in zip(sent["curves"], lines.curves):
            same &= (sent_curve["texts"], sent_curve["steps"], sent_curve["projected"], sent_curve["quoted"]) == (
                list(curve.texts), curve.steps, curve.projected, curve.quoted)
            for sent_run, run in zip(sent_curve["runs"], curve.runs):
                days, tenths = np.cumsum(sent_run["d"]), np.cumsum(sent_run["y"])
                same &= (days.tolist() == run.day.tolist() and sent_run["t"] == run.text.tolist()
                         and bool(np.all(np.abs(tenths / 10 - run.y) <= 0.05)) and bool(np.all(np.abs(origin + days * per_day - run.x) < 0.01))
                         and sent_run["until"] == (run.until[1] if run.until is not None else None))
        check(f"{key}: the data unpacked are the lines the window reads (days, x from the day, y, texts)", same)
    check("the data a tenth of the PNG or less", all(extra * 10 <= png for png, extra in sizes.values()), sizes)
    at.checkbox(key="flag:hover").uncheck().run()
    check("unticked: nothing drawn again, each chart its PNG alone", at.session_state["drawing"].finished == finished
          and all(chart_bytes(at, key) == drawing.pngs[key] for key in ("cdn", "us")))
    check("unticked: --no-hover in the address and the command line shown", "--no-hover" in address(at)
          and any("--no-hover" in code.value for code in at.code), address(at))
    at.checkbox(key="flag:hover").check().run()
    check("ticked again: the data again, nothing drawn, the address as before", at.session_state["drawing"].finished == finished
          and len(chart_bytes(at, "us")) > len(drawing.pngs["us"]) and address(at) == chart, address(at))
    page = new_page("--no-hover -s:2015 --e:2026-09-01")
    check("from an address with --no-hover: unticked, each chart its PNG alone", not page.checkbox(key="flag:hover").value
          and all(chart_bytes(page, key) == page.session_state["drawing"].pngs[key] for key in ("cdn", "us")))


class _Rerun(Exception):
    """``st.rerun`` stood in for: the fragment's run ends there, as Streamlit ends it."""


class _StandIn:
    """What ``web._auto_refresh_tick`` uses of Streamlit, stood in for: AppTest runs no fragment's timer."""

    def __init__(self, state: dict) -> None:
        self.session_state = state
        self.captions: list[str] = []
        self.reruns: list[str] = []

    def caption(self, text: str) -> None:
        self.captions.append(text)

    def rerun(self, scope: str = "app") -> None:
        self.reruns.append(scope)
        raise _Rerun


def check_batch113(check: Check) -> None:
    """Batch 11.3 (decision 23): the "Automatic refresh" box and its minutes; greyed where it brings nothing new; the fragment's ticks."""
    # A chart ending today with the quotes on, so a refresh can bring something new: the U.S. yields alone.
    at = new_page("-u --yield -s:2026-06-01")
    check("no exception", not at.exception, at.exception)
    box, minutes = at.checkbox(key="flag:auto-refresh"), at.text_input(key="value:auto-refresh-minutes")
    check("a box Automatic refresh, unticked, not greyed", box.label == "Automatic refresh" and not box.value and not box.disabled,
          (box.label, box.value, box.disabled))
    check("its minutes, greyed while unticked, 5 in grey", minutes.label == "interval, minutes" and minutes.disabled
          and minutes.placeholder == "5", (minutes.label, minutes.disabled, minutes.placeholder))

    def note() -> list[str]:
        return [caption.value for caption in at.caption if caption.value.startswith("Refreshed automatically")]

    check("unticked: no refresh under the charts", not note())
    finished = at.session_state["drawing"].finished
    at.checkbox(key="flag:auto-refresh").check().run()
    check("ticked: nothing drawn again", at.session_state["drawing"].finished == finished)
    check("ticked: --auto-refresh in the address", "--auto-refresh" in address(at), address(at))
    check("ticked: every 5 minutes, said under the charts", note() and "every 5 minutes" in note()[0], note())
    check("ticked: its minutes active", not at.text_input(key="value:auto-refresh-minutes").disabled)
    at.text_input(key="value:auto-refresh-minutes").input("10").run()
    check("10: --au:10 in the address, every 10 minutes, nothing drawn", "--au:10" in address(at)
          and note() and "every 10 minutes" in note()[0] and at.session_state["drawing"].finished == finished, address(at))
    at.text_input(key="value:auto-refresh-minutes").input("4").run()
    check("4: refused, its label red", any("Not drawn" in error.value for error in at.error)
          and at.text_input(key="value:auto-refresh-minutes").label.startswith(":red["), [error.value for error in at.error])
    at.text_input(key="value:auto-refresh-minutes").input("10").run()

    at.checkbox(key="flag:current").uncheck().run()
    box, minutes = at.checkbox(key="flag:auto-refresh"), at.text_input(key="value:auto-refresh-minutes")
    check("the quotes off: both greyed, the tick and text kept", box.disabled and minutes.disabled and box.value and minutes.value == "10")
    check("greyed: --auto-refresh in the address, its minutes not, no refresh", "--auto-refresh" in address(at)
          and "--au:10" not in address(at) and not note(), address(at))
    at.checkbox(key="flag:current").check().run()
    check("the quotes on again: active", not at.checkbox(key="flag:auto-refresh").disabled and note())
    page = new_page("-u --yield -s:2026-01-01 --e:2026-09-01 --au:10")
    check("from an address with an End before today: ticked, greyed", page.checkbox(key="flag:auto-refresh").value
          and page.checkbox(key="flag:auto-refresh").disabled and page.text_input(key="value:auto-refresh-minutes").value == "10")

    # The fragment's own runs, with Streamlit stood in for: AppTest runs no timer.
    original, original_ticker = web.st, web.Ticker
    state = {"auto_laid_out": True}
    stand_in = web.st = _StandIn(state)
    try:
        today = cli.parse_args(["--no-gui", "-u", "--au:10"])

        class Ticker:
            answer = True

            def due(self, _config, now=None):
                return Ticker.answer

        state["auto_ticker"] = Ticker()
        web._auto_refresh_tick(today)
        check("laid out with the page: says how often, nothing else", stand_in.captions and "every 10 minutes" in stand_in.captions[0]
              and not stand_in.reruns and "force_redraw" not in state, stand_in.captions)
        state["auto_laid_out"] = False
        try:
            web._auto_refresh_tick(today)
        except _Rerun:
            pass
        check("its own run, due: a drawing forced, the whole page run again", state.get("force_redraw") is True
              and stand_in.reruns == ["app"], (state.get("force_redraw"), stand_in.reruns))
        state.pop("force_redraw")
        Ticker.answer = False
        web._auto_refresh_tick(today)
        check("its own run, not due: nothing drawn, said", stand_in.reruns == ["app"] and "force_redraw" not in state
              and "no quote has changed" in stand_in.captions[-1], stand_in.captions[-1])
        try:
            web._auto_refresh_tick(replace(today, end=pd.Timestamp("2020-01-02")))
        except _Rerun:
            pass
        check("the day changed (a blank End): the whole page run again, nothing forced", stand_in.reruns == ["app", "app"]
              and "force_redraw" not in state, stand_in.reruns)
    finally:
        web.st, web.Ticker = original, original_ticker


def check_batch117(check: Check) -> None:
    """Batch 11.7: the version, small, first in the sidebar (its top left), as the window shows it."""
    at = new_page("-c -s:2015 --e:2026-09-01")
    check("no exception", not at.exception, at.exception)
    captions = [caption.value for caption in at.sidebar.caption]
    check("the sidebar's first caption: the version and the commit", bool(captions) and captions[0] == version_text()
          and captions[0].startswith(f"v{ratesplot.__version__}"), captions[:1])


def check_batch247(check: Check) -> None:
    """Batch 24.7: the wheel over a chart; its script's numbers and words, and what a turn sent from the browser does."""
    at = new_page("-c -s:2001 --e:2019")
    check("no exception", not at.exception, at.exception)
    check("the hint names the wheel, in the window's words", any("wheel: zoom dates" in caption.value for caption in at.caption),
          [caption.value for caption in at.caption][:3])
    script = web._CHART_JS
    check("the chart's script takes the wheel, its pause, factor and words the window's",
          'addEventListener("wheel"' in script and "__" not in script
          and f"const WHEEL_DELAY_MS = {frontend.WHEEL_DELAY_MS};" in script
          and f"const WHEEL_ZOOM_FACTOR = {frontend.WHEEL_ZOOM_FACTOR};" in script
          and f"const WHEEL_WORDS = {json.dumps(frontend.WHEEL_ZOOM_WORDS)};" in script)
    drawing = at.session_state["drawing"]
    width = struct.unpack(">II", drawing.pngs["cdn"][16:24])[0]
    geometry = drawing.geometry["cdn"]
    turn = web._wheel_zoom(drawing, "cdn", {"notches": 2, "x": 0.3})
    check("two notches away, 30 % across: the window's zoom about that place",
          turn == frontend.scale_dates_updates(geometry, frontend.WHEEL_ZOOM_FACTOR ** 2, 0.3 * width), turn)
    check("a place past the picture's edge: kept to it",
          web._wheel_zoom(drawing, "cdn", {"notches": 1, "x": 7}) == web._wheel_zoom(drawing, "cdn", {"notches": 1, "x": 1.0}))
    flung = web._wheel_zoom(drawing, "cdn", {"notches": -1e9, "x": 0.5})
    check("a fling of a billion notches out: the earliest date (1600 since batch 28.3) to today", flung == {"start": "1600-01-01", "end": ""}, flung)
    nothing = [None, {}, {"notches": 2}, {"notches": "two", "x": 0.3}, {"notches": [2], "x": 0.3},
               {"notches": float("nan"), "x": 0.3}, {"notches": float("inf"), "x": 0.3}, {"notches": 0, "x": 0.5}, "2"]
    results = [web._wheel_zoom(drawing, "cdn", wheel) for wheel in nothing]
    check("a turn that is not two numbers, or none: nothing", results == [None] * len(nothing), results)


def check_batch246(check: Check) -> None:
    """Batch 24.6: the settings remembered in a cookie, behind "Remember my settings"; an address with a chart first (written as the visit ends since 34.2)."""
    year = 365 * 86400

    def told(line: str, remember: bool) -> dict:
        """What the cookie's keeper is given on every run (batch 34.2): the last good line and the box; it writes as the visit ends."""
        return {"name": web.REMEMBER_COOKIE, "by": web.REMEMBER_BY_COOKIE, "line": line, "remember": remember, "seconds": year}

    at = new_page("-c -s:2015 --e:2026-09-01", heard=None)
    check("no exception", not at.exception, at.exception)
    box = at.checkbox(key="remember")
    check("the box, ticked at a visit's start, in the window's words", box.value is True and box.label == frontend.REMEMBER_LABEL
          and box.help == frontend.REMEMBER_HELP, box.label)
    check("one line under it saying what is kept", web._REMEMBER_LINE in [caption.value for caption in at.sidebar.caption])
    check("an address with a chart: drawn at once, the cookie not asked for", charts_shown(at) == 1 and not reader_asked(at))
    kept = memory(at)
    check("the chart drawn given to the keeper, ticked: the address's line, for a year",
          kept == told(address(at), True) and address(at) == "--C --start:2015 --end:2026-09-01", kept)
    at.checkbox(key="remember").uncheck().run()
    check("unticked: the keeper told so, nothing drawn again", memory(at) == told(address(at), False) and charts_shown(at) == 1
          and not at.exception, memory(at))
    at.checkbox(key="flag:per-capita").check().run()
    check("a chart drawn while unticked: given with the box unticked; the address follows it",
          memory(at) == told(address(at), False) and "--per-capita" in address(at), memory(at))
    at.checkbox(key="remember").check().run()
    check("ticked again: the chart shown given, ticked", memory(at) == told(address(at), True), memory(at))
    script = web._MEMORY_JS
    check("the writer encodes the line, writes only a change, and removes the cookie for the defaults' empty line",
          "encodeURIComponent(memory.line)" in script and "if (now.value !== value) write(name, value)" in script and "max-age=0" in script)
    reader = web._READER_JS
    check("the reader sends the cookie's value once, \"\" for none or where cookies are refused",
          'setTriggerValue("line", value)' in reader and "catch (error)" in reader and reader.count("setTriggerValue(") == 1)
    check("a value sent decoded as the writer encoded it (spaces, colons, commas, %)",
          web.remembered_line("--C%20--top%3A150%25%20--cdn%3Ayields%3A3m%2C10y") == "--C --top:150% --cdn:yields:3m,10y")
    check("anything but text sent: nothing remembered", [web.remembered_line(sent) for sent in (None, 7, {"line": "-u"}, ["-u"])] == [""] * 4)
    waiting = new_page(heard=None)
    check("no address, the reader not yet answered: it alone is laid out, nothing drawn",
          not waiting.exception and reader_asked(waiting) and charts_shown(waiting) == 0 and memory(waiting) is None
          and any("Looking for the settings remembered" in caption.value for caption in waiting.caption)
          and ("ready" not in waiting.session_state or not waiting.session_state["ready"]),
          [caption.value for caption in waiting.caption])
    waiting.session_state["heard"] = urllib.parse.quote("-u -s:2010 --e:2026-09-01")
    waiting.run()
    check("its answer: the visit starts from it", address(waiting) == "--U --start:2010 --end:2026-09-01" and charts_shown(waiting) == 1,
          address(waiting))
    alone = new_page("pizza")
    warnings = [w.value for w in alone.warning]
    check("nothing remembered, an address that cannot be read: the defaults, and the warning says so",
          address(alone) == "" and any("The default charts are shown instead." in w for w in warnings), warnings)
    kept_line = urllib.parse.quote("-u -s:2010 --e:2026-09-01")
    bare = new_page(heard=kept_line)
    check("no address: the settings remembered drawn, and the address rewritten to them",
          address(bare) == "--U --start:2010 --end:2026-09-01" and charts_shown(bare) == 1 and not bare.warning,
          (address(bare), [w.value for w in bare.warning]))
    shared = new_page("-c -s:2015 --e:2026-09-01", heard=kept_line)
    check("an address with a chart wins over them, so a shared link draws what it says",
          address(shared) == "--C --start:2015 --end:2026-09-01" and (memory(shared) or {}).get("line") == address(shared), address(shared))
    bad = new_page("pizza", heard=kept_line)
    warnings = [w.value for w in bad.warning]
    check("an address that cannot be read: reported, and the settings remembered drawn",
          address(bad) == "--U --start:2010 --end:2026-09-01"
          and any("pizza: not an option" in w and "settings remembered from your last visit are shown" in w for w in warnings), warnings)
    both = urllib.parse.quote("-s:2010 --e:2026-09-01")  # both nations, as the defaults
    reset = new_page(heard=both)
    button(reset, "Reset").click().run()
    check("Reset while ticked: the defaults given, their empty line removing the cookie as the visit ends",
          address(reset) == "" and memory(reset) == told("", True), memory(reset))
    unticked = new_page(heard=both)
    unticked.checkbox(key="remember").uncheck().run()
    button(unticked, "Reset").click().run()
    check("Reset while unticked: given with the box unticked, so the settings remembered survive it",
          address(unticked) == "" and memory(unticked) == told("", False), memory(unticked))
    spoilt = new_page(heard=urllib.parse.quote("--pizza -c"))
    warnings = [w.value for w in spoilt.warning]
    check("settings remembered that cannot be read: reported, and the defaults drawn (whose drawing removes the cookie)",
          any("--pizza: not an option" in w and "remembered from your last visit could not be used" in w for w in warnings)
          and charts_shown(spoilt) == 2 and (memory(spoilt) or {}).get("line") == "", warnings)


def check_batch281(check: Check) -> None:
    """Batch 28.1: greyed where they would do nothing (Min and Max, the yield terms, the quotes, the projections, Blank, All, Reset); an unticked nation's section gone, its values kept."""

    def disabled(page: AppTest, key: str) -> bool:
        return next(box for box in [*page.checkbox, *page.text_input, *page.button] if box.key == key).disabled

    at = new_page("--no-yield --e:2026-09-01")
    check("no exception", not at.exception, at.exception)
    terms = [box.key for box in at.checkbox if box.key and box.key.startswith("choice:cdn:yields:")]
    check("yields off: the yield terms greyed", terms and all(disabled(at, key) for key in terms), terms)
    check("nothing on the yield axis: Min and Max greyed, Top active",
          disabled(at, "value:cdn:min") and disabled(at, "value:cdn:max") and not disabled(at, "value:cdn:top"))
    at.checkbox(key="flag:policy").check()
    at.run()
    check("the policy rate on: Min and Max active, the yield terms still greyed",
          not disabled(at, "value:cdn:min") and not disabled(at, "value:cdn:max") and all(disabled(at, key) for key in terms))
    check("the policy rate drawn", not at.exception and charts_shown(at) == 2, at.exception)

    at = new_page("--C --e:2025-01-01")
    check("Canada alone, End too early for quotes: the quotes and the projections greyed",
          disabled(at, "flag:current") and disabled(at, "flag:projections"), flags(at))
    check("greyed, their ticks kept, and written as they are", flags(at)["current"] and flags(at)["projections"], address(at))
    at.checkbox(key="flag:us").check()
    at.run()
    check("the U.S. ticked: the quotes active (the Treasury's daily values), the projections still greyed",
          not disabled(at, "flag:current") and disabled(at, "flag:projections"))

    at = new_page("--na:ca,us --e:2026-09-01")
    check("the calendar's Blank: greyed for the blank Start, active for the End",
          disabled(at, "blank:start") and not disabled(at, "blank:end"))
    at.text_input(key="value:us:max").input("abc")
    at.run()
    check("a Max its field refuses: its All greyed", disabled(at, "all:us:max") and at.error, [e.value for e in at.error])
    at.text_input(key="value:us:max").input("8")
    at.run()
    check("a Max it takes: its All active", not disabled(at, "all:us:max") and not at.error, [e.value for e in at.error])
    at.checkbox(key="flag:us").uncheck()
    at.run()
    shown = {widget.key for widget in at.text_input}
    check("the U.S. unticked: its section gone, its Max kept",
          "value:us:max" not in shown and at.session_state["value:us:max"] == "8", sorted(key for key in shown if key.startswith("value:")))
    at.checkbox(key="flag:us").check()
    at.run()
    check("ticked again: its section back with its Max", at.text_input(key="value:us:max").value == "8", address(at))
    check("drawn with it", not at.exception and "--us:max:8" in address(at) and charts_shown(at) == 2, address(at))

    check("not at the defaults: Reset active", not button(at, "Reset").disabled)
    button(at, "Reset").click()
    at.run()
    check("at the defaults after Reset: Reset greyed", button(at, "Reset").disabled and not at.exception, address(at))
    check("End blank: the quotes and the projections active",
          not disabled(at, "flag:current") and not disabled(at, "flag:projections"))
    at.checkbox(key="size:lock").uncheck()
    at.run()
    check("the size lock unticked: Reset active", not button(at, "Reset").disabled)


def check_batch28s2(check: Check) -> None:
    """Batch 28.2 and 28.5: Unzoom goes back no further than the last Reset; the chart's link in a box, ?chart= for the defaults."""
    at = new_page("-c -s:2001 --e:2019")
    at.checkbox(key="flag:gdp").uncheck()
    at.run()
    button(at, "Reset").click()
    at.run()
    check("a Reset with no zoom or move since: Unzoom greyed, Back possible",
          not enabled(at)["Unzoom"] and enabled(at)["◀ Back"] and flags(at)["gdp"], enabled(at))
    button(at, "◀ Earlier").click()
    at.run()
    moved = address(at)
    check("a move after it: Unzoom possible", enabled(at)["Unzoom"] and "--start:" in moved, moved)
    button(at, "Unzoom").click()
    at.run()
    check("Unzoom: the charts after the Reset, GDP still ticked, Unzoom greyed",
          address(at) == "--C" and flags(at)["gdp"] and not enabled(at)["Unzoom"], (address(at), enabled(at)))
    button(at, "◀ Back").click()
    at.run()
    check("Back after Unzoom: the move again", address(at) == moved, address(at))
    for _ in range(2):  # the move, then the Reset
        button(at, "◀ Back").click()
        at.run()
    check("Back undoes the Reset as before: GDP unticked, the dates typed; nothing left to undo",
          not flags(at)["gdp"] and "--start:2001" in address(at) and not enabled(at)["◀ Back"] and not enabled(at)["Unzoom"],
          (address(at), enabled(at)))

    page = new_page("-c -s:2001 --e:2019")
    link = "?" + urllib.parse.urlencode({"chart": address(page)})
    codes = [code.value for code in page.code]
    command = next((index for index, value in enumerate(codes) if value.startswith("python ratesplot.py ")), None)
    labels = [block.value for block in page.markdown]
    check("the chart's link in a box with a copy button, under the command line since 31.3, both labelled as in the window"
          " (AppTest gives no address: its query alone)",
          link in codes and command is not None and codes.index(link) == command + 1
          and "**Command line:**" in labels and "**Web page:**" in labels
          and not any("link, to share it" in caption.value for caption in page.caption), (link, codes))
    defaults = new_page("", heard=urllib.parse.quote("-u -s:2010 --e:2026-09-01"))
    check("?chart= with nothing after it: the defaults, not the settings remembered, and its link ?chart=",
          not defaults.exception and address(defaults) == "" and charts_shown(defaults) == 2 and "?chart=" in [c.value for c in defaults.code],
          (address(defaults), charts_shown(defaults)))
    check("chart_link: the cloud frame's /~/+/ cut; a slash after a bare host; the window's address the page's own",
          web.chart_link("https://ratesplot.streamlit.app/~/+/", ("--C", "--end:2025-01-01"))
          == "https://ratesplot.streamlit.app/?chart=--C+--end%3A2025-01-01"
          and web.chart_link("http://localhost:8501", ()) == "http://localhost:8501/?chart="
          and web.chart_link(WEB_PAGE_ADDRESS, ("--C",)) == "https://ratesplot.streamlit.app/?chart=--C")


def check_batch283(check: Check) -> None:
    """Batch 28.3: dates from 1600-01-01, the earliest a chart shows; one before it refused in its field."""
    at = new_page("-c -s:1600 --e:2019")
    check("Start 1600: drawn, from 1600", not at.exception and charts_shown(at) == 1 and not at.error
          and "--start:1600" in address(at), address(at))
    check("at 1600, Earlier greyed", not enabled(at)["◀ Earlier"], enabled(at))
    check("the calendar's first day 1600-01-01", web._CALENDAR_FIRST == datetime.date(1600, 1, 1), web._CALENDAR_FIRST)
    at.text_input(key="value:start").input("1599-12")
    at.run()
    errors = [e.value for e in at.error]
    check("Start 1599-12: not drawn, its field red, the line naming the earliest date",
          any("Start: start date '1599-12' is before 1600-01-01, the earliest a chart shows" in e for e in errors)
          and at.text_input(key="value:start").label.startswith(":red["), errors)
    at = new_page("-c -s:1700 --e:2019")
    button(at, "Zoom out").click()
    at.run()
    # Twice 319 years is more than the 426 from 1600 to today: the whole of it, and nothing further out.
    check("Zoom out from 1700: from 1600 to today, Earlier and Zoom out greyed", not at.exception
          and address(at) == "--C --start:1600-01-01" and not enabled(at)["◀ Earlier"] and not enabled(at)["Zoom out"],
          (address(at), enabled(at)))


def check_batch284(check: Check) -> None:
    """Batch 28.4: on a narrow screen the page draws 800x800 with larger text; the link and the command line unchanged."""

    def png_size(page: AppTest) -> tuple[int, int]:
        png = chart_elements(page)[0].proto.bytes
        return int.from_bytes(png[16:20], "big"), int.from_bytes(png[20:24], "big")

    at = new_page("-c --e:2026-09-01")
    check("no word from the page's script: the full-size chart", png_size(at) == (2048, 1536), png_size(at))
    at.session_state[web._NARROW] = True  # what the width's script says on a phone (AppTest runs no script)
    at.run()
    drawn = at.session_state[web._DRAWING].drawn
    captions = [caption.value for caption in at.caption if caption.value.startswith("Drawn")]
    check("narrow: 800x800, the font floor 0.8, the line under the chart saying so",
          png_size(at) == (800, 800) and drawn.font_floor == 0.8 and drawn.font_scale == 0.8
          and any("for a narrow screen: 800 x 800 px and larger text" in c for c in captions), (png_size(at), captions))
    check("narrow: the address, the link and the command line as before", address(at) == "--C --end:2026-09-01"
          and "?chart=--C+--end%3A2026-09-01" in [code.value for code in at.code]
          and "python ratesplot.py --C --end:2026-09-01" in [code.value for code in at.code], address(at))
    at.text_input(key="size:width").input("1200")
    at.run()
    drawn = at.session_state[web._DRAWING].drawn
    check("narrow, a size given: that size, the floor kept",
          (drawn.width_px, drawn.height_px, drawn.font_floor) == (1200, 900, 0.8) and "1200x900" in address(at),
          (drawn.width_px, drawn.height_px, address(at)))
    at.session_state[web._NARROW] = False
    at.run()
    check("wide again: the size given, the floor 0.5", at.session_state[web._DRAWING].drawn.font_floor == 0.5, address(at))
    config = PlotConfig()
    check("for_screen: unchanged when wide; the phone's canvas only at the default size",
          web.for_screen(config, False) == config
          and (web.for_screen(config, True).width_px, web.for_screen(config, True).height_px) == (800, 800)
          and web.for_screen(replace(config, width_px=1600, height_px=1200), True).width_px == 1600)


def check_batch291(check: Check) -> None:
    """Batch 29.1: each "?" tip says why its control is greyed and ends with More to its topic; the help view at ?help=TOPIC."""
    at = new_page("-c --no-debt --no-GDP --no-interest --e:2026-09-01")
    box = at.checkbox(key="flag:relative")
    check("% of GDP greyed, its tip saying why, then More to its topic", not at.exception and box.disabled
          and "Greyed because it acts on debt and interest, and neither is ticked." in box.help
          and box.help.endswith("[More](?help=percent-of-gdp)"), box.help)
    # Every control names a topic since batch 29.3: Regression segments' is
    # taken away for the while, so a control with nothing more to say is
    # still seen to have no More.
    with topic_taken_away("regression"):
        segments = new_page("-c --no-debt --no-GDP --no-interest --e:2026-09-01").checkbox(key="flag:regression")
    check("Regression segments greyed, with no topic: why, and no More",
          "Greyed because" in segments.help and "[More]" not in segments.help, segments.help)
    check("a box not greyed: no reason", "Greyed because" not in at.checkbox(key="flag:cdn").help)
    margin = at.text_input(key="value:cdn:prime-margin")
    check("a nation's field greyed: why", margin.disabled and "Greyed because “Mortgage rates” is unticked." in margin.help, margin.help)
    at = new_page("-c --max:8 --e:2026-09-01")
    tick = at.checkbox(key="all:cdn:max")
    check("an “All” with one nation ticked: why it is greyed", tick.disabled and "Greyed because only one nation is ticked." in tick.help,
          tick.help)
    check("the style rule that makes each “?” an ⓘ, scoped to the help buttons",
          'button[aria-label^="Help for"]' in web._INFO_ICON_STYLE and "ⓘ" in web._INFO_ICON_STYLE)
    view = AppTest.from_file(APP, default_timeout=TIMEOUT_SECONDS)
    view.query_params["help"] = "percent-of-gdp"
    view.run()
    check("?help=percent-of-gdp: the topic alone, no control, no chart", not view.exception
          and any("## When the ratio is taken" in block.value for block in view.markdown)
          and not view.checkbox and not view.text_input and not chart_elements(view), [block.value[:40] for block in view.markdown])
    view = AppTest.from_file(APP, default_timeout=TIMEOUT_SECONDS)
    view.query_params["help"] = "../README"
    view.run()
    check("?help= a name no option gives: refused, nothing else read", not view.exception
          and [error.value for error in view.error] == ["There is no help page named “../README”."], [e.value for e in view.error])


def check_batch292(check: Check) -> None:
    """Batch 29.2: the curves', per capita's, the rates' and the nations' tips end with More to their topics; a topic's link to another is the page's own address."""
    at = new_page("-c --e:2026-09-01")
    flags = {"gdp": "gdp", "debt": "debt", "interest": "interest", "per-capita": "per-capita",
             # 29.2b, the rates' pages.
             "yield": "yields", "policy": "policy-rates", "mortgages": "mortgages", "spreads": "spreads", "current": "latest-values",
             # 29.2c, the nations' pages.
             "cdn": "canada", "us": "united-states", "uk": "united-kingdom", "germany": "germany"}
    for name, topic in flags.items():
        box = at.checkbox(key=f"flag:{name}")
        check(f"{name}: its tip ends with More to its topic", not at.exception and box.help.endswith(f"[More](?help={topic})"), box.help[-50:])
    # A list's help is on its label alone (``web._choices_control``).
    labels = {block.value: block.proto.help for block in at.markdown}
    for label, topic in (("Yield terms", "yields"), ("Policy rates", "policy-rates"), ("Mortgage terms", "mortgages"), ("Spreads", "spreads")):
        tip = labels.get(f"**{label}**", "")
        check(f"{label}: its label's tip ends with More to its topic", tip.endswith(f"[More](?help={topic})"), tip[-50:])
    margin = at.text_input(key="value:cdn:prime-margin")
    check("the prime margin, greyed: its reason, then More to the mortgages' topic",
          "Greyed because" in margin.help and margin.help.endswith("[More](?help=mortgages)"), margin.help[-80:])
    tracker = new_page("--na:gb --e:2026-09-01").text_input(key="value:uk:tracker-margin")
    check("the UK's tracker margin: More to the mortgages' topic", tracker.help.endswith("[More](?help=mortgages)"), tracker.help[-50:])
    for topic, link in (("yields", "[Yield spreads](?help=spreads)"), ("policy-rates", None), ("mortgages", "[Policy rates](?help=policy-rates)"),
                        ("spreads", "[Bond yields](?help=yields)")):
        view = AppTest.from_file(APP, default_timeout=TIMEOUT_SECONDS)
        view.query_params["help"] = topic
        view.run()
        shown = "\n".join(block.value for block in view.markdown)
        check(f"?help={topic}: a section per nation" + (", its link to another topic" if link else ""), not view.exception
              and all(f"## {nation}" in shown for nation in ("Canada", "The United States", "The United Kingdom", "Germany"))
              and (link is None or link in shown), shown[:60])
    view = AppTest.from_file(APP, default_timeout=TIMEOUT_SECONDS)
    view.query_params["help"] = "latest-values"
    view.run()
    shown = "\n".join(block.value for block in view.markdown)
    check("?help=latest-values: the Treasury's figures and the quotes", not view.exception and "# Latest values" in shown
          and "## The day's quotes, every nation" in shown, shown[:60])
    view = AppTest.from_file(APP, default_timeout=TIMEOUT_SECONDS)
    view.query_params["help"] = "debt"
    view.run()
    shown = "\n".join(block.value for block in view.markdown)
    check("?help=debt: the topic alone, a section per nation", not view.exception and not view.checkbox
          and all(f"## {nation}" in shown for nation in ("Canada", "The United States", "The United Kingdom", "Germany")), shown[:60])
    check("its links to other topics written as the page's own address, followed in the help view",
          "[By level](?help=levels)" in shown and "[As % of GDP](?help=percent-of-gdp)" in shown)
    # 29.2c: each nation's page links to the curves' pages and to the other nations'.
    for topic, title in (("canada", "Canada"), ("united-states", "The United States"), ("united-kingdom", "The United Kingdom"),
                         ("germany", "Germany")):
        view = AppTest.from_file(APP, default_timeout=TIMEOUT_SECONDS)
        view.query_params["help"] = topic
        view.run()
        shown = "\n".join(block.value for block in view.markdown)
        others = [other for other in ("canada", "united-states", "united-kingdom", "germany") if other != topic]
        check(f"?help={topic}: its title, its links to the curves' pages and to the other nations'", not view.exception
              and f"# {title}" in shown.splitlines() and all(f"(?help={curve})" in shown for curve in ("gdp", "debt", "yields", "mortgages"))
              and all(f"(?help={other})" in shown for other in others), shown[:60])


def check_batch293(check: Check) -> None:
    """Batch 29.3: the other controls' tips end with More to their topics; the eight pages, their titles and links."""
    at = new_page("-c --e:2026-09-01")
    tips = {
        "flag:label": "end-labels", "flag:regression": "regression", "value:regression-tolerance": "regression",
        "size:width": "layout", "size:height": "layout", "value:start": "dates", "value:end": "dates",
        "value:cdn:min": "limits", "value:cdn:max": "limits", "value:cdn:top": "limits", "value:cdn:bottom": "limits",
        "flag:auto-refresh": "auto-refresh", "value:auto-refresh-minutes": "auto-refresh", "flag:projections": "projections",
        "flag:hover": "hover",
    }
    for key, topic in tips.items():
        widget = at.checkbox(key=key) if key.startswith("flag:") else at.text_input(key=key)
        check(f"{key}: its tip ends with More to its topic", not at.exception and widget.help.endswith(f"[More](?help={topic})"),
              widget.help[-50:])
    pages = (
        ("dates", "Start and End (-s, -e)", ("policy-rates", "projections", "latest-values", "limits")),
        ("limits", "Axis limits: Min, Max, Top and Bottom", ("per-capita", "percent-of-gdp", "canada", "germany")),
        ("regression", "Regression segments (--reg)", ("projections", "percent-of-gdp", "limits")),
        ("projections", "Projections to today", ("end-labels", "hover", "per-capita", "latest-values")),
        ("auto-refresh", "Automatic refresh (--au)", ("latest-values",)),
        ("hover", "Hover read-out", ("projections", "latest-values", "end-labels")),
        ("end-labels", "Latest value for each line (-l)", ("projections", "limits", "layout")),
        ("layout", "Size and layout of the chart", ("dates", "limits", "end-labels", "regression")),
        # The earlier pages, linking to the new ones.
        ("latest-values", "Latest values", ("projections", "auto-refresh")),
        ("percent-of-gdp", "As % of GDP (-r)", ("regression", "limits")),
    )
    for topic, title, links in pages:
        view = AppTest.from_file(APP, default_timeout=TIMEOUT_SECONDS)
        view.query_params["help"] = topic
        view.run()
        shown = "\n".join(block.value for block in view.markdown)
        check(f"?help={topic}: its title and its links to {', '.join(links)}", not view.exception
              and f"# {title}" in shown.splitlines() and all(f"(?help={link})" in shown for link in links), shown[:60])


def check_batch31(check: Check) -> None:
    """Batch 31: the window and the page alike: the buttons' tips from frontend, as the window's; Export chart, not Download."""
    at = new_page("-c -s:2001 --e:2019")
    tips = {b.label: b.help for b in at.button}
    page = {"Refresh": web.REFRESH_HELP, "Reset": web.RESET_HELP, "◀ Back": web.BACK_HELP, "Unzoom": web.UNZOOM_HELP}
    check("Refresh, Reset, Back and Unzoom: the tips the window's buttons have", not at.exception
          and all(tips.get(label) == text for label, text in page.items()), {label: tips.get(label) for label in page})
    exports = [b for b in at.get("download_button") if b.proto.id.endswith("download:cdn")]
    check("the chart's download button reads Export chart, with the window's tip", len(exports) == 1
          and exports[0].proto.label == "Export chart" and exports[0].proto.help == web.EXPORT_HELP,
          [(b.proto.label, b.proto.help) for b in exports])
    # 31.1: the sidebar width's keeper, whose script puts back the width kept while the box is unticked (AppTest runs no script).
    check("31.1: the sidebar width's keeper laid out, told the box is ticked", sidebar_keeper(at) == {"remember": True, "key": "sidebarWidth"},
          sidebar_keeper(at))
    at.checkbox(key="remember").uncheck().run()
    check("the box unticked: the keeper told so, nothing drawn", not at.exception and sidebar_keeper(at) == {"remember": False, "key": "sidebarWidth"},
          sidebar_keeper(at))
    at.checkbox(key="remember").check().run()
    check("ticked again: told so", sidebar_keeper(at) == {"remember": True, "key": "sidebarWidth"}, sidebar_keeper(at))
    check("the line under the box names the width", any("sidebar's width" in caption.value for caption in at.sidebar.caption))
    bad = new_page("-c --start:nonsense")
    check("laid out with no chart drawn too", sidebar_keeper(bad) is not None, sidebar_keeper(bad))


def check_batch324(check: Check) -> None:
    """Batch 32.4: a blank End reaches the last day a policy target drawn holds; the field shows it in grey, the address writes none."""
    later = pd.Timestamp.today().normalize() + pd.Timedelta(days=2)
    real = plotting.policy_target_last
    try:
        plotting.policy_target_last = lambda _key: later  # a change announced to take effect in two days
        at = new_page("-c --policy -s:2026-06-01")
        drawn = at.session_state[web._DRAWING].drawn
        end = at.text_input(key="value:end")
        check("a blank End, Canada's target drawn: the chart ends on the target's last day, two days ahead",
              not at.exception and drawn.end == later, drawn.end)
        check("… the End field blank, its grey value that day; the address without an End",
              end.value == "" and end.proto.placeholder == f"{later:%Y-%m-%d}" and "--end" not in address(at),
              (end.proto.placeholder, address(at)))
        at = new_page("-c -s:2026-06-01")
        check("no policy rate drawn: the blank End is today",
              at.session_state[web._DRAWING].drawn.end == pd.Timestamp.today().normalize(), at.session_state[web._DRAWING].drawn.end)
    finally:
        plotting.policy_target_last = real


def check_batch34(check: Check) -> None:
    """Batch 34: Go beside the chart's link, opening it in a new tab (34.1); the settings remembered as the visit ends (34.2)."""

    def goes(page: AppTest) -> list[tuple[str, str, str]]:
        return [(go.proto.label, go.proto.url, go.proto.help) for go in page.get("link_button")]

    at = new_page("-c -s:2001 --e:2019")
    link = "?" + urllib.parse.urlencode({"chart": address(at)})
    check("34.1: one Go, opening the link in the box beside it (AppTest gives no address: its query alone), with its tip",
          not at.exception and goes(at) == [("Go", link, web._GO_HELP)] and link in [code.value for code in at.code], goes(at))
    at.checkbox(key="flag:per-capita").check().run()
    check("a new drawing: Go opens its link", goes(at) == [("Go", "?" + urllib.parse.urlencode({"chart": address(at)}), web._GO_HELP)]
          and "--per-capita" in address(at), goes(at))
    defaults = new_page("")
    check("the defaults: Go opens ?chart=, not the bare address that draws what the browser remembers",
          [url for _, url, _ in goes(defaults)] == ["?chart="], goes(defaults))

    # 34.2: the settings remembered written as the visit ends (AppTest runs no
    # script: what the keeper is given, and the script's parts; its behaviour
    # was checked in Chrome, tab by tab).
    at = new_page("-c -s:2001 --e:2019")
    good = address(at)
    at.text_input(key="value:start").input("nonsense").run()
    kept = memory(at)
    check("34.2: a field that stops the drawing: the keeper still given the last good chart's line",
          not at.exception and at.error and kept is not None and kept["line"] == good and kept["remember"] is True, kept)
    real = web._draw
    try:
        web._draw = lambda config, tokens: replace(real(config, tokens), error="a drawing made to fail")
        failed = new_page("-u -s:2001 --e:2019")
    finally:
        web._draw = real
    check("… and a visit whose only drawing failed: laid out, given none, so nothing is written",
          (memory(failed) or {}).get("line", "x") is None and failed.session_state[web._DRAWING].error is not None, memory(failed))
    script = web._MEMORY_JS
    check("the script acts when the page is hidden or left, never on a run", 'addEventListener("pagehide", memory.leave)' in script
          and 'document.visibilityState === "hidden"' in script and script.count("memory.leave()") == 1
          and "if (memory.leave !== undefined) return;" in script, good)
    check("it puts back what was there only while its own mark and write are there; the mark a cookie of its own",
          "now.by === memory.mark && now.value === memory.written" in script and "write(by, memory.before.by)" in script
          and web.REMEMBER_BY_COOKIE != web.REMEMBER_COOKIE and web.REMEMBER_BY_COOKIE.startswith(web.REMEMBER_COOKIE))
    check("the box's ⓘ and the line under it say the box as you leave decides",
          "as you leave" in frontend.REMEMBER_HELP and "as you leave the page" in web._REMEMBER_LINE)


PAGE_CHECKS = (
    check_address_and_rules, check_bad_addresses, check_view_buttons, check_calendars, check_size_boxes, check_choice_boxes,
    check_nation_sections, check_uk, check_de, check_none_lists, check_greyed, check_wrong_fields, check_layout,
    check_batch15, check_batch9, check_batch94, check_batch17, check_batch91, check_batch95, check_batch104, check_batch105,
    check_batch18, check_batch11, check_batch112, check_batch113, check_batch117, check_batch247, check_batch246,
    check_batch281, check_batch28s2, check_batch283, check_batch284, check_batch291, check_batch292, check_batch293,
    check_batch31, check_batch324, check_batch34,
)


def check_page() -> bool:
    """Run every page check; True if all passed."""
    results: list[bool] = []

    def check(label: str, condition: bool, detail: object = "") -> None:
        results.append(bool(condition))
        shown = f"  [{detail}]" if not condition or detail else ""
        print(f"{'ok  ' if condition else 'FAIL'} {label}{shown}", flush=True)

    for group in PAGE_CHECKS:
        print(f"-- {group.__doc__.splitlines()[0]}")
        group(check)
    print(f"[web] page: {sum(results)} of {len(results)} checks passed")
    return all(results)


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--only", choices=("charts", "page"), help="run one part only")
    parser.add_argument("--cases", help=f"comma-separated regression cases for the charts part (default: {','.join(CHART_CASES)})")
    parser.add_argument("--cache", type=Path, help="download cache folder (default out/download_cache, or out/synthetic_cache)")
    parser.add_argument(
        "--synthetic", action="store_true", help="made-up data in each source's format, no network (tools/synthetic_sources.py)"
    )
    parser.add_argument("--verbose", action="store_true", help="show the command line's own output")
    parser.add_argument("--list", action="store_true", help="list the chart cases and exit")
    args = parser.parse_args(argv)
    cases = tuple(args.cases.split(",")) if args.cases else CHART_CASES
    if args.list:
        for name in cases:
            print(f"{name:12} {' '.join(CASES[name])}")
        return 0

    if args.synthetic:
        print("[web] SYNTHETIC data (tools/synthetic_sources.py): proves code paths, not the data")
        synthetic_sources.install_synthetic_cache(http, (args.cache or ROOT / "out" / "synthetic_cache").resolve())
    else:
        install_disk_cache(http, (args.cache or ROOT / "out" / "download_cache").resolve())
    passed = True
    if args.only in (None, "charts"):
        print(f"[web] charts: the page's drawing against the command line's, {len(cases)} case(s)")
        same = check_charts(cases, ROOT / "out" / "verify_web", verbose=args.verbose)
        print("[web] charts: ALL IDENTICAL" if same else "[web] charts: SOME DIFFER")
        passed &= same
    if args.only in (None, "page"):
        print("[web] page: Streamlit's AppTest")
        passed &= check_page()
    return 0 if passed else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
