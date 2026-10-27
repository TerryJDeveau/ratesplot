#!/usr/bin/env python3
"""Check Export chart (``ratesplot/export.py``, batch 11.1): the data written against the lines drawn.

Each case is drawn as ``tools/regress_charts.py`` draws it (``--e:2026-09-01``
first, its download cache), through ``plotting.build_figure``, as the window,
the page and ``--save`` draw it. For each chart:

* the rows: one for each date in the window on which a curve drawn
  (``export.mark_series``) has a point, and no other;
* the columns: one for each curve marked, and none for a line not marked
  (the regression segments, their bands, the leaders);
* every point of every curve in the window: its value is the exported value
  on its date (where two pieces of a curve meet on a date, either's);
* every date between two points: a curve drawn as steps exports its value in
  force (its last point's), one drawn as a line a value between the two
  points' (the straight line between them, as the axis draws it);
* the CSV read back: every value to 15 significant digits, every blank where
  the chart has no line.

Then ``--save`` itself, from the command line (``cli.main``): a PNG per chart
equal, byte for byte, to the regression set's (``tools/verify_charts.run``),
and an XLSX read back with openpyxl, a sheet per chart named as its tab, the
dates as dates, each value the chart's.

Usage (from the project root)::

    python tools/check_export.py                    # the cases below
    python tools/check_export.py --cases default,c_p
    python tools/check_export.py --synthetic        # made-up data, no network

Exit code 1 on any FAIL.
"""

from __future__ import annotations

import argparse
import contextlib
import csv
import io
import math
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / "tools")]

import verify_charts  # noqa: E402  (first: it sets the Agg backend before pyplot is imported)
import matplotlib.dates as mdates  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

import synthetic_sources  # noqa: E402
from regress_charts import CASES, PINNED_END, install_disk_cache  # noqa: E402

from ratesplot import cli, export, http  # noqa: E402
from ratesplot.config import nation_by_key  # noqa: E402
from ratesplot.console import this_thread_output_to  # noqa: E402
from ratesplot.plotting import COUNTRIES, build_figure, resolve_start  # noqa: E402

# The defaults; Canada from 1867 (its monthly yields, then its daily ones) with
# --reg; -r with --reg on a log axis; levels under -r; per capita; the UK and
# Germany with their rates, spreads and mortgages; the four policy rates with
# the U.S. target range as a band; all four nations.
DEFAULT_CASES = ("default", "c_reg1867", "u_reg_r", "c_lv_r", "c_p", "uk_rates", "de_rates", "policy_both", "all4")
# --save is run from the command line on these: two nations, then four with their own levels.
SAVE_CASES = ("default", "nat_levels")
OUT = ROOT / "out" / "check_export"
DIGITS = 15  # the significant digits a CSV's numbers are written to (export._text)


def _drawn(tokens: list[str]) -> list[tuple[str, object, object]]:
    """Each chart a case draws: (nation key, figure, the nation's config), as the window draws them, their output dropped."""
    config = cli.parse_args(["--no-gui", *tokens])
    with this_thread_output_to(None):
        drawn = resolve_start(config)
        return [
            (country.key, build_figure(country, drawn), drawn.for_nation(country.key))
            for country in COUNTRIES
            if getattr(drawn, country.show_field)
        ]


def _same(a: float, b: float, digits: int = DIGITS) -> bool:
    return a == b or math.isclose(a, b, rel_tol=10.0 ** (1 - digits), abs_tol=1e-300)


def check_chart(check, label: str, figure, config, key: str) -> None:
    """The checks of one chart's exported data against its drawn lines (the module docstring's list)."""
    data = export.chart_data(figure, config, key)
    pieces = export._pieces(figure)
    start, end = (mdates.date2num(pd.Timestamp(value).to_pydatetime()) for value in (config.start, config.end))
    unmarked = [line for ax in figure.axes for line in ax.get_lines()
                if getattr(line, export._SERIES_ATTR, None) is None and len(line.get_xdata()) > 1]
    check(f"{label}: a column for each curve marked, none for the {len(unmarked)} lines not marked",
          data.names == tuple(name for name, _axis in pieces) and len(set(data.names)) == len(data.names), data.names)
    points = np.unique(np.concatenate([piece.x for each in pieces.values() for piece in each]))
    points = points[(points >= start) & (points <= end)]
    wanted = [pd.Timestamp(mdates.num2date(day).replace(tzinfo=None)).normalize() for day in points]
    check(f"{label}: a row for each date a curve has a point on, {len(wanted)}, and no other", list(data.dates) == wanted,
          (len(data.dates), len(wanted)))
    row_of = {day: row for row, day in enumerate(points)}
    wrong: list[str] = []
    between: list[str] = []
    for column, ((name, _axis), each) in enumerate(pieces.items()):
        exported = np.array([np.nan if values[column] is None else values[column] for values in data.values])
        on_date: dict[float, list[float]] = {}
        for piece in each:
            for x, y in zip(piece.x, piece.y):
                if start <= x <= end and np.isfinite(y):
                    on_date.setdefault(x, []).append(float(y))
        for x, ys in on_date.items():
            value = exported[row_of[x]]
            if not any(_same(value, y) for y in ys):
                wrong.append(f"{name} {mdates.num2date(x):%Y-%m-%d}: {value} not {ys}")
        # Between two points of a piece: the value in force (steps), or one between the two (a line).
        for piece in each:
            inside = (points > piece.x[0]) & (points < piece.x[-1])
            for day in points[inside]:
                if day in on_date:
                    continue  # a point of this curve on the date: checked above
                at = np.searchsorted(piece.x, day, side="right") - 1
                if not (np.isfinite(piece.y[at]) and np.isfinite(piece.y[at + 1])):
                    continue  # a break in the line
                value = exported[row_of[day]]
                low, high = sorted((piece.y[at], piece.y[at + 1]))
                ok = _same(value, piece.y[at]) if piece.steps else low - 1e-12 * abs(low) <= value <= high + 1e-12 * abs(high)
                if not ok:
                    between.append(f"{name} {mdates.num2date(day):%Y-%m-%d}: {value} ({piece.y[at]}, {piece.y[at + 1]})")
    check(f"{label}: every point's value exported on its date", not wrong, wrong[:3])
    check(f"{label}: every date between two points: the step's value, or the line's", not between, between[:3])
    rows = list(csv.reader(io.StringIO(export.delimited_bytes(data, ",").decode("utf-8-sig"))))
    body = rows[4:]
    read_back = all(
        (cell == "") == (value is None) and (cell == "" or _same(float(cell), value))
        for row, values in zip(body, data.values) for cell, value in zip(row[1:], values)
    )
    check(f"{label}: the CSV read back: {len(body)} rows, each value to {DIGITS} digits, each blank a blank",
          len(body) == len(data.dates) and read_back and [row[0] for row in body] == [f"{d:%Y-%m-%d}" for d in data.dates])


def check_save(check, name: str) -> None:
    """``--save`` from the command line: each chart's PNG the regression set's, and the XLSX read back."""
    from openpyxl import load_workbook

    tokens = [PINNED_END, *CASES[name]]
    folder = OUT / name
    shutil.rmtree(folder, ignore_errors=True)
    folder.mkdir(parents=True)
    with contextlib.redirect_stdout(io.StringIO()):
        regression = verify_charts.run(folder / "regression", tokens)
        cli.main([*tokens, f"--save:{folder / 'rates.png'}"])
        cli.main([*tokens, f"--save:{folder / 'rates.xlsx'}"])
    drawn = list(_drawn(tokens))
    codes = [nation_by_key(key).codes[0] for key, _figure, _config in drawn]
    saved = [folder / f"rates_{code}.png" for code in codes]
    check(f"{name}: --save:rates.png wrote {', '.join(path.name for path in saved)}", all(path.exists() for path in saved),
          sorted(path.name for path in folder.iterdir()))
    same = [path.exists() and path.read_bytes() == chart.read_bytes() for path, chart in zip(saved, regression)]
    check(f"{name}: each PNG the regression set's, byte for byte", len(saved) == len(regression) and all(same), same)
    book = load_workbook(folder / "rates.xlsx")
    titles = [nation_by_key(key).name for key, _figure, _config in drawn]
    check(f"{name}: --save:rates.xlsx, a sheet per chart named as its tab", book.sheetnames == titles, book.sheetnames)
    for (key, figure, config), title in zip(drawn, titles):
        if title not in book.sheetnames:
            continue
        data = export.chart_data(figure, config, key)
        sheet = book[title]
        rows = list(sheet.iter_rows(values_only=True))
        header = rows[0][0] == data.title and rows[2][1:] == data.names and rows[3][0] == "(YYYY-MM-DD)"
        dates = [row[0] for row in rows[4:]] == [date.to_pydatetime() for date in data.dates]
        styled = all(cell.number_format == "yyyy-mm-dd" for cell in sheet["A"][4:])
        # Each number to the CSV's 15 digits, exactly (export._cell).
        values = all(
            tuple(row[1:1 + len(data.names)]) == tuple(None if v is None else float(format(v, f".{DIGITS}g")) for v in values)
            for row, values in zip(rows[4:], data.values)
        )
        check(f"{name}: sheet {title}: title, names, {len(data.dates)} dates shown YYYY-MM-DD, every value the CSV's",
              header and dates and styled and values and len(rows) == 4 + len(data.dates),
              (header, dates, styled, values, len(rows)))


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--cases", help=f"comma-separated regression cases (default: {','.join(DEFAULT_CASES)})")
    parser.add_argument("--cache", type=Path, help="download cache folder (default out/download_cache, or out/synthetic_cache)")
    parser.add_argument(
        "--synthetic", action="store_true", help="made-up data in each source's format, no network (tools/synthetic_sources.py)"
    )
    args = parser.parse_args(argv)
    cases = tuple(args.cases.split(",")) if args.cases else DEFAULT_CASES
    unknown = [name for name in cases if name not in CASES]
    if unknown:
        parser.error(f"unknown case: {', '.join(unknown)}")
    if args.synthetic:
        print("[export] SYNTHETIC data (tools/synthetic_sources.py): proves code paths, not the data")
        synthetic_sources.install_synthetic_cache(http, (args.cache or ROOT / "out" / "synthetic_cache").resolve())
    else:
        install_disk_cache(http, (args.cache or ROOT / "out" / "download_cache").resolve())

    results: list[bool] = []

    def check(label: str, condition: bool, detail: object = "") -> None:
        results.append(bool(condition))
        shown = f"  [{detail}]" if not condition and detail != "" else ""
        print(f"{'ok  ' if condition else 'FAIL'} {label}{shown}", flush=True)

    for name in cases:
        print(f"-- {name}: {' '.join(CASES[name]) or '(the defaults)'}")
        for key, figure, config in _drawn([PINNED_END, *CASES[name]]):
            check_chart(check, f"{name} {key}", figure, config, key)
    for name in SAVE_CASES if not args.cases else ():
        print(f"-- --save on {name}: {' '.join(CASES[name]) or '(the defaults)'}")
        check_save(check, name)
    print(f"[export] {sum(results)} of {len(results)} checks passed")
    return 0 if all(results) else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
