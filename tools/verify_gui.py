#!/usr/bin/env python3
"""Check the window (``ratesplot/gui.py``): its charts against the command line's, and how it behaves.

Two parts, both run by default:

* ``charts``: for a few regression cases the window is opened on the case as
  the command line opens it (``cli.parse_choices``, ``config_from_choices``,
  ``frontend.starting_choices``, ``RatesPlotApp``), and each chart it draws
  is hashed against the command line's (``tools/verify_charts.run``, into
  ``out/verify_gui/NAME/``): IDENTICAL or DIFFER. The command line the window
  writes must draw what it came from (``options.drawn_config`` per nation).
* ``window``: the window driven as a person drives it, through the calls its
  controls make (a box ticked: its variable set, then ``_flag_changed``; a
  field typed: its variable set, the redraw waiting as after a keystroke; a
  drag on a chart: ``_zoom_to_box``), group by group. Each check prints ``ok``
  or ``FAIL``: the controls' states and grey-outs, the fields a view writes,
  the message log's lines, and the written command line drawing what the
  window shows.

The window runs in this process, with a scratch ``APPDATA`` set before the
program is imported, so the settings a person's own window remembers are
never read or written. Downloads come from the regression tool's disk cache
(``--cache``, default ``out/download_cache``, shared with
``tools/regress_charts.py``) and every case pins its end (``--e:2026-09-01``),
so nothing is fetched live once the cache is warm, but for the ``projections``
group, whose chart must end today for anything to be projected (it draws with
``--no-cur``, so no quote of the moment can make two drawings differ);
``--synthetic`` draws from made-up data instead (``tools/synthetic_sources.py``). Needs Tk and a display:
Windows has both; on Linux, ``xvfb-run``.

The groups grew from the scripts that tested the window in batches 5 to 18
(2026-10-13 to 2026-10-23, kept outside the repository), each brought up to
date with the program. Batch 11 adds its window features' tests here (the
export dialog, the hover read-out, automatic refresh). Run it after any change
to ``gui.py`` or ``frontend.py``, as ``tools/verify_web.py`` after one to
``web.py``; it is not part of every regression run.

Usage (from the project root)::

    python tools/verify_gui.py                          # both parts (a few minutes once the cache is warm)
    python tools/verify_gui.py --only window            # or --only charts
    python tools/verify_gui.py --group limits,levels    # some of the window's groups
    python tools/verify_gui.py --list                   # the chart cases and the groups
    python tools/verify_gui.py --synthetic              # made-up data, no network

Exit code 1 on any DIFFER or FAIL.
"""

from __future__ import annotations

import argparse
import contextlib
import csv
import dataclasses
import hashlib
import io
import os
import shutil
import sys
import tempfile
import time
import types
from pathlib import Path
from typing import Callable

# The window's buttons are labelled "◀ Back" and the like; a console or pipe
# in the Windows code page cannot encode them. Escape what it cannot show
# rather than fail (before ratesplot is imported: its output router writes
# through the stream that is sys.stdout now).
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="backslashreplace")

# A scratch APPDATA before anything reads it: the window loads and saves its
# settings there (frontend.state_folder), never in the person's own.
SCRATCH_APPDATA = tempfile.mkdtemp(prefix="ratesplot_verify_gui_")
os.environ["APPDATA"] = SCRATCH_APPDATA

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / "tools")]

import verify_charts  # noqa: E402  (first: it sets the Agg backend before pyplot is imported)
import tkinter as tk  # noqa: E402
from tkinter import ttk  # noqa: E402

import matplotlib.dates as mdates  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from matplotlib.backend_bases import MouseEvent  # noqa: E402
from matplotlib.backends.backend_agg import FigureCanvasAgg  # noqa: E402
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg  # noqa: E402
from matplotlib.figure import Figure  # noqa: E402

import synthetic_sources  # noqa: E402
from regress_charts import CASES, PINNED_END, install_disk_cache  # noqa: E402

from ratesplot import cli, export, frontend, gui, hover, http, options  # noqa: E402
from ratesplot.axes import format_yield  # noqa: E402
from ratesplot.config import NATIONS  # noqa: E402
from ratesplot.console import this_thread_output_to  # noqa: E402
from ratesplot.plotting import COUNTRIES, build_figure, resolve_start  # noqa: E402

# Regression cases drawn both ways: the defaults, a later start, -r with --reg,
# levels under -r, the UK and Germany with their curves, all four nations,
# both policy terms, each nation's own levels, Top and Bottom in a currency.
CHART_CASES = (
    "default", "c2015", "u_reg_r", "c_lv_r", "uk_rates", "de_rates", "all4", "policy_both", "nat_levels", "nat_fx",
)

NATION_KEYS = tuple(nation.key for nation in NATIONS)  # in drawing order, as the window's tabs
# Terry's window on Ambergris (its gui_state.json, 2026-09-29): the size a new control is checked at.
SIZE = "1728x972+40+40"
SETTLE_SECONDS = 300  # a drawing whose downloads are not cached yet takes about 15 s; Germany's longer
END = PINNED_END  # every window here ends on the regression set's pinned date

Check = Callable[..., None]


# ---------------------------------------------------------------------------
# The window, opened in this process
# ---------------------------------------------------------------------------


class Window:
    """A window opened on a command line, as ``gui.run_gui`` opens it, and the calls its controls make."""

    def __init__(self, tokens: str, size: str = SIZE, settle: bool = True) -> None:
        payloads, flags, _parser = cli.parse_choices(tokens.split())
        config = cli.config_from_choices(payloads, flags)
        choices, notes = frontend.starting_choices(config, payloads, flags, {})
        self.root = tk.Tk()
        self.root.geometry(size)
        self.app = gui.RatesPlotApp(self.root, choices, notes, {}, normal_geometry=size)
        if settle:
            self.settle()

    def settle(self, timeout: float = SETTLE_SECONDS) -> None:
        """Pump Tk until no drawing runs or waits and no typed text waits to be drawn."""
        app, root = self.app, self.root
        began = time.monotonic()
        while time.monotonic() - began < timeout:
            root.update()
            if not app.rendering and app.pending is None and app._text_after_id is None:
                # A finished drawing can queue another (a Refresh asked for meanwhile): look again.
                for _ in range(5):
                    root.update()
                    time.sleep(0.02)
                if not app.rendering and app.pending is None and app._text_after_id is None:
                    return
            time.sleep(0.05)
        raise TimeoutError(f"the window did not settle in {timeout} s")

    def pump(self, times: int = 10) -> None:
        for _ in range(times):
            self.root.update()

    def tick(self, name: str, on: bool) -> None:
        """Tick or untick a box of the column (a curve, a measure, a nation, an All box) as a click does."""
        self.app.flag_vars[name].set(on)
        self.app._flag_changed(name)

    def typed(self, key: str, text: str) -> None:
        """Type into a field as a person does: its variable written, its trace firing (one write for the whole text)."""
        self.app.payload_vars[key].set(text)
        self.root.update()

    def text(self, key: str) -> str:
        return self.app.payload_vars[key].get()

    def command_tokens(self) -> list[str]:
        """The command line the window writes, without ``python ratesplot.py``."""
        return self.app.command.get().split()[2:]

    def written_draws_the_same(self, keys: tuple[str, ...] | None = None) -> bool:
        """True if the window's written command line draws, byte for byte, the charts it shows (``keys``: all shown)."""
        tokens = self.command_tokens()
        keys = tuple(self.app.pngs) if keys is None else keys
        return bool(keys) and all(self.app.pngs[key] == cli_png(tokens, key) for key in keys)

    def close(self) -> None:
        self.app._close()


def cli_png(tokens: list[str], key: str) -> bytes:
    """The chart the command line draws for ``tokens``, through the same ``build_figure``, in this process."""
    config = cli.parse_args(["--no-gui", *tokens])
    with this_thread_output_to(None):
        drawn = resolve_start(config)
        country = next(c for c in COUNTRIES if c.key == key)
        return frontend.png_bytes(build_figure(country, drawn))


@contextlib.contextmanager
def window(tokens: str, size: str = SIZE):
    """A ``Window``, closed afterwards whatever happens."""
    opened = Window(tokens, size)
    try:
        yield opened
    finally:
        opened.close()


def enabled(widget) -> bool:
    return not widget.instate(["disabled"])


def all_widgets(widget):
    yield widget
    for child in widget.winfo_children():
        yield from all_widgets(child)


# ---------------------------------------------------------------------------
# Part 1: the window's charts against the command line's
# ---------------------------------------------------------------------------


def sha256(data: bytes) -> bytes:
    return hashlib.sha256(data).digest()


def check_charts(cases: tuple[str, ...], outdir: Path, verbose: bool = False) -> bool:
    """Draw each case in the window and on the command line; True if every chart is identical and the written line draws the same."""
    all_same = True
    for name in cases:
        tokens = [END, *CASES[name]]
        began = time.monotonic()
        try:
            opened = Window(" ".join(tokens))
        except Exception as exc:  # a case the window cannot open is a failure of this check, not of the run
            print(f"{name:12} FAIL: the window did not draw it: {type(exc).__name__}: {exc}")
            all_same = False
            continue
        try:
            app = opened.app
            seconds = time.monotonic() - began
            if app.shown_config is None:
                print(f"{name:12} FAIL: not drawn: {app.status.get()}")
                all_same = False
                continue
            output = io.StringIO()
            with contextlib.nullcontext() if verbose else contextlib.redirect_stdout(output):
                written = verify_charts.run(outdir / name, tokens)
            keys = [key for key in NATION_KEYS if key in app.pngs]
            verdicts = []
            if len(keys) != len(written):
                verdicts.append(f"FAIL: the window drew {len(keys)} chart(s), the command line {len(written)}")
                all_same = False
            for key, path in zip(keys, written):
                same = sha256(app.pngs[key]) == sha256(path.read_bytes())
                verdicts.append(f"{key}:{'IDENTICAL' if same else 'DIFFER'}")
                all_same &= same
            # The line the window writes, read back as tools/check_round_trip.py reads it, draws the same charts.
            payloads, flags, unmatched, misspelled = cli.sort_tokens(opened.command_tokens())
            back = options.config_from_choices(payloads, flags)
            round_trip = not unmatched and not misspelled and all(
                options.drawn_config(back, key) == options.drawn_config(app.shown_config, key) for key in keys
            )
            verdicts.append("written line: same" if round_trip else f"FAIL: the written line draws otherwise: {' '.join(opened.command_tokens())}")
            all_same &= round_trip
            print(f"{name:12} {' '.join(CASES[name]):28} {' '.join(verdicts)}  (window {seconds:.1f} s)")
            for line in output.getvalue().splitlines():
                if line.strip().startswith("Warning:") and "requested start" not in line:
                    print(f"    {line.strip()}")
        finally:
            opened.close()
    return all_same


# ---------------------------------------------------------------------------
# Part 2: the window's behaviour
# ---------------------------------------------------------------------------


def view_fields(w: Window, key: str = "cdn") -> tuple[str, ...]:
    return tuple(w.text(name) for name in ("start", "end", f"{key}:min", f"{key}:max", f"{key}:top", f"{key}:bottom"))


def zoom_middle(w: Window, key: str = "cdn", vertical: bool = True) -> None:
    """A drag across the middle of a chart's plot area: the dates' middle 40 %, and the values' too if ``vertical``."""
    left, top, right, bottom = w.app.geometry[key].box
    y0, y1 = (top + (bottom - top) * 0.3, top + (bottom - top) * 0.7) if vertical else (top, bottom)
    w.app._zoom_to_box(key, left + (right - left) * 0.3, y0, left + (right - left) * 0.7, y1)
    w.settle()


def check_view(check: Check) -> None:
    """Batch 5 (2026-10-13): a zoom fills the fields; Unzoom and Back; the first view kept however long the history."""
    with window(f"-c -s:2001 {END}") as w:
        app = w.app
        check("at first: Back and Unzoom greyed", not enabled(app.back_button) and not enabled(app.unzoom_button))
        first = view_fields(w)
        zoom_middle(w)
        zoomed, zoomed_png = view_fields(w), app.pngs["cdn"]
        check("a zoom writes every field", all(a != b for a, b in zip(first, zoomed)), zoomed)
        check("zoomed: Back and Unzoom active", enabled(app.back_button) and enabled(app.unzoom_button))
        check("zoomed: the written line draws the chart shown", w.written_draws_the_same(("cdn",)), w.command_tokens())
        app._unzoom()
        w.settle()
        check("Unzoom: the fields as at first", view_fields(w) == first, view_fields(w))
        check("Unzoom: Unzoom greyed, Back active", enabled(app.back_button) and not enabled(app.unzoom_button))
        app._back()
        w.settle()
        check("Back after Unzoom: the zoom again", view_fields(w) == zoomed, view_fields(w))
        check("Back after Unzoom: the zoomed chart, byte for byte", app.pngs["cdn"] == zoomed_png)
        check("Back after Unzoom: Unzoom active again", enabled(app.unzoom_button))
        app._back()
        w.settle()
        check("Back again: the first view, nothing left to go back to",
              view_fields(w) == first and not enabled(app.back_button) and not enabled(app.unzoom_button))
        # The first view is kept however many steps are taken.
        for step in range(gui._HISTORY_LIMIT + 5):
            app._apply_view({"start": f"{1960 + step % 40}-01-01"}, "Zoomed")
            app._loading = False
        w.settle()
        check("the history capped, its first view kept",
              len(app.history) == gui._HISTORY_LIMIT and app.history[0][0][0]["start"] == first[0],
              (len(app.history), app.history[0][0][0]["start"]))
        app._unzoom()
        w.settle()
        check("Unzoom after many steps: the first view", w.text("start") == first[0], w.text("start"))


def check_reset(check: Check) -> None:
    """Batches 5 and 17.5: Reset sets the defaults but leaves the nation boxes, the preview's settings too; Back undoes it."""
    with window(f"-s:2015 --no-gdp --no-debt --no-interest {END} -dim:1600x900") as w:
        app = w.app
        editor = app.size_editors["dimensions"]
        w.tick("germany", True)
        w.settle()
        w.tick("us", False)
        w.settle()
        w.typed("cdn:max", "9")
        w.tick("policy", True)
        w.settle()
        nations = {name: app.flag_vars[name].get() for name in frontend.nation_boxes()}
        before = app._snapshot()
        app.zoom.set("actual")
        app._refresh_all_views()
        editor.keep_ratio.set(False)
        app._reset()
        w.settle()
        check("Reset: the nation boxes as they were", {name: app.flag_vars[name].get() for name in nations} == nations
              == {"cdn": True, "us": False, "uk": False, "germany": True}, nations)
        check("Reset: the rest at the defaults (the policy rate off, Canada's Max blank, the default size)",
              not app.flag_vars["policy"].get() and w.text("cdn:max") == "" and w.text("dimensions") == "2048x1536",
              (app.flag_vars["policy"].get(), w.text("cdn:max"), w.text("dimensions")))
        # reset_choices names every nation's every field; the window has only the controls each nation has (no prime margin for the UK).
        (payloads, flags), (reset_payloads, reset_flags) = app._snapshot(), frontend.reset_choices(nations)
        check("Reset: every control as frontend.reset_choices gives",
              flags == reset_flags and all(reset_payloads.get(key) == value for key, value in payloads.items()),
              {key: (value, reset_payloads.get(key)) for key, value in payloads.items() if reset_payloads.get(key) != value})
        check("Reset: Fit to window, the shape locked", app.zoom.get() == "fit" and editor.keep_ratio.get(),
              (app.zoom.get(), editor.keep_ratio.get()))
        # An unticked nation's last chart is kept off screen, its tab hidden: the tabs say what is shown.
        tabs = {key: app.notebook.tab(app.tabs[key], "state") for key in NATION_KEYS}
        check("Reset: Canada's and Germany's tabs shown, the others hidden",
              tabs == {"cdn": "normal", "us": "hidden", "uk": "hidden", "de": "normal"}, tabs)
        app._back()
        w.settle()
        check("Back after Reset: the choices as before", app._snapshot() == before)
        check("Back after Reset: Actual size and the lock unticked again", app.zoom.get() == "actual" and not editor.keep_ratio.get(),
              (app.zoom.get(), editor.keep_ratio.get()))
        # A zoom's step keeps no preview settings: Unzoom and Back leave the preview alone.
        zoom_middle(w, vertical=False)
        app.zoom.set("fit")
        app._refresh_all_views()
        app._unzoom()
        w.settle()
        app._back()
        w.settle()
        check("Unzoom and Back leave the preview's mode as set", app.zoom.get() == "fit", app.zoom.get())


def drawings(w: Window) -> int:
    return w.app.log.get("1.0", "end").count("--- drawing")


def check_refresh(check: Check) -> None:
    """Batch 5: one Refresh (F5) button, which draws the same choices again; one asked for during a drawing is kept."""
    with window(f"-c -s:2001 {END}") as w:
        app = w.app
        texts = [str(widget.cget("text")) for widget in all_widgets(w.root) if widget.winfo_class() == "TButton"]
        check("one Refresh (F5) button, no Redraw or Reload data", "Refresh (F5)" in texts
              and not any(text.startswith(("Redraw", "Reload")) for text in texts), texts)
        before = drawings(w)
        app._refresh()
        w.settle()
        check("Refresh draws the same choices again", drawings(w) == before + 1, drawings(w) - before)
        check("the log says so", "--- refresh:" in app.log.get("1.0", "end"))
        before = drawings(w)
        app._refresh()
        check("drawing", app.rendering)
        app._refresh()  # during that drawing: kept, and drawn after it
        w.settle()
        check("a Refresh during a drawing draws once more after it", drawings(w) == before + 2, drawings(w) - before)


def greyed(w: Window) -> set[str]:
    """Which of the controls a right-axis curve governs are greyed (batch 5's set)."""
    app = w.app
    names = {name for name in ("relative", "per-capita", "regression", "gdp") if not enabled(app.flag_boxes[name])}
    names |= {key for key in ("regression-tolerance", "us:top", "us:bottom") if not enabled(app.value_entries[key])}
    if all(not enabled(widget) for widget in app.level_editors["us:debt-parts"].widgets):
        names.add("levels")
    return names


def check_greyed(check: Check) -> None:
    """Batch 5 (Terry, 2026-10-08): what has no curve to act on is greyed, its tick or text kept, and left off the command line."""
    with window(f"-u --yield --debt --reg:2 --d:fp --us:top:40t -s:2001 {END}") as w:
        app = w.app
        first_png = app.pngs["us"]
        check("debt drawn: nothing greyed", greyed(w) == set(), greyed(w))
        w.typed("regression-tolerance", "x")
        w.settle()
        check("a bad tolerance: not drawn", app.status.get().startswith("Not drawn"), app.status.get())
        w.tick("debt", False)
        w.settle()
        every = {"relative", "per-capita", "regression", "regression-tolerance", "us:top", "us:bottom", "levels"}
        check("no right-axis curve: all greyed but GDP", greyed(w) == every, greyed(w))
        check("drawn despite the bad text in a greyed field", app.status.get().startswith("Drawn"), app.status.get())
        levels = app.level_editors["us:debt-parts"].boxes
        check("ticks and texts kept", app.flag_vars["regression"].get() and levels["f"].get() and levels["p"].get()
              and w.text("us:top") == "40t" and w.text("regression-tolerance") == "x")
        command = " ".join(w.command_tokens())
        check("the command line: --regression kept, no tolerance, levels or Top", "--regression" in command and "--reg:" not in command
              and "debt:" not in command and "top" not in command, command)
        check("drawn as the command line's yields alone", app.pngs["us"] == cli_png(["-u", "--yield", "-s:2001", END], "us"))
        w.tick("gdp", True)
        w.settle()
        check("GDP: -p, --reg, its tolerance, Top and Bottom active; -r and the levels greyed", greyed(w) == {"relative", "levels"}, greyed(w))
        w.typed("regression-tolerance", "2")
        w.tick("gdp", False)
        w.tick("debt", True)
        w.settle()
        check("debt again: the chart as at first, byte for byte", app.pngs["us"] == first_png)

    with window(f"-u -r --debt --gdp -s:2001 {END}") as w:
        app = w.app
        gdp = app.flag_boxes["gdp"]
        check("-r with debt: GDP greyed (the divisor)", not enabled(gdp))
        w.typed("us:top", "150%")
        w.settle()
        check("Top 150% drawn", app.shown_config is not None and app.shown_config.for_nation("us").macro_top == 150.0, app.status.get())
        w.tick("debt", False)
        w.tick("interest", False)
        w.settle()
        check("debt and interest off: GDP active again, -r still ticked", enabled(gdp) and app.flag_vars["relative"].get())
        check("Top in % cleared (GDP is in dollars now)", w.text("us:top") == "", w.text("us:top"))
        check("drawn: GDP in dollars, as -u --gdp", app.pngs["us"] == cli_png(["-u", "--gdp", "-s:2001", END], "us"))
        check("the command line keeps -r, and draws the same", "--relative" in w.command_tokens() and w.written_draws_the_same(("us",)),
              w.command_tokens())
        w.tick("debt", True)
        w.settle()
        check("debt again: GDP greyed, -r in effect", not enabled(gdp) and app.shown_config.relative)


class _Event:
    """What ``_image_point`` reads of a pointer event."""

    def __init__(self, x: float, y: float) -> None:
        self.x, self.y = x, y


def check_fit(check: Check) -> None:
    """Batch 5: the preview fitted to its tab, enlarged or shrunk, centred, its shape kept; the pointer mapped back to the chart."""
    with window(f"-c -s:2001 {END} -dim:1100x825") as w:
        app, key = w.app, "cdn"
        canvas = app.canvases[key]
        w.root.geometry("3400x2000+20+20")
        w.pump(5)
        app._refresh_view(key)
        scale, left, top = app.placement[key]
        image = app.images[key]
        width, height = canvas.winfo_width(), canvas.winfo_height()
        check("a 1100x825 chart in a larger tab: enlarged", scale > 1.0, (scale, width, height))
        shown_w, shown_h = round(image.width * scale), round(image.height * scale)
        check("it fills the tab one way and fits the other, centred",
              (abs(shown_w - width) <= 1 or abs(shown_h - height) <= 1) and shown_w <= width + 1 and shown_h <= height + 1
              and abs(left - (width - shown_w) // 2) <= 1 and abs(top - (height - shown_h) // 2) <= 1,
              (shown_w, shown_h, width, height, left, top))
        check("its shape kept", abs(shown_w / shown_h - image.width / image.height) < 0.01)
        # Tk rounds an event's place to a whole screen pixel, so the point is found within 1/scale.
        point = app._image_point(key, _Event(left + 400 * scale, top + 300 * scale))
        check("the pointer maps back to the chart's own pixels", point is not None
              and abs(point[0] - 400) < 1 / scale and abs(point[1] - 300) < 1 / scale, point)
        app.zoom.set("actual")
        app._refresh_all_views()
        check("Actual size: pixel for pixel", app.placement[key][0] == 1.0)
        app.zoom.set("fit")
        app._refresh_all_views()
        w.root.geometry("1100x700+20+20")
        w.pump(5)
        app._refresh_view(key)
        check("a small window: shrunk", app.placement[key][0] < 1.0, app.placement[key][0])


def styles(w: Window, key: str) -> tuple[str, str]:
    return str(w.app.value_labels[key].cget("style")), str(w.app.value_entries[key].cget("style"))


PLAIN = ("TLabel", "TEntry")
RED = ("Error.TLabel", "Error.TEntry")


def check_wrong_fields(check: Check) -> None:
    """Batch 6: a field that stops the drawing is red, label and text, the one changed last; a greyed field never."""
    with window(f"-c {END} --cdn:max:3") as w:
        app = w.app
        check("drawn at first", app.shown_config is not None, app.status.get())
        w.typed("cdn:min", "5")
        w.settle()
        check("Min typed above Max: Min's label and text red", styles(w, "cdn:min") == RED, styles(w, "cdn:min"))
        check("Max not red", styles(w, "cdn:max") == PLAIN, styles(w, "cdn:max"))
        check("the status line names Canada's Min and the rule", app.status.get().startswith("Not drawn: Canada, Min")
              and "must be less than max" in app.status.get(), app.status.get())
        w.typed("cdn:max", "5")
        w.settle()
        check("Max then typed equal: Max red (changed last), Min not", styles(w, "cdn:max") == RED and styles(w, "cdn:min") == PLAIN,
              (styles(w, "cdn:min"), styles(w, "cdn:max")))
        w.typed("cdn:max", "6")
        w.settle()
        check("put right: drawn, nothing red", not app.status.get().startswith("Not drawn")
              and styles(w, "cdn:min") == PLAIN and styles(w, "cdn:max") == PLAIN, app.status.get())
        # A zoom writes both limits, in order: it never trips the rule.
        zoom_middle(w)
        check("a zoom: drawn, nothing red", not app.status.get().startswith("Not drawn")
              and all(styles(w, k) == PLAIN for k in ("cdn:min", "cdn:max", "cdn:top", "cdn:bottom")), app.status.get())
        app._back()
        w.settle()
        w.typed("cdn:top", "0")
        w.settle()
        check("a Top of zero: red by itself", styles(w, "cdn:top") == RED, app.status.get())
        w.typed("cdn:top", "")
        w.settle()
        # The size: the lock off, then the width alone too small.
        editor = app.size_editors["dimensions"]
        editor.keep_ratio.set(False)
        editor.lock_changed()
        editor.width.set("600")
        w.settle()
        width_entry, height_entry = app.size_entries["dimensions"]
        check("the width alone too small: its box and the label red, the height's box not",
              str(width_entry.cget("style")) == "Error.TEntry" and str(height_entry.cget("style")) == "TEntry"
              and str(app.value_labels["dimensions"].cget("style")) == "Error.TLabel", app.status.get())
        editor.width.set("1500")
        w.settle()
        check("the width put right: no box red", str(width_entry.cget("style")) == "TEntry" and str(height_entry.cget("style")) == "TEntry",
              app.status.get())
        # A greyed field is never red: every right-axis curve unticked, Top keeps a bad text.
        w.typed("cdn:top", "-1")
        for name in ("debt", "gdp", "interest"):
            w.tick(name, False)
        w.settle()
        check("Top greyed with a bad text: not red, drawn", styles(w, "cdn:top") == PLAIN and not app.status.get().startswith("Not drawn"),
              app.status.get())


def hint(w: Window, key: str) -> tuple[bool, str]:
    """Whether a field's grey value is shown, and its text."""
    label = w.app.hints[key].label
    w.root.update()
    return bool(label.winfo_ismapped()), str(label.cget("text"))


def check_grey_values(check: Check) -> None:
    """Batches 7 and 8: a blank field shows the value drawn in grey, hidden by typing, back when cleared, hidden while greyed."""
    with window(f"-u --reg {END}", size="1600x1000+40+40") as w:
        w.pump()
        check("Start: the automatic start in grey", hint(w, "start") == (True, "2000-01-01"), hint(w, "start"))
        check("End typed: no grey value", not hint(w, "end")[0], hint(w, "end"))
        check("the tolerance: 0.3 in grey (batch 8)", hint(w, "regression-tolerance") == (True, "0.3"), hint(w, "regression-tolerance"))
        shown, text = hint(w, "us:max")
        check("the U.S.' Max: the value drawn", shown and float(text) > 0, (shown, text))
        check("Canada not drawn: no grey value", not hint(w, "cdn:max")[0], hint(w, "cdn:max"))
        w.typed("start", "2001")
        check("typing hides it", not hint(w, "start")[0], hint(w, "start"))
        w.settle()
        w.typed("start", "")
        check("clearing brings it back", hint(w, "start")[0], hint(w, "start"))
        w.settle()
        w.tick("regression", False)
        check("greyed (--reg off): hidden", not hint(w, "regression-tolerance")[0], hint(w, "regression-tolerance"))
        w.tick("regression", True)
        check("active again: shown", hint(w, "regression-tolerance")[0], hint(w, "regression-tolerance"))
        w.settle()
        command = " ".join(w.command_tokens())
        check("grey values never on the command line", "--start" not in command and "--reg:" not in command, command)


def check_typing(check: Check) -> None:
    """Batch 15.5: a number typed without its 0 gets it on Return or on leaving the field."""
    cases = {".3": "0.3", "-.25": "-0.25", "+.5": "+0.5", ".5t": "0.5t", "-.5%": "-0.5%", " .3": " 0.3",
             "0.3": "0.3", "3": "3", "1.5": "1.5", "": "", ".": ".", "-.": "-.", "2000-01": "2000-01", "x.5": "x.5", "800x600": "800x600"}
    wrong = {typed: frontend.with_leading_zero(typed) for typed, wanted in cases.items() if frontend.with_leading_zero(typed) != wanted}
    check(f"the rule on {len(cases)} texts", not wrong, wrong)
    with window(f"--reg {END}") as w:
        for key, typed, how, wanted in (
            ("regression-tolerance", ".3", "<Return>", "0.3"),
            ("cdn:min", "-.5", "<FocusOut>", "-0.5"),
            ("us:top", ".5t", "<FocusOut>", "0.5t"),
        ):
            entry = w.app.value_entries[key]
            entry.focus_force()
            w.root.update()
            w.typed(key, typed)
            entry.event_generate(how)
            w.root.update()
            check(f"{key}: {typed!r} then {how} gives {wanted!r}", w.text(key) == wanted, w.text(key))
            w.settle()
        command = w.command_tokens()
        check("the command line writes them with their 0", {"--reg:0.3", "--cdn:min:-0.5", "--us:top:0.5t"} <= set(command), command)


def check_save_log(check: Check) -> None:
    """Batch 15: Save log file writes the log as shown, in UTF-8 with Windows line ends (the dialog answered by the check)."""
    with window(f"-c -s:2015 {END}") as w:
        app = w.app
        app._append_log("Ünïcode → check, and a second line\n")
        target = Path(SCRATCH_APPDATA) / "saved_log.txt"
        asked: dict = {}

        def answer(**kwargs) -> str:
            asked.update(kwargs)
            return str(target)

        dialog = gui.filedialog.asksaveasfilename
        gui.filedialog.asksaveasfilename = answer
        try:
            app._save_log()
        finally:
            gui.filedialog.asksaveasfilename = dialog
        check("the dialog was asked, for a .txt file", bool(asked) and str(asked.get("defaultextension", "")) == ".txt",
              {k: v for k, v in asked.items() if k != "parent"})
        data = target.read_bytes() if target.exists() else b""
        shown = app.log.get("1.0", "end-1c")
        check("the file is the log as shown", data.decode("utf-8").replace("\r\n", "\n") == shown, len(data))
        check("Windows line ends throughout", data.count(b"\n") == data.count(b"\r\n") > 0, (data.count(b"\n"), data.count(b"\r\n")))
        check("the status line says where", str(target.name) in app.status.get(), app.status.get())


def check_export(check: Check) -> None:
    """Batch 11.1: Export chart writes the chart as each type chosen, the data as shown, refusing a redrawn image mid-drawing."""
    from openpyxl import load_workbook

    folder = Path(SCRATCH_APPDATA) / "export"
    folder.mkdir(exist_ok=True)
    with window(f"-s:2015 {END}") as w:
        app = w.app
        button = app._bar_buttons[1]
        check("the pinned bar's button: Export chart… (Ctrl+S)", str(button.cget("text")) == "Export chart… (Ctrl+S)",
              button.cget("text"))
        asked: dict = {}
        answers: list[tuple[str, str]] = []  # (the type to choose, the name to give), one per dialog

        def answer(**kwargs) -> str:
            asked.clear()
            asked.update(kwargs)
            caption, name = answers.pop(0)
            if caption:
                kwargs["typevariable"].set(caption)
            return str(folder / name) if name else ""

        def export(caption: str, name: str, how: str = "button") -> Path:
            answers.append((caption, name))
            dialog = gui.filedialog.asksaveasfilename
            gui.filedialog.asksaveasfilename = answer
            try:
                if how == "Ctrl+S":
                    # A key reaches the window only while it has the focus (as in check_typing).
                    w.root.focus_force()
                    w.root.update()
                    w.root.event_generate("<Control-s>")
                    w.pump()
                else:
                    button.invoke()
            finally:
                gui.filedialog.asksaveasfilename = dialog
                if answers:  # the dialog was not opened: its answer is dropped, so the next export's is its own
                    answers.clear()
                    asked.clear()
            return folder / name

        app.notebook.select(app.tabs["us"])
        w.pump()
        figure, drawn = app.figures["us"]
        written = export("PNG image", "chart_png", how="Ctrl+S")
        check("Ctrl+S opens the dialog: titled Export chart, the eight types, the name without an extension, PNG's",
              asked.get("title") == "Export chart" and len(asked.get("filetypes", ())) == 8
              and asked.get("initialfile") == frontend.saved_png_name("us", drawn).removesuffix(".png")
              and asked.get("defaultextension") == ".png",
              {k: v for k, v in asked.items() if k not in ("parent", "typevariable")})
        check("PNG: the extension added, the bytes drawn", written.with_suffix(".png").read_bytes() == app.pngs["us"]
              if written.with_suffix(".png").exists() else False, sorted(p.name for p in folder.iterdir()))
        check("the status line says where", "chart_png.png" in app.status.get(), app.status.get())
        starts = {"JPG image": b"\xff\xd8\xff", "PDF document": b"%PDF", "SVG drawing": b"<?xml"}
        for caption, extension, _every in gui.EXPORT_CHOICES[1:4]:
            written = export(caption, f"chart_{extension[1:]}")
            data = written.with_suffix(extension).read_bytes() if written.with_suffix(extension).exists() else b""
            check(f"{caption}: written as {extension}", data.startswith(starts[caption]), data[:8])
        written = export("CSV, comma-separated", "chart_csv")
        check("the dialog opens on the type last chosen, with its extension", app.export_type.get() == "CSV, comma-separated"
              and asked.get("defaultextension") == ".svg", asked.get("defaultextension"))
        rows = list(csv.reader(io.StringIO(written.with_suffix(".csv").read_text(encoding="utf-8-sig"))))
        data = gui.export.chart_data(figure, drawn.for_nation("us"), "us")
        check("CSV: the chart's data, its title first, a row a date", rows[0] == [figure._suptitle.get_text()]
              and rows[2][0] == "Date" and len(rows) == 4 + len(data.dates) and len(data.dates) > 1000,
              (rows[:3], len(rows), len(data.dates)))
        check("CSV: the file is export's bytes", written.with_suffix(".csv").read_bytes() == gui.export.delimited_bytes(data, ","))
        written = export("TSV, tab-separated", "chart_tsv")
        check("TSV: tab-separated", written.with_suffix(".tsv").read_bytes() == gui.export.delimited_bytes(data, "\t"))
        written = export("XLSX workbook, this chart", "chart_one")
        book = load_workbook(written.with_suffix(".xlsx"))
        check("XLSX: one sheet, the tab's name", book.sheetnames == ["United States"], book.sheetnames)
        written = export("XLSX-global workbook, every chart", "chart_all")
        book = load_workbook(written.with_suffix(".xlsx"))
        check("XLSX-global: a sheet for each chart, named as its tab", book.sheetnames == ["Canada", "United States"], book.sheetnames)
        check("XLSX-global: each sheet its chart's title", [book[name]["A1"].value for name in book.sheetnames]
              == [app.figures[key][0]._suptitle.get_text() for key in ("cdn", "us")])
        export("PNG image", "named.csv")
        check("a name's own extension wins over the type chosen", (folder / "named.csv").exists()
              and not (folder / "named.csv.png").exists(), sorted(p.name for p in folder.iterdir()))
        export("XLSX-global workbook, every chart", "named_global.csv")
        check("XLSX-global with a .csv name: this chart's CSV", (folder / "named_global.csv").read_bytes()
              == gui.export.delimited_bytes(data, ","))
        before = sorted(p.name for p in folder.iterdir())
        export("PNG image", "")
        check("the dialog cancelled: nothing written", sorted(p.name for p in folder.iterdir()) == before)
        app.rendering = True  # as while the worker draws
        try:
            export("PDF document", "drawing_pdf")
            check("while drawing: a PDF refused, the status line says why",
                  not (folder / "drawing_pdf.pdf").exists() and "being drawn" in app.status.get(), app.status.get())
            export("PNG image", "drawing_png")
            check("while drawing: a PNG written (the bytes drawn)", (folder / "drawing_png.png").exists())
            export("CSV, comma-separated", "drawing_csv")
            check("while drawing: the data written", (folder / "drawing_csv.csv").exists())
        finally:
            app.rendering = False
        # Unticking a nation hides its tab: XLSX-global then holds the charts shown.
        w.tick("cdn", False)
        w.settle()
        export("XLSX-global workbook, every chart", "shown_only")
        book = load_workbook(folder / "shown_only.xlsx")
        check("a nation unticked: XLSX-global without it", book.sheetnames == ["United States"], book.sheetnames)


def check_tabs(check: Check) -> None:
    """Batch 17.4: a nation ticked gets its tab once drawn; unticking the current one makes the first shown current."""
    with window(f"{END} -s:2015 --no-gdp --no-debt --no-interest") as w:
        app = w.app
        check("the first tab current at first", app._current_tab_key() == "cdn", app._current_tab_key())
        w.tick("uk", True)
        check("the UK ticked: its tab not current before its chart is drawn", app._current_tab_key() == "cdn", app._current_tab_key())
        w.settle()
        check("the UK drawn: its tab current", app._current_tab_key() == "uk", app._current_tab_key())
        w.tick("germany", True)
        w.settle()
        check("Germany drawn: its tab current", app._current_tab_key() == "de", app._current_tab_key())
        w.tick("germany", False)
        w.settle()
        check("the current nation unticked: the first tab current, not the nearest", app._current_tab_key() == "cdn", app._current_tab_key())
        app.notebook.select(app.tabs["us"])
        w.root.update()
        w.tick("uk", False)
        w.settle()
        check("another nation unticked: the current tab kept", app._current_tab_key() == "us", app._current_tab_key())
        w.tick("germany", True)
        w.tick("germany", False)
        w.settle()
        check("ticked and unticked before drawn: the current tab kept", app._current_tab_key() == "us", app._current_tab_key())


def check_policy_terms(check: Check) -> None:
    """Batch 9.1: each nation's panel has its policy terms, the target ticked, greyed while the policy rate is off."""
    with window(f"{END} -s:2015 --no-gdp --no-debt --no-interest") as w:
        app = w.app

        def editor(nation: str):
            return app.choice_editors[f"{nation}:policy-terms"]

        def ticked(nation: str) -> list[str]:
            return [value for value, variable in editor(nation).boxes.items() if variable.get()]

        def tick(nation: str, value: str, on: bool) -> None:
            editor(nation).boxes[value].set(on)
            editor(nation).box_changed(value)

        def drawn(nation: str) -> tuple[bool, tuple[str, ...]]:
            config = options.drawn_config(cli.parse_args(["--no-gui", *w.command_tokens()]), nation)
            return config.policy_rates, config.policy_terms

        check("each panel: target and market", all(list(editor(n).boxes) == ["target", "market"] for n in NATION_KEYS),
              {n: list(editor(n).boxes) for n in NATION_KEYS})
        check("the target ticked in each", all(ticked(n) == ["target"] for n in NATION_KEYS), {n: ticked(n) for n in NATION_KEYS})
        check("greyed while the policy rate is off", all(not enabled(widget) for n in NATION_KEYS for widget in editor(n).widgets))
        w.tick("policy", True)
        w.settle()
        check("the policy rate on: the boxes active", all(enabled(widget) for n in NATION_KEYS for widget in editor(n).widgets))
        check("the target drawn", drawn("us") == (True, ("target",)), drawn("us"))
        tick("us", "market", True)
        w.settle()
        check("the U.S.' market ticked: both there, the target alone in Canada",
              (drawn("us"), drawn("cdn")) == ((True, ("target", "market")), (True, ("target",))), (drawn("us"), drawn("cdn")))
        check("its charts the written line's", w.written_draws_the_same(), w.command_tokens())
        tick("cdn", "target", False)
        check("Canada's last box stays ticked, with the message", ticked("cdn") == ["target"] and "Policy rate" in app.status.get(),
              (ticked("cdn"), app.status.get()))
        w.settle()
        app._reset()
        w.settle()
        check("Reset: the target alone ticked, the policy rate off",
              all(ticked(n) == ["target"] for n in NATION_KEYS) and not app.flag_vars["policy"].get(), {n: ticked(n) for n in NATION_KEYS})


# Named as the legends name them since batch 18.1 (Terry, 2026-10-23: "Rename the level boxes to match the legends (Länder, State & Local)").
LEVEL_CAPTIONS = {
    "cdn": ["Federal", "Provincial", "Municipal", "Non-federal"],
    "us": ["Federal", "State", "Local", "State & Local"],
    "uk": ["Central", "Local"],
    "de": ["Federal", "Länder", "Local", "Länder & Local"],
}


def check_levels(check: Check) -> None:
    """Batch 10.5: each nation's own levels of government in its panel, named for its governments, written per nation."""
    with window(f"{END} -s:2005 --na:ca,us,gb,de --no-yield") as w:
        app = w.app

        def tick(key: str, letter: str, on: bool) -> None:
            editor = app.level_editors[f"{key}:debt-parts"]
            editor.boxes[letter].set(on)
            editor.box_changed()

        def levels_greyed(key: str) -> bool:
            return all(not enabled(widget) for widget in app.level_editors[f"{key}:debt-parts"].widgets)

        tokens = w.command_tokens
        check("no By level in the column", "debt-parts" not in app.level_editors, sorted(app.level_editors))
        for key, wanted in LEVEL_CAPTIONS.items():
            editor = app.level_editors.get(f"{key}:debt-parts")
            captions = [str(widget.cget("text")) for widget in editor.widgets] if editor else None
            check(f"{key}'s levels named for its governments", captions == wanted, captions)
        check("none ticked at first", not any(var.get() for editor in app.level_editors.values() for var in editor.boxes.values()))
        # Each panel's boxes fit its tab at Terry's two window sizes (Ambergris's, and Buddug's large screen).
        for size in ("1728x972", "3743x1931"):
            w.root.geometry(size + "+0+0")
            w.pump(30)
            for key, tab in app.tabs.items():
                app.notebook.select(tab)
                w.pump(5)
                panel = next(child for child in tab.winfo_children() if isinstance(child, ttk.LabelFrame))
                flow = app.panel_flows[key]
                used = flow.width_for(flow.columns) + gui._block_overhead(panel)
                check(f"{size}: {key}'s panel fits its tab", used <= tab.winfo_width(), (used, tab.winfo_width(), f"{flow.columns} across"))
        w.root.geometry(SIZE)
        app.notebook.select(app.tabs["cdn"])
        w.settle()
        tick("cdn", "f", True)
        tick("cdn", "p", True)
        w.settle()
        check("Canada's federal and provincial: --cdn:debt:fp", "--cdn:debt:fp" in tokens() and "--debt" in tokens(), tokens())
        for key in ("us", "de"):
            tick(key, "f", True)
            tick(key, "p", True)
        w.settle()
        check("the U.S.' and Germany's the same, the UK's none: each nation's own",
              {"--cdn:debt:fp", "--us:debt:fp", "--ger:debt:fp"} <= set(tokens()) and "--uk:debt" not in " ".join(tokens()), tokens())
        check("the UK's aggregate debts active while it has no level",
              all(enabled(widget) for widget in app.choice_editors["uk:aggregate"].widgets))
        tick("uk", "f", True)
        w.settle()
        check("the UK's central level too: --debt:fp once", "--debt:fp" in tokens()
              and not any(token.startswith(("--cdn:debt", "--uk:debt")) for token in tokens()), tokens())
        check("the UK's aggregate debts greyed by its level", all(not enabled(widget) for widget in app.choice_editors["uk:aggregate"].widgets))
        check("each chart the written line's", w.written_draws_the_same(NATION_KEYS), tokens())
        for name in ("debt", "interest"):
            w.tick(name, False)
        w.settle()
        check("no debt or interest: every panel's levels greyed, ticks kept",
              all(levels_greyed(key) for key in LEVEL_CAPTIONS) and app.level_editors["cdn:debt-parts"].boxes["f"].get())
        check("no debt or interest: the levels left out", not any("debt:" in token or "interest:" in token for token in tokens()), tokens())
        app._reset()
        w.settle()
        check("Reset: no level ticked", not any(var.get() for editor in app.level_editors.values() for var in editor.boxes.values()))

    # A command line's own levels in each panel; the UK's n is its local level.
    with window(f"{END} -s:2005 --na:gb,de --debt:fn --ger:interest:pm") as w:
        app = w.app
        ticked = {key: [letter for letter, var in app.level_editors[f"{key}:debt-parts"].boxes.items() if var.get()] for key in ("uk", "de")}
        check("--debt:fn --ger:interest:pm: the UK's Central and Local, Germany's States and Local",
              ticked == {"uk": ["f", "m"], "de": ["p", "m"]}, ticked)
        check("both charts the command line's", all(app.pngs[key] == cli_png([END, "-s:2005", "--na:gb,de", "--debt:fn", "--ger:interest:pm"], key)
                                                     for key in ("uk", "de")))


def check_limits(check: Check) -> None:
    """Batch 18: each limit's unit beside its box, and its All tick, which copies a limit to every nation (in a currency for Top and Bottom)."""
    with window(f"{END} -s:2015 --na:ca,us") as w:
        app = w.app
        tokens = w.command_tokens

        def ticked(name: str) -> bool:
            return app.flag_vars[f"all-{name}"].get()

        def all_greyed(key: str) -> bool:
            return not enabled(app.all_boxes[key])

        def unit(key: str) -> str:
            return str(app.unit_labels[key].cget("text"))

        def texts(name: str) -> list[str]:
            return [w.text(f"{key}:{name}") for key in NATION_KEYS]

        def tick_all(key: str) -> None:
            app.flag_vars[f"all-{key.split(':')[1]}"].set(True)
            app._all_ticked(key)

        labels = {name: str(app.value_labels[f"us:{name}"].cget("text")) for name in ("max", "min", "top", "bottom")}
        check("the labels: Max, Min, Top, Bottom", labels == {"max": "Max", "min": "Min", "top": "Top", "bottom": "Bottom"}, labels)
        units = {f"{key}:{name}": unit(f"{key}:{name}") for key in NATION_KEYS for name in ("max", "top")}
        check("% right of Max, the currency right of Top",
              units == {"cdn:max": "%", "cdn:top": "C$", "us:max": "%", "us:top": "$", "uk:max": "%", "uk:top": "£", "de:max": "%", "de:top": "€"},
              units)
        check("every All box greyed while its box is blank", all(all_greyed(key) for key in app.all_boxes))
        w.typed("us:max", "8")
        check("a value typed: its All box active, the others greyed", not all_greyed("us:max") and all_greyed("cdn:max"))
        tick_all("us:max")
        w.settle()
        check("All ticked: every nation's Max 8", texts("max") == ["8"] * 4, texts("max"))
        check("written once, --max:8", "--max:8" in tokens() and not any(t.startswith(("--us:", "--cdn:")) for t in tokens()), tokens())
        check("its charts the written line's", w.written_draws_the_same(("cdn", "us")))
        w.typed("cdn:max", "7")
        check("a value typed in another nation's box goes into every one", texts("max") == ["7"] * 4, texts("max"))
        w.settle()
        w.typed("cdn:max", "")
        check("blanked by typing: still ticked, its box active, the others' values kept",
              ticked("max") and not all_greyed("cdn:max") and texts("max")[1:] == ["7"] * 3)
        app._limit_settled("cdn:max")  # Return, or the box left (decision 47)
        check("the blank settled: unticked, its own All greyed, the others' active and still 7", not ticked("max") and all_greyed("cdn:max")
              and not any(all_greyed(f"{key}:max") for key in ("us", "uk", "de")) and texts("max")[1:] == ["7"] * 3)
        w.settle()
        check("written per nation now: --us:max:7, no --max", "--us:max:7" in tokens() and not any(t.startswith("--max") for t in tokens()), tokens())
        # A blank settled by a click elsewhere; a click in the box itself settles nothing.
        tick_all("us:max")
        w.typed("us:max", "")
        event = tk.Event()
        event.widget = app.canvases["us"]
        app._clicked_anywhere(event)
        check("a click on the chart settles the blank", not ticked("max") and all_greyed("us:max"))
        event.widget = app.value_entries["us:max"]
        w.typed("us:max", "6")
        w.typed("us:max", "")
        app._clicked_anywhere(event)
        check("a click in the box itself does not", "us:max" in app._blanked)
        app._limit_settled("us:max")
        # Min, then a zoom with Min's All ticked and Max's not.
        w.typed("us:max", "6")
        w.typed("us:min", "1")
        tick_all("us:min")
        w.settle()
        check("All on Min: every nation's Min 1", texts("min") == ["1"] * 4, texts("min"))
        app._apply_view({"max": "5", "min": "0.5"}, "Zoomed", "us")
        w.settle()
        check("a zoom: Min, whose All is ticked, into every nation's box", texts("min") == ["0.5"] * 4, texts("min"))
        check("a zoom: Max, whose All is not, the U.S.' alone", texts("max") == ["7", "5", "7", "7"], texts("max"))
        check("its charts the written line's", w.written_draws_the_same(("cdn", "us")), tokens())
        # Top and Bottom: an amount goes with its currency (18.3); under -r, percentages as they are.
        w.typed("us:top", "40t")
        check("Top in dollars: its All box active", not all_greyed("us:top"))
        tick_all("us:top")
        w.settle()
        check("All on Top: the U.S.' 40t as USD40t in the other boxes", texts("top") == ["USD40t", "40t", "USD40t", "USD40t"], texts("top"))
        check("--top:USD40t once", "--top:USD40t" in tokens(), tokens())
        check("its charts the written line's (Canada's Top converted)", w.written_draws_the_same(("cdn", "us")))
        log = app.log.get("1.0", "end")
        check("the log names the conversion", "Top USD40t is C$" in log, [line for line in log.splitlines() if "USD40t" in line][-2:])
        w.typed("cdn:top", "50t")
        check("Canada's 50t typed: CAD50t in the others", texts("top") == ["50t", "CAD50t", "CAD50t", "CAD50t"], texts("top"))
        w.tick("relative", True)
        w.settle()
        check("-r: Top's unit blank, Max's still %", unit("us:top") == "" and unit("us:max") == "%", (unit("us:top"), unit("us:max")))
        check("-r: Top cleared", w.text("us:top") == "")
        w.typed("us:top", "150%")
        tick_all("us:top")
        w.settle()
        check("-r: All on Top copies 150% to every nation", texts("top") == ["150%"] * 4, texts("top"))
        check("-r: --top:150% once", "--top:150%" in tokens(), tokens())
        w.tick("relative", False)
        w.settle()
        check("the measure changed: Top cleared, its All unticked; Min's All kept",
              not ticked("top") and w.text("us:top") == "" and ticked("min"), (ticked("top"), w.text("us:top"), ticked("min")))
        # One nation: every All unticked and greyed.
        w.tick("cdn", False)
        w.settle()
        check("one nation: every All unticked and greyed",
              not any(ticked(n) for n in ("max", "min", "top", "bottom")) and all(all_greyed(key) for key in app.all_boxes))
        w.tick("cdn", True)
        w.settle()
        check("two again: the boxes with values active, unticked", not all_greyed("us:min") and not ticked("min"))
        tick_all("us:min")
        app._reset()
        w.settle()
        check("Reset: no All ticked, every box greyed",
              not any(ticked(n) for n in ("max", "min", "top", "bottom")) and all(all_greyed(key) for key in app.all_boxes))
        app._back()
        w.settle()
        check("Back after Reset: Min's All ticked again", ticked("min"))

    # Started from a command line: the limit written once ticks All, where every nation drawn shares it.
    for tokens_text, wanted in (
        (f"{END} -s:2015 --na:ca,us --max:8", True),
        (f"{END} -s:2015 -u --max:8", False),
        (f"{END} -s:2015 --na:ca,us --max:8 --us:max:9", False),
    ):
        with window(tokens_text) as w:
            check(f"started with {tokens_text.split(' ', 2)[2]}: Max's All {'ticked' if wanted else 'not ticked'}",
                  w.app.flag_vars["all-max"].get() == wanted)


def check_projections(check: Check) -> None:
    """Batch 11 (decisions 68 and 72): "Projections to today", on by default, a box and a switch with no connection to the quotes'."""
    # A chart ending today, so the projection applies; --no-cur, so no quote of the moment makes two drawings differ.
    with window("-u --no-cur -s:2020") as w:
        app = w.app
        box, quotes = app.flag_boxes["projections"], app.flag_boxes["current"]

        def block(widget):
            """The control column's block a control sits in (each row of a block is a frame of its own)."""
            while widget is not None and widget.winfo_class() != "TLabelframe":
                widget = widget.master
            return widget

        check("the box: Projections to today, in the quotes' block", str(box.cget("text")) == "Projections to today"
              and block(box) is not None and block(box) is block(quotes), (box.cget("text"), str(block(box)), str(block(quotes))))
        check("ticked by default", app.flag_vars["projections"].get())
        log = app.log.get("1.0", "end")
        check("--no-cur: the curves still projected to today", "Projecting to" in log and "(--projections)" in log,
              [line for line in log.splitlines() if "Projecting" in line])
        check("the command line: --no-current alone", "--no-current" in w.command_tokens() and "--no-projections" not in w.command_tokens(),
              w.command_tokens())
        first = app.pngs["us"]
        start = app.log.index("end-1c")
        w.tick("projections", False)
        w.settle()
        check("unticked: --no-projections written", "--no-projections" in w.command_tokens(), w.command_tokens())
        check("unticked: nothing projected, the chart changed", "Projecting to" not in app.log.get(start, "end") and app.pngs["us"] != first)
        check("unticked: the written line draws the same", w.written_draws_the_same(("us",)), w.command_tokens())
        w.tick("projections", True)
        w.settle()
        check("ticked again: the first chart, byte for byte", app.pngs["us"] == first)
        # No right-axis curve: nothing to project, so greyed (batch 5's rule), its tick kept, on the command line as
        # a greyed --regression's is (a greyed field's text is what is left out).
        w.tick("projections", False)
        for name in ("debt", "gdp", "interest"):
            w.tick(name, False)
        w.settle()
        check("no right-axis curve: the box greyed, its tick kept off", not enabled(box) and not app.flag_vars["projections"].get())
        check("greyed: --no-projections kept on the command line, which draws the same",
              "--no-projections" in w.command_tokens() and w.written_draws_the_same(("us",)), w.command_tokens())


def _wait_until(w: Window, condition: Callable[[], bool], seconds: float = SETTLE_SECONDS) -> bool:
    """Pump Tk until ``condition`` holds (a worker's message handled), then until the window settles; False if it never held."""
    began = time.monotonic()
    while not condition():
        if time.monotonic() - began > seconds:
            return False
        w.root.update()
        time.sleep(0.05)
    w.settle()
    return True


def check_auto_refresh(check: Check) -> None:
    """Batch 11.3 (decision 23): "Automatic refresh", off by default; every few minutes what Refresh does; greyed where it brings nothing new."""
    # A chart ending today with the quotes on, so a refresh can bring something new: the yields alone, few downloads.
    with window("-u --yield -s:2026-01-01") as w:
        app = w.app
        box, quotes = app.flag_boxes["auto-refresh"], app.flag_boxes["current"]
        minutes = app.value_entries["auto-refresh-minutes"]

        def block(widget):
            while widget is not None and widget.winfo_class() != "TLabelframe":
                widget = widget.master
            return widget

        check("the box: Automatic refresh, in the quotes' block", str(box.cget("text")) == "Automatic refresh,"
              and block(box) is block(quotes), box.cget("text"))
        check("off by default, not greyed; its minutes greyed while it is off", not app.flag_vars["auto-refresh"].get()
              and enabled(box) and not enabled(minutes))
        check("no timer while off", app._auto_after_id is None and app._auto_minutes is None)
        before = drawings(w)
        w.tick("auto-refresh", True)
        w.settle()
        check("ticked: nothing drawn again", drawings(w) == before, drawings(w) - before)
        check("ticked: its timer runs, every 5 minutes", app._auto_minutes == 5 and app._auto_after_id is not None, app._auto_minutes)
        check("ticked: its minutes active, 5 in grey", enabled(minutes) and hint(w, "auto-refresh-minutes") == (True, "5"),
              hint(w, "auto-refresh-minutes"))
        check("the command line: --auto-refresh", "--auto-refresh" in w.command_tokens(), w.command_tokens())
        w.typed("auto-refresh-minutes", "10")
        w.settle()
        check("10 minutes: the timer at 10, nothing drawn", app._auto_minutes == 10 and drawings(w) == before, app._auto_minutes)
        # The minutes by the first spelling --help shows, as --reg:TOL is written.
        check("the command line: --au:10", "--au:10" in w.command_tokens(), w.command_tokens())
        w.typed("auto-refresh-minutes", "4")
        w.settle()
        check("4 minutes: refused, the field red, no timer", app.status.get().startswith("Not drawn")
              and styles(w, "auto-refresh-minutes")[1] == "Error.TEntry" and app._auto_minutes is None, app.status.get())
        w.typed("auto-refresh-minutes", "10")
        w.settle()

        # A tick: what Refresh does when it is due (within market hours, the Ticker stubbed), else the minutes again.
        app._auto_ticker.due = lambda config, now=None: True
        app._auto_refresh_tick()
        drew = _wait_until(w, lambda: drawings(w) == before + 1)
        log = app.log.get("1.0", "end")
        check("a tick due: drawn again as Refresh draws", drew and "--- automatic refresh, every 10 minutes" in log and "--- refresh:" in log,
              drawings(w) - before)
        check("after it the timer runs again", app._auto_after_id is not None and app._auto_minutes == 10)
        app._auto_ticker.due = lambda config, now=None: False
        if app._auto_after_id is not None:  # as when its minutes are up: the timer has fired
            app.root.after_cancel(app._auto_after_id)
        app._auto_refresh_tick()
        check("a tick not due: nothing drawn, the timer again", _wait_until(w, lambda: app._auto_after_id is not None)
              and drawings(w) == before + 1, drawings(w) - before)

        # Greyed where it can bring nothing new, its tick and minutes kept; the minutes left off the command line.
        def greyed_now() -> bool:
            return not enabled(box) and not enabled(minutes)

        w.tick("current", False)
        w.settle()
        check("quotes off: greyed, no timer, its tick and text kept", greyed_now() and app._auto_minutes is None
              and app.flag_vars["auto-refresh"].get() and w.text("auto-refresh-minutes") == "10")
        command = w.command_tokens()
        check("greyed: --auto-refresh written, its minutes not", "--auto-refresh" in command and "--au:10" not in command, command)
        w.tick("current", True)
        w.settle()
        check("quotes on again: active, every 10 minutes", not greyed_now() and app._auto_minutes == 10)
        w.typed("end", "2026-09-01")
        w.settle()
        check("an End before today: greyed", greyed_now() and app._auto_minutes is None)
        w.typed("auto-refresh-minutes", "x")
        w.settle()
        check("greyed: a bad text in its minutes stops nothing", app.status.get().startswith("Drawn"), app.status.get())
        w.typed("auto-refresh-minutes", "10")
        w.typed("end", "")
        w.settle()
        check("End blank (today) again: active", not greyed_now() and app._auto_minutes == 10)
        w.tick("debt", True)
        w.tick("yield", False)
        w.settle()
        check("no yield shown (debt alone): greyed", greyed_now())
        w.tick("spreads", True)
        w.settle()
        check("a spread shown, which runs on to the quotes: active", not greyed_now() and app._auto_minutes == 10)
        w.tick("spreads", False)
        w.tick("yield", True)
        w.tick("debt", False)
        w.settle()

        # Remembered as it stands, as the hover read-out's box is: its change draws nothing.
        app._remember()
        state, _note, _path = frontend.load_state()
        check("remembered: ticked, every 10 minutes", state.get("flags", {}).get("auto-refresh") is True
              and state.get("payloads", {}).get("auto-refresh-minutes") == "10", (state.get("flags", {}).get("auto-refresh"),
              state.get("payloads", {}).get("auto-refresh-minutes")))
        w.tick("auto-refresh", False)
        w.settle()
        check("unticked: no timer, its minutes greyed", app._auto_after_id is None and app._auto_minutes is None and not enabled(minutes))


def _made_up_chart():
    """A small chart with known lines, marked as plotting marks its own, for the read-out's rules (``hover``).

    Steps: 1 from 2020-01-01 (carried on a row of 2020-02-01), 2 from
    2020-04-01, 1.5 from 2020-07-01, held to 2020-09-01 where a point with no
    value breaks the line. A twin of the first
    step 3 px above it, 2020-01-20 to 2020-03-20. A line through 2020-01-01
    (3), 2020-01-11 (3.4), 2020-03-01 (3), quoted, with the quote on
    2020-03-01. A projection from 2020-10-01 (0.5) to 2020-10-11 (1.5). A line
    over the y axis's top (3.6), shown only up to it.
    """
    figure = Figure(figsize=(8, 6), dpi=100)
    FigureCanvasAgg(figure)
    ax = figure.subplots()
    days = pd.to_datetime
    ax.set_xlim(days("2019-12-01"), days("2020-11-01"))
    ax.set_ylim(0, 3.6)
    three_px = 3 / (ax.transData.transform((0, 1.0))[1] - ax.transData.transform((0, 0.0))[1])
    steps, = ax.plot(
        days(["2020-01-01", "2020-02-01", "2020-04-01", "2020-07-01", "2020-09-01"]), [1.0, 1.0, 2.0, 1.5, np.nan], drawstyle="steps-post"
    )
    twin, = ax.plot(days(["2020-01-20", "2020-03-20"]), [1.0 + three_px, 1.0 + three_px])
    line, = ax.plot(days(["2020-01-01", "2020-01-11", "2020-03-01"]), [3.0, 3.4, 3.0])
    projection, = ax.plot(days(["2020-10-01", "2020-10-11"]), [0.5, 1.5])
    high, = ax.plot(days(["2020-05-01", "2020-06-01"]), [3.5, 3.9])  # crosses the axis's top
    two = lambda value: f"{value:.2f}"  # noqa: E731
    for drawn, name, how in (
        (steps, "Steps", {}), (twin, "Twin", {}), (line, "Line", {"quoted": True}),
        (projection, "Projection", {"projected": True}), (high, "High", {}),
    ):
        export.mark_series(drawn, name, export.YIELD_AXIS)
        hover.mark(drawn, (lambda value: "≈" + two(value)) if how.get("projected") else two, **how)
    hover.mark_quote(figure, pd.Timestamp("2020-03-01 14:32", tz="America/Toronto"), "America/Toronto")
    figure.savefig(io.BytesIO(), format="png", dpi=figure.dpi)
    return figure, ax


def _image_point(figure: Figure, ax, when: str, value: float) -> tuple[float, float]:
    """Where a date and value fall on ``figure``'s picture (image pixels, origin top left)."""
    x, y = ax.transData.transform((mdates.date2num(pd.Timestamp(when)), value))
    return float(x), float(figure.bbox.height - y)


def _hover_texts(app) -> list[str]:
    """The texts of the window's read-out, on any tab: [] while it is not shown."""
    return [canvas.itemcget(item, "text") for canvas in app.canvases.values()
            for item in canvas.find_withtag(gui._HOVER_TAG) if canvas.type(item) == "text"]


def _pump_for(w: Window, seconds: float) -> None:
    """Pump Tk for ``seconds``, so its timers fire."""
    until = time.monotonic() + seconds
    while time.monotonic() < until:
        w.root.update()
        time.sleep(0.01)


def check_hover(check: Check) -> None:
    """Batch 11.2 (decisions 21 and 68): the hover read-out: its rules, the window's, its box and switch, the plain windows'."""
    # The rules, on a chart whose lines are known (hover.ChartLines.reading, which the page's script follows).
    figure, ax = _made_up_chart()
    lines = hover.chart_lines(figure)

    def read(when: str, value: float, dx: float = 0.0, dy: float = 0.0, reach: float = 3.0):
        x, y = _image_point(figure, ax, when, value)
        return lines.reading(x + dx, y + dy, reach)

    def says(found) -> tuple[str, str, str] | None:
        return None if found is None else (found.name, found.value, found.when)

    check("a step's flat stretch: its value, the day under the pointer, the observation's date (not a later row's)",
          says(read("2020-02-15 12:00", 1.0, dy=2)) == ("Steps", "1.00", "2020-02-15 (observed 2020-01-01)"), says(read("2020-02-15 12:00", 1.0, dy=2)))
    check("on the observation's own day: the date alone", says(read("2020-01-01 06:00", 1.0, dy=-1)) == ("Steps", "1.00", "2020-01-01"),
          says(read("2020-01-01 06:00", 1.0, dy=-1)))
    check("a riser, 2 px to its right: the new value, from its own date",
          says(read("2020-04-01", 1.5, dx=2)) == ("Steps", "2.00", "2020-04-01"), says(read("2020-04-01", 1.5, dx=2)))
    check("the last step held to the point with no value, then nothing",
          says(read("2020-08-31 12:00", 1.5)) == ("Steps", "1.50", "2020-08-31 (observed 2020-07-01)") and read("2020-09-10", 1.5) is None,
          (says(read("2020-08-31 12:00", 1.5)), says(read("2020-09-10", 1.5))))
    check("caught within 3 px (2.9 above), not at 3.1", read("2020-01-10 12:00", 1.0, dy=-2.9) is not None and read("2020-01-10 12:00", 1.0, dy=-3.1) is None)
    check("caught 2.9 px to the left of a riser", says(read("2020-04-01", 1.5, dx=-2.9)) == ("Steps", "2.00", "2020-04-01"), says(read("2020-04-01", 1.5, dx=-2.9)))
    check("a line: the observation nearer along it (30 %: the first, 70 %: the second)",
          says(read("2020-01-04", 3.12)) == ("Line", "3.00", "2020-01-01") and says(read("2020-01-08", 3.28)) == ("Line", "3.40", "2020-01-11"),
          (says(read("2020-01-04", 3.12)), says(read("2020-01-08", 3.28))))
    check("a quoted line on the quote's day: the quote's time",
          says(read("2020-03-01", 3.0)) == ("Line", "3.00", "2020-03-01 14:32 EST (intraday, provisional)"), says(read("2020-03-01", 3.0)))
    check("a projection: the value on the line that day, its end label's ≈, (projected)",
          says(read("2020-10-06 06:00", 1.0)) == ("Projection", "≈1.00", "2020-10-06 (projected)"), says(read("2020-10-06 06:00", 1.0)))
    # The twin runs 3 px above the first step: 1 px above the step is nearer it, 2 px above nearer the twin.
    check("of two lines in reach, the nearer",
          read("2020-02-15 12:00", 1.0, dy=-1).name == "Steps" and read("2020-02-15 12:00", 1.0, dy=-2).name == "Twin",
          (says(read("2020-02-15 12:00", 1.0, dy=-1)), says(read("2020-02-15 12:00", 1.0, dy=-2))))
    check("a line past the axis's top: read inside the axes, nothing where it runs above them",
          read("2020-05-04", 3.5 + 0.4 * 3 / 31) is not None and read("2020-05-25", 3.5 + 0.4 * 24 / 31) is None)
    check("outside the axes: nothing", lines.reading(lines.box[0] - 10, (lines.box[1] + lines.box[3]) / 2, 3) is None)
    origin, per_day = lines.days_to_x
    worst = max(float(np.max(np.abs(origin + run.day * per_day - run.x))) for curve in lines.curves for run in curve.runs)
    check("each point's x is its day's (the page works x out from the day)", worst < 0.01, worst)

    # A dense line: four points kept to a pixel's column, its first, last, highest and lowest.
    xs = np.array([10.0, 10.2, 10.4, 10.5, 10.6, 10.9, 11.1, 11.5])
    ys = np.array([5.0, 1.0, 7.0, 3.0, 9.0, 4.0, 2.0, 2.5])
    check("a dense line: per column of pixels its first, last, highest and lowest kept",
          hover._decimated(xs, ys).tolist() == [0, 1, 4, 5, 6, 7], hover._decimated(xs, ys).tolist())

    # The timing, on the controller alone, with a clock of its own.
    clock: list[tuple[float, Callable[[], None]]] = []
    now = [0.0]
    events: list[str] = []
    controller = hover.Controller(
        lambda ms, callback: clock.append((now[0] + ms, callback)) or (len(clock) - 1),
        lambda handle: clock.__setitem__(handle, (float("inf"), lambda: None)),
        lambda found, _x, _y: events.append(f"show {found}"),
        lambda: events.append("hide"),
    )

    def advance(ms: float) -> None:
        now[0] += ms
        for index, (due, callback) in enumerate(list(clock)):
            if due <= now[0]:
                clock[index] = (float("inf"), lambda: None)
                callback()

    controller.moved("a", 0, 0)
    advance(400)
    shown_early = list(events)
    controller.moved("b", 0, 0)  # moving along the line, still on it
    advance(150)
    check("controller: nothing at 0.4 s; at 0.5 s what the pointer is on then", shown_early == [] and events == ["show b"], events)
    controller.moved(None, 0, 0)
    advance(1900)
    controller.moved("c", 0, 0)
    advance(2500)
    check("controller: off the line under 2 s and back: kept, and follows", events == ["show b", "show c"], events)
    controller.moved(None, 0, 0)
    advance(1999)
    still = list(events)
    advance(2)
    check("controller: gone 2 s after leaving every line", still == ["show b", "show c"] and events[-1] == "hide", events)
    controller.moved("d", 0, 0)
    controller.moved(None, 0, 0)
    advance(600)
    check("controller: off the line before 0.5 s: nothing shown", events[-1] == "hide", events)

    # The window: a chart's lines read with each drawing; the read-out shown, followed and hidden.
    with window(f"-u -s:2020 {END}") as w:
        app = w.app
        box, projections = app.flag_boxes["hover"], app.flag_boxes["projections"]

        def block(widget):
            """The control column's block a control sits in (each row of a block is a frame of its own)."""
            while widget is not None and widget.winfo_class() != "TLabelframe":
                widget = widget.master
            return widget

        check("the box: Hover read-out, ticked", str(box.cget("text")) == "Hover read-out" and app.flag_vars["hover"].get(), str(box.cget("text")))
        check("in the block of Projections to today", block(box) is not None and block(box) is block(projections),
              (str(block(box)), str(block(projections))))
        check("on by default: not on the command line", not any("hover" in token for token in w.command_tokens()), w.command_tokens())
        figure_us, drawn = app.figures["us"]
        chart = app.hover_lines["us"]
        data = export.chart_data(figure_us, drawn.for_nation("us"), "us")
        names = list(dict.fromkeys(curve.name for curve in chart.curves))
        check("every curve Export chart writes has its lines read, and no other", names == list(data.names), (names, data.names))
        origin, per_day = chart.days_to_x
        worst = max(float(np.max(np.abs(origin + run.day * per_day - run.x))) for curve in chart.curves for run in curve.runs)
        check("every point on a whole day (the page's x from its day)", worst < 0.01, worst)

        # A point on the 10-year yield with no other line within 12 px, read where its line runs (and checked
        # against the exported value of that day).
        ten = next(curve for curve in chart.curves if curve.name == "10-Year Yield")
        run = max(ten.runs, key=lambda each: len(each.x))
        others = [dataclasses.replace(chart, curves=(curve,)) for curve in chart.curves if curve.name != ten.name]
        found = None
        for k in range(len(run.x) // 2, len(run.x)):
            x, y = float(run.x[k]), float(run.y[k])
            if all(other.reading(x, y, 12) is None for other in others):
                found = chart.reading(x, y, 3)
                break
        day = pd.Timestamp(hover.date_text(int(run.day[k])))
        exported = data.values[data.dates.index(day)][data.names.index("10-Year Yield")]
        check("a yield read: its value as the end label writes it, the exported value that day",
              found is not None and found.value == format_yield(exported) and found.when == hover.date_text(int(run.day[k])),
              (found, exported))
        debt = next(curve for curve in chart.curves if curve.name.startswith("Aggregate"))
        run_debt = max(debt.runs, key=lambda each: len(each.x))
        middle = len(run_debt.x) // 2
        flat_x = (float(run_debt.x[middle]) + float(run_debt.x[middle + 1])) / 2
        found_debt = chart.reading(flat_x, float(run_debt.y[middle]), 3)
        observed = hover.date_text(int(run_debt.day[middle]))
        check("the debt, a step: its value in force and its observation's date", found_debt is not None and found_debt.name == debt.name
              and (found_debt.when == observed or found_debt.when.endswith(f"(observed {observed})")), (found_debt, observed))

        # The window itself: the pointer on the 10-year yield's point, at the scale the chart is shown.
        scale, left, top = app.placement["us"]

        def pointer(ix: float, iy: float) -> types.SimpleNamespace:
            return types.SimpleNamespace(x=left + ix * scale, y=top + iy * scale)

        on_line = pointer(found.x, found.y)
        app._hovered("us", on_line)
        _pump_for(w, 0.3)
        early = _hover_texts(app)
        _pump_for(w, 0.4)
        check("the window: nothing at 0.3 s, the read-out at 0.7 s", early == [] and _hover_texts(app) == ["\n".join(found.lines)],
              (early, _hover_texts(app)))
        canvas = app.canvases["us"]
        kinds = sorted(canvas.type(item) for item in canvas.find_withtag(gui._HOVER_TAG))
        check("a box, its text and a ring on the point read", kinds == ["oval", "rectangle", "text"], kinds)
        # 2 screen px off the line: still caught at any scale; then well off every line.
        app._hovered("us", types.SimpleNamespace(x=on_line.x, y=on_line.y + 2))
        w.pump()
        check("2 screen px off the line: still read", _hover_texts(app) != [], _hover_texts(app))
        empty = pointer(chart.box[0] + 3, chart.box[1] + 3)  # the axes' top left corner: no line there
        check("the axes' corner reads no line", chart.reading(chart.box[0] + 3, chart.box[1] + 3, 3 / scale) is None)
        app._hovered("us", empty)
        _pump_for(w, 1.5)
        kept = _hover_texts(app)
        _pump_for(w, 0.8)
        check("off every line: kept at 1.5 s, gone by 2.3 s", kept != [] and _hover_texts(app) == [], (kept, _hover_texts(app)))
        app._hovered("us", on_line)
        _pump_for(w, 0.7)
        shown = _hover_texts(app)
        app._drag_start("us", on_line, "zoom")
        gone = _hover_texts(app)
        app._drag_end("us", on_line)  # a click, not a drag: nothing zoomed
        w.settle()
        check("a press of the button: the read-out goes at once", shown != [] and gone == [], (shown, gone))

        # The box: unticked, the read-out goes and nothing is drawn again; the address and the remembered settings follow.
        count = drawings(w)
        png = app.pngs["us"]
        app._hovered("us", on_line)
        _pump_for(w, 0.7)
        w.tick("hover", False)
        w.settle()
        check("unticked: the read-out gone, no drawing, the chart the same", _hover_texts(app) == [] and drawings(w) == count and app.pngs["us"] == png,
              (_hover_texts(app), drawings(w), count))
        check("unticked: --no-hover written, which draws the same", "--no-hover" in w.command_tokens() and w.written_draws_the_same(("us",)),
              w.command_tokens())
        app._hovered("us", on_line)
        _pump_for(w, 0.7)
        check("unticked: the pointer on a line shows nothing", _hover_texts(app) == [], _hover_texts(app))
        app._remember()
        state, _note, _path = frontend.load_state()
        check("unticked: remembered so", state.get("flags", {}).get("hover") is False, state.get("flags", {}).get("hover"))
        w.tick("hover", True)
        w.settle()
        app._hovered("us", on_line)
        _pump_for(w, 0.7)
        check("ticked again: shown again, still no drawing", _hover_texts(app) != [] and drawings(w) == count, (_hover_texts(app), drawings(w)))

    # The switch: --no-hover, at least "--ho"; "--h" stays argparse's --help.
    spelled = {token: cli.match_token(token) for token in ("--ho", "--hover", "--no-ho", "--no-hover")}
    check("--ho and --hover turn it on, --no-ho and --no-hover off",
          [(option.name, value) for option, value in spelled.values()] == [("hover", True), ("hover", True), ("hover", False), ("hover", False)],
          spelled)
    check("--h is not the read-out", cli.match_token("--h") is None)
    check("--no-hover: PlotConfig.hover off, the same chart", not cli.parse_args(["--no-hover"]).hover
          and cli.parse_args(["--no-hover"]).same_chart(cli.parse_args([])))

    # The plain matplotlib windows (--no-gui): matplotlib's own mouse events and timers, on a Tk canvas here.
    root = tk.Tk()
    try:
        figure, ax = _made_up_chart()
        canvas = FigureCanvasTkAgg(figure, master=root)
        canvas.get_tk_widget().pack()
        hover.attach(figure)
        canvas.draw()
        root.update()
        plain = figure._ratesplot_hover_window
        x, y = _image_point(figure, ax, "2020-02-15 12:00", 1.0)
        height = figure.bbox.height
        canvas.callbacks.process("motion_notify_event", MouseEvent("motion_notify_event", canvas, x, height - y))
        began = time.monotonic()
        while time.monotonic() - began < 0.7:
            root.update()
            time.sleep(0.01)
        check("a plain window: the read-out after 0.5 s", plain.text.get_visible()
              and plain.text.get_text() == "Steps\n1.00  2020-02-15 (observed 2020-01-01)", plain.text.get_text())
        canvas.callbacks.process("motion_notify_event", MouseEvent("motion_notify_event", canvas, lines.box[0] + 3, height - lines.box[1] - 3))
        began = time.monotonic()
        while time.monotonic() - began < 2.3:
            root.update()
            time.sleep(0.01)
        check("a plain window: gone 2 s after leaving the lines", not plain.text.get_visible())
    finally:
        root.destroy()
    # Its text and ring are drawn only over the window's picture: a chart saved from such a window is the same.
    attached, _ax = _made_up_chart()
    hover.attach(attached)
    attached.canvas.draw()
    attached._ratesplot_hover_window._show(lines.reading(*_image_point(attached, _ax, "2020-02-15 12:00", 1.0), 3), 100, 100)
    saved, clean = io.BytesIO(), io.BytesIO()
    attached.savefig(saved, format="png", dpi=attached.dpi)
    _made_up_chart()[0].savefig(clean, format="png", dpi=100)
    check("a plain window's chart saved with the read-out shown: the same bytes", saved.getvalue() == clean.getvalue())


def check_not_drawn(check: Check) -> None:
    """Batch 19.4: a chart that cannot be drawn is named on a line of its own, red and bold, and in the status line, in red."""
    real = gui.build_figure

    def failing_for_canada(country, drawn):
        if country.key == "cdn":
            raise RuntimeError("made up by verify_gui")
        return real(country, drawn)

    with window(f"-s:2015 {END}") as w:
        app = w.app

        def colour() -> str:
            return str(app.status_label.cget("foreground"))

        check("both charts drawn, the status line its usual colour", sorted(app.pngs) == ["cdn", "us"] and colour() == "", repr(colour()))
        start = app.log.index("end-1c")
        gui.build_figure = failing_for_canada
        try:
            app._refresh()
            w.settle()
        finally:
            gui.build_figure = real
        log = app.log.get(start, "end")
        lines = [line for line in log.splitlines() if line.startswith("NOT DRAWN: ")]
        ranges = app.log.tag_ranges(gui._NOT_DRAWN_TAG)
        tagged = [line for i in range(0, len(ranges), 2) for line in app.log.get(ranges[i], ranges[i + 1]).splitlines() if line]
        check("one NOT DRAWN line, naming the chart and why",
              lines == ["NOT DRAWN: the Canada chart: RuntimeError: made up by verify_gui"], lines)
        check("its traceback before it (not a source's failure)", "Traceback" in log)
        check("that line red and bold (its text tagged)", tagged == lines, tagged)
        check("the tag red and bold", app.log.tag_cget(gui._NOT_DRAWN_TAG, "foreground") == gui._ERROR_FOREGROUND
              and "bold" in str(app.log.tag_cget(gui._NOT_DRAWN_TAG, "font")))
        check("Canada's chart taken off its tab, the U.S.' drawn", sorted(app.pngs) == ["us"], sorted(app.pngs))
        check("the status line names it, in red", "Canada not drawn" in app.status.get() and colour() == gui._ERROR_FOREGROUND,
              (app.status.get(), colour()))
        app._refresh()
        w.settle()
        check("drawn again: both charts, the status line its usual colour", sorted(app.pngs) == ["cdn", "us"] and colour() == "",
              (sorted(app.pngs), colour(), app.status.get()))


def check_settings_files(check: Check) -> None:
    """Batch 6: each command window keeps its own settings file; a new one starts from the newest; old ones are tidied."""
    folder_root = tempfile.mkdtemp(prefix="ratesplot_verify_gui_sessions_")
    appdata, console_window = os.environ["APPDATA"], frontend.console_window
    os.environ["APPDATA"] = folder_root
    try:
        folder = Path(folder_root) / "ratesplot"

        def save(handle: int, mark: str) -> str | None:
            frontend.console_window = lambda: handle
            return frontend.save_state({"version": frontend.STATE_VERSION, "mark": mark})

        def load(handle: int) -> tuple[str | None, str | None, str | None]:
            frontend.console_window = lambda: handle
            state, note, path = frontend.load_state()
            return state.get("mark"), note, path.name if path else None

        check("nothing remembered: nothing read", load(0xA) == (None, None, None), load(0xA))
        check("window A saves to its own file", save(0xA, "A") is None and (folder / "gui_state_a.json").exists(),
              sorted(p.name for p in folder.iterdir()))
        check("no temporary file left", not list(folder.glob("*.tmp")), sorted(p.name for p in folder.iterdir()))
        time.sleep(0.05)
        check("a new window B starts from A's (written last)", load(0xB) == ("A", None, "gui_state_a.json"), load(0xB))
        save(0xB, "B")
        time.sleep(0.05)
        check("A again: its own, not B's newer one", load(0xA) == ("A", None, "gui_state_a.json"), load(0xA))
        check("a new window C: B's, written last", load(0xC) == ("B", None, "gui_state_b.json"), load(0xC))
        check("no command window, no shared file: the newest", load(0) == ("B", None, "gui_state_b.json"), load(0))
        save(0, "shared")
        time.sleep(0.05)
        save(0xA, "A2")
        check("no command window: the shared file, its own", load(0) == ("shared", None, "gui_state.json"), load(0))
        # A command window's file not written for over a month is removed, and a stray temporary file; the shared one never.
        old = time.time() - (frontend.STATE_FILE_MAX_AGE_DAYS + 10) * 86400
        for name in ("gui_state_d.json", "gui_state.json", "gui_state_e.1234.tmp"):
            path = folder / name
            if not path.exists():
                path.write_text('{"version": 1, "mark": "D"}', encoding="utf-8")
            os.utime(path, (old, old))
        load(0xA)
        names = sorted(p.name for p in folder.iterdir())
        check("a month-old window's file and a stray temporary file removed, the shared one kept",
              "gui_state_d.json" not in names and "gui_state_e.1234.tmp" not in names and "gui_state.json" in names, names)
        os.utime(folder / "gui_state_b.json", (old, old))
        check("this window's own file kept however old", load(0xB)[2] == "gui_state_b.json" and (folder / "gui_state_b.json").exists(), load(0xB))
    finally:
        os.environ["APPDATA"], frontend.console_window = appdata, console_window
        shutil.rmtree(folder_root, ignore_errors=True)


# The groups in the order they run; --group picks some by these names.
WINDOW_CHECKS: dict[str, Callable[[Check], None]] = {
    "view": check_view,
    "reset": check_reset,
    "refresh": check_refresh,
    "greyed": check_greyed,
    "fit": check_fit,
    "wrong_fields": check_wrong_fields,
    "grey_values": check_grey_values,
    "typing": check_typing,
    "save_log": check_save_log,
    "export": check_export,
    "tabs": check_tabs,
    "policy_terms": check_policy_terms,
    "levels": check_levels,
    "limits": check_limits,
    "projections": check_projections,
    "auto_refresh": check_auto_refresh,
    "hover": check_hover,
    "not_drawn": check_not_drawn,
    "settings_files": check_settings_files,
}


def check_window(groups: list[str]) -> bool:
    """Run the window's groups of checks; True if all passed. A group that raises is a FAIL, and the rest still run."""
    results: list[bool] = []

    def check(label: str, condition: bool, detail: object = "") -> None:
        results.append(bool(condition))
        shown = f"  [{detail}]" if not condition or detail else ""
        print(f"{'ok  ' if condition else 'FAIL'} {label}{shown}", flush=True)

    for name in groups:
        group = WINDOW_CHECKS[name]
        print(f"-- {name}: {group.__doc__.splitlines()[0]}", flush=True)
        began = time.monotonic()
        try:
            group(check)
        except Exception as exc:  # the group stops; its failure is counted and shown
            check(f"{name} ran to its end", False, f"{type(exc).__name__}: {exc}")
        print(f"   ({time.monotonic() - began:.0f} s)")
    print(f"[gui] window: {sum(results)} of {len(results)} checks passed")
    return all(results)


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--only", choices=("charts", "window"), help="run one part only")
    parser.add_argument("--cases", help=f"comma-separated regression cases for the charts part (default: {','.join(CHART_CASES)})")
    parser.add_argument("--group", help=f"comma-separated groups for the window part (default: all, {','.join(WINDOW_CHECKS)})")
    parser.add_argument("--cache", type=Path, help="download cache folder (default out/download_cache, or out/synthetic_cache)")
    parser.add_argument(
        "--synthetic", action="store_true", help="made-up data in each source's format, no network (tools/synthetic_sources.py)"
    )
    parser.add_argument("--verbose", action="store_true", help="show the command line's own output")
    parser.add_argument("--list", action="store_true", help="list the chart cases and the window's groups, and exit")
    args = parser.parse_args(argv)
    cases = tuple(args.cases.split(",")) if args.cases else CHART_CASES
    groups = args.group.split(",") if args.group else list(WINDOW_CHECKS)
    unknown = [name for name in cases if name not in CASES] + [name for name in groups if name not in WINDOW_CHECKS]
    if unknown:
        parser.error(f"unknown case or group: {', '.join(unknown)}")
    if args.list:
        for name in cases:
            print(f"{name:14} {' '.join(CASES[name])}")
        for name, group in WINDOW_CHECKS.items():
            print(f"{name:14} {group.__doc__.splitlines()[0]}")
        return 0

    if args.synthetic:
        print("[gui] SYNTHETIC data (tools/synthetic_sources.py): proves code paths, not the data")
        synthetic_sources.install_synthetic_cache(http, (args.cache or ROOT / "out" / "synthetic_cache").resolve())
    else:
        install_disk_cache(http, (args.cache or ROOT / "out" / "download_cache").resolve())
    # As run_gui does before its window: Windows then reports the screen in real pixels.
    gui._enable_windows_dpi_awareness()
    passed = True
    try:
        if args.only in (None, "charts"):
            print(f"[gui] charts: the window's drawing against the command line's, {len(cases)} case(s)")
            same = check_charts(cases, ROOT / "out" / "verify_gui", verbose=args.verbose)
            print("[gui] charts: ALL IDENTICAL" if same else "[gui] charts: SOME DIFFER")
            passed &= same
        if args.only in (None, "window"):
            print(f"[gui] window: {len(groups)} group(s)")
            passed &= check_window(groups)
    finally:
        shutil.rmtree(SCRATCH_APPDATA, ignore_errors=True)
    return 0 if passed else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
